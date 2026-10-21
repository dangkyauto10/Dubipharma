"""NSX adaptable table parser for one-row-per-product catalogues.

Detects headings by semantic aliases, not fixed A/B/C columns. Original
row/column provenance is preserved and no row is silently chosen when a
document includes multiple products. This is not an AI verdict.
"""
from __future__ import annotations

from io import BytesIO, StringIO
import csv
import hashlib
import re
import unicodedata
from typing import Any

MAX_BYTES = 30 * 1024 * 1024
MAX_TABLE_ROWS = 8_000
MAX_COLUMNS = 80


def normalize(value: Any) -> str:
    value = unicodedata.normalize("NFKD", str(value or "").replace("Đ", "D").replace("đ", "d"))
    value = "".join(c for c in value if not unicodedata.combining(c)).casefold()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", value)).strip()


SCHEMA = {
    "PRODUCT_ID": ("ma san pham", "product id", "ma thuoc noi bo", "ma hang hoa"),
    "REGISTRATION_NO": ("so dang ky", "sdk", "so dang ky luu hanh", "registration number"),
    "PRODUCT_NAME": ("ten thuoc", "ten san pham", "drug name", "trade name"),
    "SUPPLIER": ("don vi cung cap", "nha cung cap", "supplier", "distributor"),
    "BIDDER": ("ten nha thau", "don vi du thau", "bidder"),
    "REGISTRATION_HOLDER": ("co so dang ky", "don vi dang ky",
                            "chu so huu so dang ky", "marketing authorization holder"),
    "INGREDIENT": ("hoat chat", "ten hoat chat", "active ingredient", "ingredient"),
    "STRENGTH": ("ham luong", "nong do", "nong do ham luong", "strength"),
    "DOSAGE_FORM": ("dang bao che", "dosage form"),
    "ROUTE": ("duong dung", "route of administration"),
    "MANUFACTURER": ("nha san xuat", "ten nha san xuat", "manufacturer"),
    "FACTORY": ("nha may san xuat", "co so san xuat", "manufacturing site", "factory"),
    "GMP_NUMBER": ("so chung nhan gmp", "gmp certificate number"),
    "QUOTE_UNIT_PRICE": ("don gia bao gia", "don gia chao", "gia bao gia", "quoted unit price"),
    "QUOTE_CURRENCY": ("don vi tien te", "dong tien bao gia", "currency"),
}
FIELD = {normalize(alias): key for key, names in SCHEMA.items() for alias in names}
ID_FIELDS = {"PRODUCT_ID", "REGISTRATION_NO", "PRODUCT_NAME"}


def _read_grid(filename: str, raw: bytes) -> tuple[list[tuple[str, list[list[Any]]]], list[str]]:
    ext = filename.lower().rsplit(".", 1)[-1]
    tables: list[tuple[str, list[list[Any]]]] = []
    warnings: list[str] = []
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        try:
            for sh in wb.worksheets[:15]:
                values = [list(row) for row in sh.iter_rows(
                    max_row=MAX_TABLE_ROWS+1, max_col=MAX_COLUMNS, values_only=True)]
                tables.append((f"Sheet {sh.title}", values))
                if sh.max_row and sh.max_row > MAX_TABLE_ROWS:
                    warnings.append(f"{sh.title}: vượt giới hạn dòng")
        finally:
            wb.close()
    elif ext == "xls":
        import xlrd
        wb = xlrd.open_workbook(file_contents=raw)
        for sh in wb.sheets()[:15]:
            values = [sh.row_values(r, 0, min(sh.ncols, MAX_COLUMNS))
                      for r in range(min(sh.nrows, MAX_TABLE_ROWS+1))]
            tables.append((f"Sheet {sh.name}", values))
            if sh.nrows > MAX_TABLE_ROWS:
                warnings.append(f"{sh.name}: vượt giới hạn dòng")
    elif ext == "docx":
        from docx import Document
        doc = Document(BytesIO(raw))
        for ti, table in enumerate(doc.tables[:30], 1):
            values = [[c.text for c in row.cells[:MAX_COLUMNS]]
                      for row in table.rows[:MAX_TABLE_ROWS+1]]
            tables.append((f"Bảng Word {ti}", values))
            if len(table.rows) > MAX_TABLE_ROWS:
                warnings.append(f"Bảng Word {ti}: vượt giới hạn dòng")
        warnings.append("Bảng Word có thể có ô gộp; cần kiểm tra bảng gốc")
    elif ext == "csv":
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("cp1258", errors="replace")
            warnings.append("CSV không chuẩn UTF-8; cần so lại dấu tiếng Việt")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        values = list(csv.reader(StringIO(text), dialect=dialect))[:MAX_TABLE_ROWS+1]
        tables.append(("CSV", values))
        if len(values) > MAX_TABLE_ROWS:
            warnings.append("CSV vượt giới hạn dòng")
    else:
        return [], []
    return tables, warnings


def extract_product_rows(filename: str, raw: bytes, *, max_products: int = 250) -> dict:
    if not raw or len(raw) > MAX_BYTES:
        raise ValueError("Tệp rỗng hoặc quá 30MB")
    tables, warnings = _read_grid(filename, raw)
    result: list[dict] = []
    sha = hashlib.sha256(raw).hexdigest()
    for label, rows in tables:
        active: dict[str, int] = {}
        header_no = 0
        for row_no, row in enumerate(rows, 1):
            matched = {}
            for col, value in enumerate(row):
                key = FIELD.get(normalize(value))
                if key and key not in matched:
                    matched[key] = col
            if len(matched) >= 3 and matched.keys() & ID_FIELDS:
                active, header_no = matched, row_no
                continue
            if not active:
                continue
            fields, evidence = {}, {}
            for name, col in active.items():
                if col >= len(row):
                    continue
                value = row[col]
                if value is None or not str(value).strip():
                    continue
                value = str(value).strip()[:300]
                fields[name] = value
                evidence[name] = {
                    "file": filename, "sha256": sha, "source": f"{label}, hàng {row_no}, cột {col+1}",
                    "raw": value, "value": value, "verified": False,
                }
            if (fields.keys() & ID_FIELDS and
                    (fields.keys() & {"INGREDIENT", "STRENGTH", "MANUFACTURER", "FACTORY"})):
                result.append({"fields": fields, "evidence": evidence,
                               "file": filename, "sha256": sha,
                               "table": label, "header_row": header_no,
                               "row": row_no, "verified": False})
            if len(result) >= max_products:
                warnings.append("Vượt giới hạn sản phẩm; chỉ đọc một phần danh mục")
                break
        if len(result) >= max_products:
            break
    return {
        "file": filename, "sha256": sha, "candidates": result,
        "total_candidates": len(result), "warnings": warnings,
        "complete": not any("Vượt giới hạn" in x for x in warnings),
        "verified": False,
    }
