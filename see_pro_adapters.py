"""Read-only adapters for CSV, legacy XLS and Word tables.

Convert structure in memory to a common parser (never upload or write user data).
Unrecognized layouts are reported rather than interpreted as valid lots.
"""
from __future__ import annotations

from io import BytesIO, StringIO
import csv
import re
from typing import Any

from see_pro_structured import MAX_BYTES, MAX_ROWS, parse_excel_lots


def parse_tabular_document(filename: str, data: bytes) -> dict:
    extension = filename.rsplit(".", 1)[-1].lower()
    if extension == "xlsx":
        return parse_excel_lots(filename, data)
    if extension not in ("csv", "xls", "docx"):
        raise ValueError("Không phải bảng XLSX/XLS/CSV/DOCX")
    if not data or len(data) > MAX_BYTES:
        raise ValueError("Tệp rỗng hoặc vượt giới hạn 30 MB")
    from openpyxl import Workbook
    workbook = Workbook()
    first = workbook.active
    tables: list[tuple[str, list[tuple]]] = []
    caveats = []

    if extension == "csv":
        try:
            txt = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            txt = data.decode("cp1258", errors="replace")
            caveats.append("Bảng CSV có ký tự không chuẩn UTF-8; cần so lại dấu tiếng Việt")
        try:
            dialect = csv.Sniffer().sniff(txt[:8192], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        table_rows = list(csv.reader(StringIO(txt), dialect=dialect))[:MAX_ROWS+1]
        tables.append(("CSV", table_rows))
    elif extension == "xls":
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        for sh in book.sheets():
            table_rows = [sh.row_values(i) for i in range(min(sh.nrows, MAX_ROWS+1))]
            tables.append((sh.name, table_rows))
        caveats.append("XLS cũ có thể làm mất kiểu hiển thị, cần so lại ngày, số và mã có số 0 đầu")
    elif extension == "docx":
        from docx import Document
        doc = Document(BytesIO(data))
        for idx, table in enumerate(doc.tables, 1):
            table_rows = [[cell.text for cell in row.cells] for row in table.rows[:MAX_ROWS+1]]
            tables.append((f"Bang Word {idx}", table_rows))
        caveats.append("Bảng Word có thể có ô gộp hoặc hàng nhiều dòng, cần so lại vị trí gốc")

    if not tables:
        return {"file": filename, "package_id": "", "total_lots": 0,
                "rows": [], "issues": [{"type": "NO_TABLE_FOUND", "file": filename}],
                "sheets": [], "verified": False, "complete": False,
                "warnings": caveats}
    for idx, (title, source_rows) in enumerate(tables):
        clean_title = re.sub(r"[\[\]\*\?/:\\]", " ", title).strip()[:29] or f"Sheet {idx+1}"
        sheet = first if idx == 0 else workbook.create_sheet()
        sheet.title = clean_title if clean_title not in workbook.sheetnames or idx == 0 else f"S{idx+1}"
        for row in source_rows:
            normalized = []
            for value in row[:80]:
                # Treat spreadsheet source strings as data, not executable formulas.
                if isinstance(value, str) and value.startswith("="):
                    value = "'" + value
                normalized.append(value)
            sheet.append(normalized)
    output = BytesIO()
    workbook.save(output)
    virtual_filename = filename + ".xlsx"
    parsed = parse_excel_lots(virtual_filename, output.getvalue())
    parsed["file"] = filename
    parsed["warnings"] = caveats
    for item in parsed["rows"]:
        item["FILE_NGUON"] = filename
        for evidence in item.get("EVIDENCE", {}).values():
            evidence["source"] = filename
    for item in parsed["issues"]:
        item["file"] = filename
    return parsed
