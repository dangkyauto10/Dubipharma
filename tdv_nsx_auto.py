"""TDV NSX automatic intake: one batch, many documents and many products.

Document roles and evidence links are *proposals*, never GMP/registration/
authorization/price approvals. Preserve original source + SHA256 + row/page.
No external API calls, no Google Drive writes and no model inference.
"""
from __future__ import annotations

NSX_AUTO_RUNTIME_REV = "2026-10-22J"

from collections import defaultdict
from io import BytesIO
import hashlib
import json
import re
from typing import Any

from see_pro_multiformat import ACCEPTED, extract_file
from tdv_nsx_adaptive import extract_product_rows, _read_grid, FIELD, normalize as norm
from tdv_nsx_multisheet import reconcile_product_rows
from tdv_nsx_source_diff import compare_conflict_sources
from tdv_manufacturer_evidence import audit_manufacturer_evidence
from tdv_nsx_evidence_candidates import discover_evidence_candidates
from tdv_nsx_intake import KINDS, _read_one
from tdv_nsx_public_guard import detect_nonprimary_document
from tdv_tender_trace import build_reverse_dossiers
from tdv_tender_context import (
    detect_procurement_document, index_tender_documents, PUBLIC_REFERENCES,
)
from see_pro_traceability import safe_cell, _format

MAX_FILES = 24
MAX_QUOTE_ROWS = 250
IDENTIFIERS = ("PRODUCT_ID", "REGISTRATION_NO")
ROLES = {"PRODUCT", "GMP_CERT", "GMP_EXPLANATION", "REGISTRATION",
         "AUTHORIZATION", "QUOTE", "OTHER"}
STATUS_MATCH = "LIEN_KET_SO_BO_CAN_XAC_MINH"
STATUS_CONFLICT = "XUNG_DOT_DINH_DANH"
STATUS_UNKNOWN = "CHUA_DU_CHUNG_CU"
STATUS_AMBIGUOUS = "NHIEU_SAN_PHAM_CO_THE_KHOP"


def _source(record: dict, field: str) -> dict:
    return record.get("evidence", {}).get(field, {})


def _usable(record: dict, field: str) -> str:
    # A conflicting field cannot become an automatic exact-match key.
    if field in record.get("conflicts", {}):
        return ""
    return str(record.get("fields", {}).get(field, "")).strip()


def _ref(doc: dict) -> dict:
    return {"file": doc["filename"], "sha256": doc["sha256"],
            "type": doc.get("type", ""), "complete": doc.get("complete_read", False)}


def classify_document(filename: str, raw: bytes, *, enable_ocr: bool = False) -> dict:
    """Use evidence inside the document, never filename-only role labeling.

    A tie or a document that blends multiple primary categories is OTHER and
    explicitly flagged for review. Classification itself makes no claim that
    a certificate, registration or quote is genuine.
    """
    if filename.rsplit(".", 1)[-1].lower() not in ACCEPTED:
        raise ValueError("Định dạng không được hỗ trợ")
    source = extract_file(filename, raw, ocr=enable_ocr)
    tender = detect_procurement_document(source["lines"])
    if tender:
        return {
            "kind": "OTHER", "kind_label": KINDS["OTHER"],
            "filename": filename, "sha256": source["sha256"],
            "content_type": source["type"], "complete_read": bool(source["complete"]),
            "role_confidence": "TAI_LIEU_GOI_THAU",
            "role_candidates": [], "score": 0, "reasons": [tender["warning"]],
            "tender_document": tender,
            "catalogue_products_found": 0, "quotation_table_found": False,
            "_extracted_source": source, "verified": False,
        }
    nonprimary = detect_nonprimary_document(source["lines"])
    if nonprimary:
        # A government GMP list, a registration decision, a quote invitation
        # or a blank quotation template is not a supplier's original proof.
        # Preserve the source, but never turn it into a positive link.
        return {
            "kind": "OTHER", "kind_label": KINDS["OTHER"],
            "filename": filename, "sha256": source["sha256"],
            "content_type": source["type"], "complete_read": bool(source["complete"]),
            "role_confidence": "KHÔNG_PHẢI_CHỨNG_CỨ_TRỰC_TIẾP",
            "role_candidates": [], "score": 0,
            "public_document_class": nonprimary["class"],
            "reasons": [nonprimary["reason"]],
            "catalogue_products_found": 0,
            "quotation_table_found": False,
            "_extracted_source": source,
            "verified": False,
        }
    text = norm(" ".join(str(line["text"]) for line in source["lines"][:1500]))
    # Use a separate evidence reader to detect labeled field/value pairs.
    parsed = _read_one(
        filename, "OTHER", raw, enable_ocr, source_document=source,
    )
    f = parsed["fields"]
    # Single strong titles take precedence over incidental GMP / product words
    # in a long company document.
    headers = norm(" ".join(str(l["text"]) for l in source["lines"][:55]))
    def has(*phrases: str) -> bool:
        return any(norm(phrase) in text for phrase in phrases)
    def heading(*phrases: str) -> bool:
        return any(norm(phrase) in headers for phrase in phrases)

    scores = {k: 0 for k in ROLES if k != "OTHER"}
    # Explicit document-type phrases, with conservative precedence.
    if heading("giải trình gmp", "gmp explanation", "công văn giải trình gmp"):
        scores["GMP_EXPLANATION"] += 12
    elif heading("giấy chứng nhận gmp", "chứng nhận thực hành tốt sản xuất",
                 "certificate of gmp", "good manufacturing practice certificate",
                 "gmp certificate"):
        scores["GMP_CERT"] += 12
    if heading("giấy ủy quyền", "thư ủy quyền", "letter of authorization",
               "authorization letter", "power of attorney"):
        scores["AUTHORIZATION"] += 12
    if heading("giấy đăng ký lưu hành", "giấy phép lưu hành", "marketing authorization",
               "registration certificate", "giấy chứng nhận đăng ký lưu hành"):
        scores["REGISTRATION"] += 12
    if heading("báo giá", "thư báo giá", "quotation", "price quotation", "commercial offer"):
        scores["QUOTE"] += 12
    if heading("danh mục thuốc", "hồ sơ sản phẩm", "product dossier",
               "drug catalogue", "product catalog"):
        scores["PRODUCT"] += 8

    scores["GMP_CERT"] += 3 * bool(f.get("GMP_NUMBER"))
    scores["GMP_CERT"] += 2 * bool(f.get("GMP_SCOPE"))
    scores["GMP_CERT"] += 2 * bool(f.get("GMP_EXPIRY"))
    scores["AUTHORIZATION"] += 3 * bool(f.get("AUTH_EXPIRY"))
    scores["QUOTE"] += 6 * bool(f.get("QUOTE_UNIT_PRICE"))
    scores["QUOTE"] += 2 * bool(f.get("QUOTE_CURRENCY"))
    scores["QUOTE"] += 2 * bool(f.get("QUOTE_VALID_UNTIL"))
    scores["PRODUCT"] += 2 * bool(f.get("PRODUCT_NAME"))
    scores["PRODUCT"] += 3 * bool(f.get("INGREDIENT"))
    scores["PRODUCT"] += 2 * bool(f.get("STRENGTH"))
    scores["PRODUCT"] += 1 * bool(f.get("DOSAGE_FORM"))
    scores["REGISTRATION"] += 3 * bool(f.get("REGISTRATION_NO"))
    # Do not treat a plain product catalogue with SĐK as registration certificate.
    if scores["PRODUCT"] >= 7 and scores["REGISTRATION"] < 12:
        scores["REGISTRATION"] = 0
    # Real drug catalogues are commonly rectangular tables with no
    # "hồ sơ sản phẩm" title or key:value rows. Inspect semantic column
    # headings to discover the category, not filenames or user selection.
    catalogue_count = 0
    quote_table_present = False
    if source["type"] in ("xlsx", "xls", "csv", "docx"):
        catalogue = extract_product_rows(filename, raw)
        catalogue_count = catalogue["total_candidates"]
        quote_table_present = bool(_table_items(
            filename, raw, source["sha256"], "QUOTE"
        ))
        if catalogue_count:
            scores["PRODUCT"] += 12
        if quote_table_present:
            # Price-per-drug table: QUOTE is the document role, even when the
            # same table includes ingredient/strength/manufacturer columns.
            scores["QUOTE"] += 18
    if scores["GMP_EXPLANATION"] >= 12:
        scores["GMP_CERT"] = 0
    # Quotation tables can contain ingredient/strength; explicit quote wins.
    if scores["QUOTE"] >= 12 and scores["PRODUCT"] < 12:
        scores["PRODUCT"] = 0
    ranked = sorted(scores.items(), key=lambda item: item[1], reverse=True)
    winner, best = ranked[0]
    second = ranked[1][1]
    confident = best >= 7 and best - second >= 3
    role = winner if confident else "OTHER"
    candidates = [k for k, v in ranked if v >= 7 and best - v < 3]
    reasons = (
        [] if confident else
        ["Nội dung chưa đủ dấu hiệu để xác định duy nhất loại tài liệu."]
    )
    return {
        "kind": role, "kind_label": KINDS[role],
        "filename": filename, "sha256": source["sha256"],
        "content_type": source["type"],
        "complete_read": bool(source["complete"]),
        "role_confidence": "GỢI_Ý_CẦN_XÁC_MINH" if confident else "CHƯA_PHÂN_LOẠI",
        "role_candidates": candidates,
        "score": best, "reasons": reasons,
        "catalogue_products_found": catalogue_count,
        "quotation_table_found": quote_table_present,
        "_extracted_source": source,
        "verified": False,
    }


def _candidate_from_document(doc: dict) -> list[dict]:
    product_list = doc.get("catalogue_candidates") or []
    if product_list:
        return [
            {
                "fields": dict(row["fields"]),
                "evidence": dict(row["evidence"]),
                "conflicts": {},
                "file": doc["filename"], "sha256": doc["sha256"],
                "origin": f"{row['table']}, dòng {row['row']}",
                "complete_read": doc["complete_read"],
                "verified": False,
            } for row in product_list
        ]
    if "error" in doc:
        return []
    fields = doc.get("fields") or {}
    if not any(fields.get(k) for k in ("PRODUCT_ID", "REGISTRATION_NO", "PRODUCT_NAME")):
        return []
    return [{
        "fields": dict(fields), "evidence": dict(doc.get("evidence") or {}),
        "conflicts": dict(doc.get("conflicts") or {}),
        "file": doc["filename"], "sha256": doc["sha256"],
        "origin": "Trích trường có nhãn từ tài liệu",
        "complete_read": doc["complete_read"], "verified": False,
    }]


def _table_items(filename: str, raw: bytes, sha: str, kind: str) -> list[dict]:
    """Extract exact quote product rows (one table may contain many drugs)."""
    if kind != "QUOTE" or filename.rsplit(".", 1)[-1].lower() not in ("xlsx", "xls", "csv", "docx"):
        return []
    tables, _warnings = _read_grid(filename, raw)
    items = []
    truncated = False
    for title, rows in tables:
        heads: dict[str, int] = {}
        for ri, row in enumerate(rows, 1):
            matched = {FIELD[norm(v)]: i for i, v in enumerate(row) if norm(v) in FIELD}
            if "QUOTE_UNIT_PRICE" in matched and any(k in matched for k in IDENTIFIERS):
                heads = matched
                continue
            if not heads:
                continue
            fields = {}; evidence = {}
            for field, pos in heads.items():
                if pos >= len(row) or row[pos] is None or not str(row[pos]).strip():
                    continue
                value = str(row[pos]).strip()[:350]
                fields[field] = value
                evidence[field] = {
                    "file": filename, "sha256": sha,
                    "source": f"{title}, hàng {ri}, cột {pos+1}",
                    "raw": value, "value": value, "verified": False,
                }
            if "QUOTE_UNIT_PRICE" in fields and any(fields.get(k) for k in IDENTIFIERS):
                items.append({
                    "fields": fields, "evidence": evidence, "conflicts": {},
                    "filename": filename, "sha256": sha,
                    "kind": kind, "complete_read": True,
                    "origin": f"{title}, hàng {ri}", "verified": False,
                })
            if len(items) >= MAX_QUOTE_ROWS:
                truncated = True
                break
        if truncated:
            break
    if truncated:
        # We have not proved the rest of the source contains no further rows.
        # All extracted price evidence is still usable for review, not
        # automatically linkable/complete.
        for item in items:
            item["complete_read"] = False
    return items


def _match(product: dict, proof: dict) -> dict:
    """Conservative one-product/one-evidence comparison with source anchors."""
    k = proof["kind"]
    pf, df = product["fields"], proof["fields"]
    anchors = {
        "GMP_CERT": ("FACTORY",), "GMP_EXPLANATION": ("FACTORY",),
        "QUOTE": IDENTIFIERS, "REGISTRATION": ("REGISTRATION_NO",),
        "AUTHORIZATION": (*IDENTIFIERS, "MANUFACTURER"),
    }.get(k, ())
    agreed = []; different = []
    for field in anchors:
        pv, dv = _usable(product, field), _usable(proof, field)
        if pv and dv:
            (agreed if norm(pv) == norm(dv) else different).append(field)
    if k in ("GMP_CERT", "GMP_EXPLANATION"):
        required = "FACTORY" in agreed
    elif k == "REGISTRATION":
        required = "REGISTRATION_NO" in agreed
    elif k == "AUTHORIZATION":
        required = "MANUFACTURER" in agreed and any(x in agreed for x in IDENTIFIERS)
    elif k == "QUOTE":
        required = any(x in agreed for x in IDENTIFIERS)
    else:
        required = False
    if different:
        status = STATUS_CONFLICT
    elif not required:
        status = STATUS_UNKNOWN
    elif not product.get("complete_read") or not proof.get("complete_read"):
        status = "CAN_KIEM_TRA_BAN_GOC"
    else:
        status = STATUS_MATCH
    return {
        "product_id": product["uid"], "product": pf.get("PRODUCT_NAME", ""),
        "source_product": product["file"], "source_evidence": proof["filename"],
        "evidence_kind": k, "evidence_origin": proof.get("origin", ""),
        "match_fields": agreed, "conflict_fields": different,
        "status": status,
        "source_sha256": proof["sha256"],
        "evidence": dict(proof.get("evidence") or {}),
        "reason": ("Trường định danh không khớp: " + ", ".join(different))
                  if different else ("Trùng định danh nguồn, vẫn cần xác minh chứng cứ gốc"
                  if status == STATUS_MATCH else "Thiếu định danh chung xác thực"),
        "verified": False, "gmp_confirmed": False,
        "quote_approved": False, "bid_eligible": False,
    }


def analyze_auto(files: list[tuple[str, bytes]], *, enable_ocr: bool = True) -> dict:
    """One click: detect roles; discover products; link proofs only when unique."""
    if not files or len(files) > MAX_FILES:
        raise ValueError("Chọn từ 1 đến 24 tệp trong một bộ hồ sơ")
    classified = []; detailed = []; original = []
    warnings = []
    for i, (name, raw) in enumerate(files):
        if not isinstance(raw, bytes):
            raise ValueError("Tệp tải lên không có dạng nhị phân hợp lệ")
        try:
            meta = classify_document(name, raw, enable_ocr=enable_ocr)
            kind = meta["kind"]
            # Never require users to pre-label types or choose product rows.
            # Uncertain labels go to the review queue, not arbitrary guesses.
            meta["kind_label"] = KINDS[kind]
            extracted_source = meta.pop("_extracted_source")
            details = _read_one(
                name, kind, raw, enable_ocr,
                source_document=extracted_source,
            )
            meta["complete_read"] = details["complete_read"]
            original.append((name, raw))
            detailed.append(details)
            classified.append(meta)
            warnings.extend(
                f"{name}: {w}" for w in details["warnings"]
                if "phải chọn đúng một dòng" not in w
            )
        except Exception as exc:
            sha = hashlib.sha256(raw).hexdigest()
            classified.append({
                "kind": "OTHER", "kind_label": KINDS["OTHER"],
                "filename": name, "sha256": sha, "complete_read": False,
                "role_confidence": "KHÔNG_ĐỌC_ĐƯỢC",
                "role_candidates": [], "reasons": [type(exc).__name__],
                "verified": False,
            })
            detailed.append({"kind": "OTHER", "filename": name, "sha256": sha,
                             "error": type(exc).__name__})
            original.append((name, raw))
            warnings.append(f"{name}: không đọc được ({type(exc).__name__})")
    products = []
    evidence_items = []
    incomplete_attachments = []
    for classification, document, (filename, raw) in zip(classified, detailed, original):
        kind = classification["kind"]
        if "error" in document:
            continue
        # Registration dossiers can be the ONLY product-identifying source in
        # a supplier batch. Do not require a separately uploaded catalogue.
        # Promote only a same-document, non-conflicting labeled NAME + SĐK,
        # with both source citations; never derive identities from filenames.
        if kind in ("PRODUCT", "REGISTRATION"):
            for p in _candidate_from_document(document):
                if kind == "REGISTRATION":
                    fs, evidence = p["fields"], p["evidence"]
                    if (not fs.get("PRODUCT_NAME") or not fs.get("REGISTRATION_NO")
                        or "PRODUCT_NAME" in p["conflicts"]
                        or "REGISTRATION_NO" in p["conflicts"]
                        or "PRODUCT_NAME" not in evidence
                        or "REGISTRATION_NO" not in evidence):
                        continue
                    p["origin"] = "Tên thuốc và SĐK trong giấy đăng ký lưu hành: " + (
                        str(evidence["PRODUCT_NAME"]["source"]) + " / " +
                        str(evidence["REGISTRATION_NO"]["source"]))
                    p["registration_document_source"] = True
                p["uid"] = "P" + str(len(products) + 1).zfill(4)
                products.append(p)
        if kind in ("GMP_CERT", "GMP_EXPLANATION", "REGISTRATION", "QUOTE", "AUTHORIZATION"):
            rows = _table_items(filename, raw, classification["sha256"], kind)
            if rows and len(rows) >= MAX_QUOTE_ROWS and kind == "QUOTE":
                classification["complete_read"] = False
                document["complete_read"] = False
                warnings.append(
                    f"{filename}: bảng báo giá có thể còn dòng ngoài giới hạn "
                    f"{MAX_QUOTE_ROWS}; kết quả chưa được coi là đọc đủ"
                )
                outstanding_warning = {
                    "file": filename, "kind": kind, "origin": "",
                    "status": "CHUA_DOC_DU_BANG_BAO_GIA",
                    "candidate_products": [],
                    "reason": "Bảng báo giá chạm giới hạn phân tích; không tự xác nhận các dòng thiếu.",
                }
                # Review queue is built below; retain this warning in
                # a separate list, then merge into review queue.
                incomplete_attachments.append(outstanding_warning)
            if rows:
                for price_row in rows:
                    price_row["complete_read"] = (
                        price_row.get("complete_read", True)
                        and document["complete_read"]
                    )
            evidence_items.extend(rows or [{
                **document, "filename": filename, "kind": kind,
                "origin": "Các trường có nhãn trong tệp",
            }])
    # Product candidates extracted from different tabs of a SINGLE source
    # are not necessarily distinct drugs. Reconcile only with an independent,
    # corroborating manufacturer/site/registration anchor. Preserve every
    # original field and its worksheet cell as source_rows.
    products, catalogue_audit = reconcile_product_rows(products)
    for i, p in enumerate(products, 1):
        p["uid"] = "P" + str(i).zfill(4)
    if catalogue_audit["excluded_ghost_rows"]:
        warnings.append(
            f'{len(catalogue_audit["excluded_ghost_rows"])} dòng tiêu đề/số thứ tự '
            'bị loại khỏi danh mục thuốc; vẫn có chứng cứ gốc trong kiểm toán.')
    if catalogue_audit["merged_secondary_rows"]:
        warnings.append(
            f'Đã gộp {catalogue_audit["merged_secondary_rows"]} dòng lặp giữa '
            'các sheet CÙNG TỆP có ít nhất một định danh NSX/nhà máy trùng. '
            'Mọi giá trị khác nhau vẫn được lưu và đánh dấu cần xác minh.')
    if catalogue_audit["ambiguous_name_groups"]:
        warnings.append(
            f'{catalogue_audit["ambiguous_name_groups"]} nhóm trùng tên nhưng '
            'không đủ điều kiện xác định duy nhất; KHÔNG tự chọn/gộp.')
    links = []
    outstanding = list(incomplete_attachments)
    for product in products:
        if product.get("conflicts"):
            outstanding.append({
                "file": product["file"], "kind": "PRODUCT",
                "origin": product["origin"],
                "status": "MAU_THUAN_GIUA_CAC_SHEET",
                "candidate_products": [product["uid"]],
                "reason": "Trường khác nhau: " +
                          ", ".join(sorted(product["conflicts"])) +
                          ". Giữ cả hai giá trị và chứng cứ gốc; chưa được "
                          "chọn một giá trị để kết luận đạt/không đạt.",
            })
    for name in catalogue_audit["review_names"]:
        ids = [p["uid"] for p in products
               if norm(p["fields"].get("PRODUCT_NAME")) == norm(name)]
        outstanding.append({
            "file": "", "kind": "PRODUCT", "origin": "",
            "status": "TRUNG_TEN_CHUA_GHEP",
            "candidate_products": ids,
            "reason": "Một tên xuất hiện nhiều lần trong cùng sheet hoặc thiếu "
                      "định danh độc lập; không ghép ngẫu nhiên các dòng nguồn.",
        })
    for proof in evidence_items:
        # A registration certificate is the source of the identity it provides,
        # not an independent second document proving that identity. Exclude
        # the source product itself from matching and do not create an empty
        # warning when no other distinct product can be matched.
        self_sourced = (proof["kind"] == "REGISTRATION" and any(
            p.get("registration_document_source")
            and p["sha256"] == proof["sha256"]
            for p in products))
        assessments = [_match(p, proof) for p in products
                       if not (self_sourced
                               and p["sha256"] == proof["sha256"])]
        candidates = [x for x in assessments if x["status"] == STATUS_MATCH]
        if len(candidates) == 1:
            links.append(candidates[0])
        elif len(candidates) > 1:
            # Never choose an arbitrary one when multiple products share
            # the same factory or identification number.
            outstanding.append({
                "file": proof["filename"], "kind": proof["kind"],
                "origin": proof.get("origin", ""),
                "status": STATUS_AMBIGUOUS,
                "candidate_products": [m["product_id"] for m in candidates],
                "reason": "Nhiều thuốc có cùng định danh trong tài liệu chứng cứ.",
            })
        else:
            if self_sourced:
                # Source registration already contributes an unverified
                # product; no self-confirmation and no spurious mismatch.
                continue
            # Mismatching identifiers of unrelated products are EXPECTED in
            # a multi-product dossier. They are not automatically data errors.
            # Only surface "conflict" if a common anchor agrees but another
            # identifying anchor contradicts it, or for the sole product.
            conflicting = [
                a for a in assessments if a["status"] == STATUS_CONFLICT
                and (a["match_fields"] or len(products) == 1)
            ]
            outstanding.append({
                "file": proof["filename"], "kind": proof["kind"],
                "origin": proof.get("origin", ""),
                "status": STATUS_CONFLICT if conflicting else STATUS_UNKNOWN,
                "candidate_products": [],
                "reason": "Không tìm được liên kết duy nhất có chứng cứ đủ mạnh.",
            })
    for c in classified:
        if c.get("public_document_class"):
            outstanding.append({
                "file": c["filename"], "kind": c["kind"],
                "origin": "", "status": "TAI_LIEU_THAM_KHAO",
                "candidate_products": [],
                "reason": "Phân loại " + c["public_document_class"] + ": "
                          + "; ".join(c["reasons"]),
            })
        if not c.get("complete_read", False) and c["role_confidence"] != "KHÔNG_ĐỌC_ĐƯỢC":
            outstanding.append({
                "file": c["filename"], "kind": c["kind"],
                "origin": "", "status": "CAN_KIEM_TRA_DO_PHU",
                "candidate_products": [],
                "reason": "Nguồn chưa đọc đủ hoặc có giới hạn trích xuất.",
            })
        if c["role_confidence"] in ("CHƯA_PHÂN_LOẠI", "KHÔNG_ĐỌC_ĐƯỢC"):
            outstanding.append({
                "file": c["filename"], "kind": c["kind"],
                "origin": "", "status": "CAN_PHAN_LOAI_LAI",
                "candidate_products": [], "reason": "; ".join(c["reasons"]),
            })
    tender_documents = [
        {"filename": c["filename"], "sha256": c["sha256"],
         "complete_read": c["complete_read"], **c["tender_document"]}
        for c in classified if c.get("tender_document")
    ]
    tender_index = index_tender_documents(tender_documents)
    for doc in tender_index["unassigned"]:
        outstanding.append({
            "file": doc["filename"], "kind": "OTHER", "origin": "",
            "status": doc["status"], "candidate_products": [],
            "reason": "Không tự ghép chứng cứ qua các gói: " + doc["reason"],
        })
    if not products:
        warnings.append(
            "Không có danh mục/sản phẩm đủ định danh để liên kết chứng cứ; "
            "không tự tạo thuốc hoặc ghép GMP/báo giá"
        )
        outstanding.append({
            "file": "", "kind": "PRODUCT", "origin": "",
            "status": "CHUA_NHAN_DIEN_DANH_MUC",
            "candidate_products": [],
            "reason": "Chưa có sản phẩm được nhận diện có nguồn kiểm chứng.",
        })
    hashes = defaultdict(list)
    for meta in classified:
        hashes[meta["sha256"]].append(meta["filename"])
    for digest, names in hashes.items():
        if len(names) > 1:
            warnings.append("Tệp có nội dung trùng SHA-256: " + ", ".join(names))
    source_candidates = discover_evidence_candidates(products, classified, detailed)
    report = {
        "mode": "NSX_AUTO_MULTI_DOCUMENT_READ_ONLY",
        "tender_documents": tender_documents,
        "tender_index": tender_index,
        "files": classified, "sources": detailed,
        "products": products, "links": links, "review_queue": outstanding,
        "catalogue_audit": catalogue_audit,
        "evidence_candidates": source_candidates,
        "warnings": warnings,
        "summary": {
            "files": len(files),
            "parsed": sum(x["role_confidence"] != "KHÔNG_ĐỌC_ĐƯỢC" for x in classified),
            "reference_documents": sum(bool(x.get("public_document_class")) for x in classified),
            "products": len(products),
            "source_product_rows": catalogue_audit["source_rows"],
            "reconciled_rows": catalogue_audit["merged_secondary_rows"],
            "excluded_ghost_rows": len(catalogue_audit["excluded_ghost_rows"]),
            "ambiguous_name_groups": catalogue_audit["ambiguous_name_groups"],
            "linked": len(links),
            "registration_hints": source_candidates["summary"]["registration_hints"],
            "factory_source_groups": source_candidates["summary"]["factory_source_groups"],
            "unresolved_evidence": source_candidates["summary"]["unresolved_evidence"],
            "review": len(outstanding),
        },
        "verified": False, "gmp_confirmed": False,
        "registration_confirmed": False, "quote_approved": False,
        "bid_eligible": False, "cloud_write_enabled": False,
    }
    report["traceability"] = build_reverse_dossiers(report)
    return report


def export_auto_xlsx(report: dict) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    overview = wb.active
    overview.title = "TU_DONG_NHAN_DIEN"
    overview.append(["Tệp", "Loại nhận diện", "Độ chắc chắn", "SHA256", "Đọc đủ", "Cần duyệt"])
    for d in report["files"]:
        overview.append([safe_cell(x) for x in (
            d["filename"], d["kind_label"], d["role_confidence"],
            d["sha256"], d["complete_read"],
            "; ".join(d.get("reasons", [])),
        )])
    rows = wb.create_sheet("DANH_MUC_THUOC")
    rows.append(["Mã dòng", "Tên thuốc", "Mã sản phẩm", "SĐK", "Hoạt chất",
                 "Hàm lượng", "NSX", "Nhà máy", "Tệp", "Vị trí", "Xác minh"])
    for p in report["products"]:
        f = p["fields"]
        rows.append([safe_cell(x) for x in (
            p["uid"], f.get("PRODUCT_NAME"), f.get("PRODUCT_ID"),
            f.get("REGISTRATION_NO"), f.get("INGREDIENT"), f.get("STRENGTH"),
            f.get("MANUFACTURER"), f.get("FACTORY"),
            p["file"], p["origin"],
            "MÂU THUẪN NGUỒN - CHƯA" if p.get("conflicts") else "CHƯA",
        )])
    tab_review = wb.create_sheet("GOP_DONG_NHIEU_SHEET")
    tab_review.append(["Mã dòng", "Tên thuốc", "Tệp gốc", "Sheet và vị trí",
                       "Số nguồn", "Trường khác nhau", "Trạng thái"])
    for p in report["products"]:
        sources = p.get("source_rows") or []
        for original in sources:
            tab_review.append([safe_cell(v) for v in (
                p["uid"], p["fields"].get("PRODUCT_NAME", ""),
                original["file"], original["origin"], len(sources),
                ", ".join(sorted(p.get("conflicts") or {})), "CHƯA XÁC MINH",
            )])
    pairs_sheet = wb.create_sheet("MAU_THUAN_GIA_TRI_GOC")
    headers = [
        "Mã NSX", "Tên thuốc", "Trường cần đối chiếu",
        "Giá trị nguồn 1", "Tệp nguồn 1", "Ô nguồn 1", "SHA256 nguồn 1",
        "Giá trị nguồn 2", "Tệp nguồn 2", "Ô nguồn 2", "SHA256 nguồn 2",
        "Kết luận",
    ]
    pairs_sheet.append(headers)
    for row in compare_conflict_sources(report.get("products") or []):
        pairs_sheet.append([safe_cell(row.get(h, "")) for h in headers])
    links = wb.create_sheet("LIEN_KET_CHUNG_CU")
    links.append(["Mã dòng", "Thuốc", "Loại chứng cứ", "Nguồn", "Vị trí", "Khớp trường",
                  "Trạng thái", "Giải thích", "Đã xác minh"])
    for item in report["links"]:
        links.append([safe_cell(x) for x in (
            item["product_id"], item["product"], KINDS[item["evidence_kind"]],
            item["source_evidence"], item["evidence_origin"],
            item["match_fields"], item["status"], item["reason"], "CHƯA",
        )])
    inspect = wb.create_sheet("CAN_XAC_MINH")
    inspect.append(["Tệp", "Loại", "Vị trí", "Trạng thái", "Ứng viên", "Lý do"])
    for item in report["review_queue"]:
        inspect.append([safe_cell(x) for x in (
            item["file"], KINDS.get(item["kind"], item["kind"]),
            item["origin"], item["status"], item["candidate_products"], item["reason"],
        )])
    suggestions = report.get("evidence_candidates") or {}
    registration_review = wb.create_sheet("GOI_Y_SDK_CHUA_XAC_MINH")
    registration_review.append([
        "Mã ứng viên", "Tên trên giấy SĐK", "SĐK trích từ giấy",
        "Nguồn danh mục", "Vị trí tên trong danh mục", "Nguồn giấy SĐK",
        "Vị trí tên trong giấy", "SHA256 giấy", "Trường cùng khớp",
        "Trường mâu thuẫn", "Trạng thái", "Xác minh",
    ])
    for item in suggestions.get("registrations") or []:
        registration_review.append([safe_cell(v) for v in (
            item["product_id"], item["product_name"], item["registration_no_in_proof"],
            item["product_source"]["file"], item["product_source"]["location"],
            item["document_source"]["file"], item["document_source"]["location"],
            item["document_sha256"], ", ".join(item["supporting_fields"]),
            ", ".join(item["conflict_fields"]), item["status"], "CHUA",
        )])
    facility_review = wb.create_sheet("GMP_THEO_NHA_MAY")
    facility_review.append([
        "Nhà máy trích trong GMP", "Nguồn GMP", "Vị trí gốc", "SHA256",
        "Số GMP ghi trong giấy", "Phạm vi ghi trong giấy", "Ngày hết hạn ghi",
        "Cảnh báo theo lịch", "Mã sản phẩm liên quan (không phải đã đạt)",
        "Số sản phẩm", "Trạng thái", "Đã xác minh",
    ])
    for item in suggestions.get("factories") or []:
        facility_review.append([safe_cell(v) for v in (
            item["factory"], item["document_file"], item["factory_source"]["location"],
            item["document_sha256"], item["gmp_number_in_proof"],
            item["gmp_scope_in_proof"], item["expiry_in_proof"],
            item["expiry_observation"], ", ".join(item["related_product_ids"]),
            item["related_product_count"], item["status"], "CHUA",
        )])
    missing_review = wb.create_sheet("CHUNG_CU_CHUA_GHEP")
    missing_review.append(["Loại", "Tệp", "SHA256", "Tình trạng", "Nguyên nhân"])
    for item in suggestions.get("unresolved") or []:
        missing_review.append([safe_cell(v) for v in (
            item["type"], item["file"], item["sha256"],
            item["status"], item["reason"],
        )])
    evidence_audit = wb.create_sheet("KIEM_TRA_CHUNG_CU_NSX")
    audit_headers = [
        "Tệp chứng cứ", "Loại chứng cứ gợi ý", "Nhà máy trong giấy",
        "SĐK trong giấy", "Số GMP trong giấy", "Phạm vi GMP trong giấy",
        "Hết hạn GMP ghi trên giấy", "Ứng viên liên kết", "Số liên kết sơ bộ",
        "Vị trí trích xuất", "SHA256 tệp", "Điểm cần kiểm tra", "Xác minh pháp lý",
    ]
    evidence_audit.append(audit_headers)
    for item in audit_manufacturer_evidence(report)["rows"]:
        evidence_audit.append([safe_cell(item.get(h, "")) for h in audit_headers])
    proof_hints = wb.create_sheet("GOI_Y_NHAN_SDK_GMP")
    proof_hints.append([
        "Tệp", "Loại tài liệu", "Trường", "Giá trị gợi ý", "Trang/ô nguồn",
        "Trích dẫn nguyên văn", "Cách đọc", "SHA256", "Trạng thái",
    ])
    for source in report.get("sources") or []:
        for hint in source.get("proof_hints") or []:
            proof_hints.append([safe_cell(v) for v in (
                source.get("filename", ""),
                KINDS.get(source.get("kind", ""), source.get("kind", "")),
                hint.get("field", ""), hint.get("value", ""),
                hint.get("source", ""), hint.get("raw", ""),
                hint.get("reading_method", ""), hint.get("sha256", ""),
                hint.get("status", "GOI_Y_CAN_SO_BAN_GOC"),
            )])
    diag_sheet = wb.create_sheet("CHAN_DOAN_DOC_SDK_GMP")
    diag_sheet.append([
        "Tệp nguồn", "Loại tài liệu gợi ý", "Trang", "Dòng chữ đọc được",
        "Dòng OCR", "Trường chưa trích chắc chắn",
        "Trường chưa tìm được gợi ý", "Số gợi ý có nguồn",
        "Trạng thái", "SHA256 tệp", "Xác minh pháp lý",
    ])
    for source in report.get("sources") or []:
        diagnosis = source.get("proof_diagnostic") or {}
        if not diagnosis:
            continue
        diag_sheet.append([safe_cell(value) for value in (
            source.get("filename", ""), KINDS.get(source.get("kind", ""), ""),
            diagnosis.get("pages", 0), diagnosis.get("source_lines", 0),
            diagnosis.get("ocr_lines", 0),
            diagnosis.get("fields_without_extracted_value", ""),
            diagnosis.get("fields_without_hint", ""),
            diagnosis.get("hint_count", 0), diagnosis.get("status", ""),
            source.get("sha256", ""), "CHUA_XAC_MINH",
        )])
    context_sheet = wb.create_sheet("DOAN_OCR_CAN_SO")
    context_sheet.append([
        "Tệp PDF", "Loại tài liệu", "Vị trí trang", "Từ khóa phát hiện",
        "Dòng văn bản gốc", "Dòng trước cùng trang", "Dòng sau cùng trang",
        "Cách đọc", "SHA256", "Trạng thái", "Xác minh",
    ])
    for source in report.get("sources") or []:
        for item in source.get("proof_contexts") or []:
            context_sheet.append([safe_cell(value) for value in (
                source.get("filename", ""),
                KINDS.get(source.get("kind", ""), source.get("kind", "")),
                item.get("source", ""), item.get("keywords", ""),
                item.get("line", ""), item.get("previous_line", ""),
                item.get("next_line", ""), item.get("reading_method", ""),
                source.get("sha256", ""), item.get("status", ""),
                "CHUA_XAC_MINH",
            )])
    rejected_sheet = wb.create_sheet("CAU_THANH_PHAN_CAN_SO")
    rejected_sheet.append([
        "Tệp PDF", "Trường phát hiện", "Văn bản chưa chuẩn hóa",
        "Vị trí nguồn", "Trích gốc", "Lý do chưa dùng", "SHA256 tệp",
        "Trạng thái",
    ])
    for source in report.get("sources") or []:
        for item in source.get("rejected_scalar_candidates") or []:
            rejected_sheet.append([safe_cell(v) for v in (
                source.get("filename", ""), item.get("field", ""),
                item.get("value", ""), item.get("source", ""),
                item.get("raw", ""), item.get("reason", ""),
                item.get("sha256", ""), "CHUA_XAC_MINH",
            )])
    trace = wb.create_sheet("TRUY_XUAT_TUNG_TRUONG")
    trace.append(["Tệp", "Vai trò", "Trường", "Giá trị", "Trang hoặc ô nguồn", "SHA256", "Trích đoạn"])
    for p in report["products"]:
        # Export ALL observed fields from all source tabs, including conflicting
        # technical values (never silently drop a supplementary sheet).
        for source_row in p.get("source_rows") or [{
            "file": p["file"], "evidence": p.get("evidence") or {},
        }]:
            for field, item in source_row["evidence"].items():
                trace.append([safe_cell(x) for x in (
                    item.get("file", source_row["file"]), "PRODUCT", field,
                    item.get("value", ""), item.get("source", ""),
                    item.get("sha256", ""), item.get("raw", ""),
                )])
    for discarded in report.get("catalogue_audit", {}).get("excluded_ghost_rows", []):
        for field, item in discarded.get("evidence", {}).items():
            trace.append([safe_cell(x) for x in (
                discarded["file"], "LOAI_TIEU_DE_GIA", field,
                item.get("value", ""), item.get("source", ""),
                item.get("sha256", ""), item.get("raw", ""),
            )])
    for d in report["sources"]:
        for field, item in d.get("evidence", {}).items():
            trace.append([safe_cell(x) for x in (
                d["filename"], d["kind"], field, item["value"], item["source"],
                item["sha256"], item["raw"],
            )])
    # Multi-line quotation prices belong to individual drug rows, not to a
    # document-level scalar field. Preserve each exact quote cell as evidence.
    for link in report["links"]:
        for field, item in link.get("evidence", {}).items():
            trace.append([safe_cell(x) for x in (
                link["source_evidence"], link["evidence_kind"], field,
                item["value"], item["source"], item["sha256"], item["raw"],
            )])
    # Bid dossier: source-anchored reverse tracing is mandatory for leadership.
    # These review categories are not automatically interpreted as statutory
    # requirements until this project's own E-HSMT is parsed and confirmed.
    reverse = report.get("traceability") or build_reverse_dossiers(report)
    leaders = wb.create_sheet("TRUY_CUU_LANH_DAO")
    leaders.append(["Mã thuốc nội bộ", "Tên thuốc", "SĐK nguồn", "Đơn vị cung cấp",
                    "Nhà thầu", "Đơn vị đứng tên đăng ký", "Nhà sản xuất",
                    "Cơ sở / nhà máy", "Tệp nguồn", "Vị trí nguồn",
                    "SHA256 nguồn", "Trạng thái xác minh"])
    for product in reverse["products"]:
        party = product["parties"]
        leaders.append([safe_cell(v) for v in (
            product["product_id"], product["product_name"], product["registration_no"],
            party["SUPPLIER"]["value"], party["BIDDER"]["value"],
            party["REGISTRATION_HOLDER"]["value"], party["MANUFACTURER"]["value"],
            party["FACTORY"]["value"], product["origin_file"], product["origin"],
            product["source_sha256"], "CHUA_XAC_MINH",
        )])
    dossier_sheet = wb.create_sheet("GIAY_TO_THEO_THUOC")
    dossier_sheet.append(["Mã", "Thuốc", "SĐK", "Đơn vị cung cấp", "Nhà sản xuất",
                          "Nhà máy", "Loại chứng cứ cần rà soát",
                          "Có chứng cứ gắn thuốc?", "Điều kiện E-HSMT",
                          "Tên tệp", "Vị trí", "SHA256", "Trích gốc", "Xác minh"])
    for item in reverse["checklist"]:
        dossier_sheet.append([safe_cell(v) for v in (
            item["product_id"], item["product_name"], item["registration_no"],
            item["supplier"], item["manufacturer"], item["factory"], item["label"],
            item["status"], item["requirement_status"], item["file"],
            item["location"], item["sha256"], item["excerpt"], "CHUA",
        )])
    party_sheet = wb.create_sheet("CHU_THE_VA_CHUNG_CU")
    party_sheet.append(["Mã", "Thuốc", "Vai trò đơn vị", "Tên ghi trong nguồn",
                        "Trạng thái", "Tệp", "Vị trí", "SHA256", "Xác minh"])
    for product in reverse["products"]:
        for item in product["parties"].values():
            party_sheet.append([safe_cell(v) for v in (
                product["product_id"], product["product_name"], item["label"],
                item["value"], item["status"], item["file"], item["source"],
                item["sha256"], "CHUA",
            )])
    tender = wb.create_sheet("TAI_LIEU_GOI_THAU")
    tender.append(["Tệp", "Loại nguồn", "Mã gói BP đọc từ nội dung",
                   "Mã TBMT IB", "Số mã G1 tham chiếu", "SHA256",
                   "Đọc đầy đủ", "Xác minh"])
    for item in report.get("tender_documents", []):
        tender.append([safe_cell(v) for v in (
            item["filename"], item["kind"], ", ".join(item["package_ids"]),
            ", ".join(item["notice_ids"]), len(item["lot_codes"]),
            item["sha256"], item["complete_read"], "CHUA",
        )])
    public_sources = wb.create_sheet("NGUON_TRA_CUU_CONG_KHAI")
    public_sources.append(["Cơ quan / nguồn", "URL", "Dùng để tra", "Giới hạn"])
    for item in PUBLIC_REFERENCES:
        public_sources.append([safe_cell(v) for v in (
            item["name"], item["url"], item["purpose"], item["authority"],
        )])
    disclaim = wb.create_sheet("GIOI_HAN_KIEM_TOAN")
    disclaim.append(["Giới hạn", "Ý nghĩa"])
    for k, v in (
        ("Phân loại", "Gợi ý tự động, không thay thế xác minh hồ sơ gốc"),
        ("Đối chiếu", "Chỉ liên kết khi định danh nguồn đủ duy nhất"),
        ("GMP", "Chưa xác thực cơ quan cấp, phạm vi hoặc hiệu lực"),
        ("SĐK", "Chưa xác minh trên cơ sở dữ liệu quản lý dược chính thức"),
        ("Giá", "Giá báo không phải giá dự thầu hoặc giá được duyệt"),
        ("Drive", "Chức năng ghi hoàn toàn khóa, chỉ xuất tệp về máy theo yêu cầu"),
    ):
        disclaim.append([k, v])
    for sh in wb:
        sh.freeze_panes = "B2"
        _format(sh, status_field="Trạng thái" if sh.title in
                ("LIEN_KET_CHUNG_CU", "CAN_XAC_MINH") else None)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
