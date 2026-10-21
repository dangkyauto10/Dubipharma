"""TDV QC7 leadership: real E-HSMT↔NSX comparison and evidence report.

Only compares data extracted within the SAME Streamlit session.
No demo fallback, external API, disk/cloud writes or bid/compliance approvals.
An ingredient match is NOT a technical match; incomplete fields are reported.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import datetime, timedelta, timezone
from io import BytesIO
import re
import unicodedata

LOT = re.compile(r"^G\d+[.．]\d{2,4}$", re.I)
VN = timezone(timedelta(hours=7))


def norm(value) -> str:
    txt = unicodedata.normalize("NFKD", str(value or "").replace("Đ", "D").replace("đ", "d"))
    txt = "".join(c for c in txt if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", txt.casefold())).strip()


CRITERIA = (
    ("HOAT_CHAT", "Hoạt chất", "HOAT_CHAT_HSMT", "INGREDIENT"),
    ("HAM_LUONG", "Hàm lượng / nồng độ", "NONG_DO_HAM_LUONG", "STRENGTH"),
    ("DANG_BAO_CHE", "Dạng bào chế", "DANG_BAO_CHE", "DOSAGE_FORM"),
    ("DUONG_DUNG", "Đường dùng", "DUONG_DUNG", "ROUTE"),
)
GREEN = "DAT_SO_BO"
YELLOW = "CHUA_DU_CAN_CU"
RED = "KHONG_KHOP_CAN_XAC_MINH"


def _comparable(value: str) -> str:
    """Only normalize conservative typography and simple single-amount units."""
    import decimal
    # Preserve the decimal separator; text normalization would erase it.
    raw = re.sub(r"\s+", "", str(value or "").lower()).replace(",", ".")
    # Only exactly a single mass unit; never change salt/base or dosage form.
    m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*(mcg|ug|mg|g)", raw)
    if m:
        amount = decimal.Decimal(m.group(1).replace(",", "."))
        mg = amount * {"mcg": decimal.Decimal("0.001"), "ug": decimal.Decimal("0.001"),
                       "mg": 1, "g": 1000}[m.group(2)]
        return "mg:" + str(mg.normalize())
    return norm(value)


def _field_evidence(lot: dict, field: str) -> dict:
    e = (lot.get("EVIDENCE") or {}).get(field)
    if isinstance(e, dict) and e.get("source"):
        return {"file": e.get("source", lot.get("FILE_NGUON", "")),
                "location": f'{e.get("sheet", "")}, dòng {e.get("row", "")}, cột {e.get("column", "")}',
                "raw": e.get("raw", ""), "sha256": lot.get("SHA256", "")}
    g = lot.get("EVIDENCE_KY_THUAT") or {}
    e = g.get(field) if isinstance(g, dict) else None
    if isinstance(e, dict):
        return {"file": lot.get("FILE_NGUON", ""),
                "location": "Trang " + ", ".join(map(str, e.get("pages") or [])),
                "raw": " ".join(map(str, e.get("raw_tokens") or [])),
                "sha256": lot.get("SHA256", "")}
    return {"file": lot.get("FILE_NGUON", ""),
            "location": str(lot.get("TRANG_VI_TRI") or ""),
            "raw": str(lot.get("TRICH_DAN") or ""), "sha256": lot.get("SHA256", "")}


def compare_sources(ehsmt: dict | None, nsx: dict | None) -> dict:
    """E-HSMT is the sole standard; NSX sources are per-field candidates.

    Evidence matches = preliminary, NEVER sufficient proof of bid eligibility.
    Missing E-HSMT fields do not count as passes or verified failures.
    """
    ehsmt = ehsmt or {}
    nsx = nsx or {}
    lots = {}
    duplicate = set()
    for row in ehsmt.get("rows") or []:
        code = str(row.get("MA_THUOC") or "").strip().upper()
        if not LOT.fullmatch(code):
            continue
        if code in lots:
            duplicate.add(code)
        else:
            lots[code] = row
    p = nsx.get("products") or []
    traces = nsx.get("traceability") or {}
    detailed = {x["product_id"]: x for x in traces.get("products") or []}
    product_list = []
    for item in p:
        product_list.append({
            "id": item["uid"], "fields": item.get("fields") or {},
            "evidence": item.get("evidence") or {},
            "trace": detailed.get(item["uid"], {}),
            "filename": item.get("file", ""), "sha256": item.get("sha256", ""),
        })
    comparison = []
    aggregate = {"green": 0, "yellow": 0, "red": 0, "evaluable": 0,
                 "total_criteria": 0, "lots_with_candidates": 0}
    for code, lot in sorted(lots.items()):
        ingredient = str(lot.get("HOAT_CHAT_HSMT") or "").strip()
        compatible = [item for item in product_list if
                      ingredient and item["fields"].get("INGREDIENT") and
                      _comparable(ingredient) == _comparable(item["fields"]["INGREDIENT"])]
        strength = str(lot.get("NONG_DO_HAM_LUONG") or "")
        narrowed = [item for item in compatible
                    if strength and item["fields"].get("STRENGTH") and
                    _comparable(strength) == _comparable(item["fields"]["STRENGTH"])]
        candidates = narrowed if narrowed else compatible
        if candidates:
            aggregate["lots_with_candidates"] += 1
        chosen = candidates[0] if len(candidates) == 1 and code not in duplicate else None
        criteria = []
        counts = {"green": 0, "yellow": 0, "red": 0, "evaluable": 0}
        for field_id, label, lhs_key, rhs_key in CRITERIA:
            expected = str(lot.get(lhs_key) or "").strip()
            actual = str(chosen["fields"].get(rhs_key) or "").strip() if chosen else ""
            right_proof = ((chosen.get("evidence") or {}).get(rhs_key) or {}) if chosen else {}
            left_proof = _field_evidence(lot, lhs_key)
            evidence_complete = bool(left_proof.get("file") and right_proof.get("source"))
            if not expected or not actual or not chosen or code in duplicate or not evidence_complete:
                status = YELLOW
            elif _comparable(expected) == _comparable(actual):
                status = GREEN
            else:
                status = RED
            counts[{GREEN:"green", YELLOW:"yellow", RED:"red"}[status]] += 1
            counts["evaluable"] += int(bool(expected))
            criteria.append({
                "field": field_id, "label": label,
                "ehsmt": expected, "nsx": actual,
                "status": status, "ehsmt_evidence": left_proof,
                "nsx_evidence": {
                    "file": chosen["filename"] if chosen else "",
                    "location": right_proof.get("source", ""),
                    "raw": right_proof.get("raw", ""),
                    "sha256": right_proof.get("sha256", chosen["sha256"] if chosen else ""),
                },
                "note": "Có trích dẫn, CHƯA xác minh chuyên môn" if status == GREEN
                        else ("Thông số nguồn chưa khớp, cần so bản gốc"
                              if status == RED else "Thiếu dữ liệu/chứng cứ hoặc nhiều ứng viên"),
            })
        for key in aggregate:
            if key in counts:
                aggregate[key] += counts[key]
        aggregate["total_criteria"] += len(CRITERIA)
        pct = round(counts["green"] * 100 / len(CRITERIA), 1)
        raw_trace = chosen["trace"] if chosen else {}
        state = ("KHONG_KHOP_SO_BO" if counts["red"] else
                 "CHUA_DU_CAN_CU" if counts["yellow"] else
                 "KHOP_KY_THUAT_SO_BO_CHUA_DU_DIEU_KIEN")
        comparison.append({
            "MA_THUOC_EHSMT": code,
            "HOAT_CHAT_EHSMT": ingredient,
            "HAM_LUONG_EHSMT": strength,
            "MA_PHAN_LO": str(lot.get("MA_PHAN_LO") or ""),
            "NGUON_EHSMT": str(lot.get("FILE_NGUON") or ""),
            "VI_TRI_EHSMT": str(lot.get("TRANG_VI_TRI") or ""),
            "TRICH_DAN_EHSMT": str(lot.get("TRICH_DAN") or "")[:1000],
            "TRANG_THAI_DOI_CHIEU": state,
            "SO_UNG_VIEN_NSX": len(candidates),
            "MA_THUOC_NSX": chosen["id"] if chosen else "",
            "TEN_THUOC_NSX": chosen["fields"].get("PRODUCT_NAME", "") if chosen else "",
            "SO_DANG_KY": chosen["fields"].get("REGISTRATION_NO", "") if chosen else "",
            "HOAT_CHAT_NSX": chosen["fields"].get("INGREDIENT", "") if chosen else "",
            "HAM_LUONG_NSX": chosen["fields"].get("STRENGTH", "") if chosen else "",
            "DON_VI_CUNG_CAP": raw_trace.get("parties", {}).get("SUPPLIER", {}).get("value", ""),
            "NHA_SAN_XUAT": raw_trace.get("parties", {}).get("MANUFACTURER", {}).get("value", ""),
            "NHA_MAY": raw_trace.get("parties", {}).get("FACTORY", {}).get("value", ""),
            "NGUON_THUOC_NSX": chosen["filename"] if chosen else "",
            "SHA256_NSX": chosen["sha256"] if chosen else "",
            "SO_TIEU_CHI_DAT": counts["green"],
            "SO_TIEU_CHI_CHUA_DAT": counts["yellow"],
            "SO_TIEU_CHI_KHONG_KHOP": counts["red"],
            "TY_LE_KHOP_SO_BO": pct,
            "criteria": criteria,
            "DA_XAC_MINH": "CHUA",
            "DAP_UNG_THAU": "CHUA_DU_CO_SO_KET_LUAN",
        })
    warnings = []
    if not lots:
        warnings.append("Chưa có danh mục E-HSMT đọc được trong cùng phiên.")
    if not product_list:
        warnings.append("Chưa có danh mục thuốc NSX có chứng cứ trong cùng phiên.")
    if duplicate:
        warnings.append(f"Có {len(duplicate)} mã E-HSMT bị trùng, không tự chọn bản.")
    warnings.extend(ehsmt.get("warnings") or [])
    warnings.extend(nsx.get("warnings") or [])
    return {
        "created_at_vn": datetime.now(VN).strftime("%d/%m/%Y %H:%M:%S +07:00"),
        "mode": "REAL_SOURCES_ONLY_NO_DEMO", "comparison": comparison,
        "summary": {
            "ehsmt_lots": len(lots), "nsx_products": len(product_list),
            "nsx_files": len(nsx.get("files") or []),
            "nsx_review": len(nsx.get("review_queue") or []),
            "green_criteria_preliminary": aggregate["green"],
            "yellow_criteria": aggregate["yellow"],
            "red_mismatch_preliminary": aggregate["red"],
            "total_criteria": aggregate["total_criteria"],
            "matching_percent_preliminary": round(
                100 * aggregate["green"] / aggregate["total_criteria"], 1)
                if aggregate["total_criteria"] else 0,
            "criteria_ehsmt_available": aggregate["evaluable"],
            "lots_with_candidate": aggregate["lots_with_candidates"],
            "fully_verified": 0,
        }, "traceability": traces,
        "ehsmt_documents": ehsmt.get("documents") or [],
        "nsx_documents": nsx.get("files") or [],
        "warnings": warnings, "verified": False, "bid_eligible": False,
        "gmp_approved": False, "requirements_checked": False,
        "note": ("Màu xanh = khớp SƠ BỘ 2 nguồn, vàng = thiếu hoặc mơ hồ, "
                 "đỏ = số liệu chưa khớp phải kiểm tra; % = tỷ lệ tiêu chí trích xuất "
                 "khớp sơ bộ trong tổng 4 tiêu chí kỹ thuật mỗi mã. "
                 "Không được xem % là mức đủ điều kiện dự thầu."),
    }

def _safe(value):
    from see_pro_traceability import safe_cell
    return safe_cell(value)


def export_leadership_xlsx(report: dict) -> bytes:
    from openpyxl import Workbook
    from see_pro_traceability import _format
    wb = Workbook()
    title = wb.active
    title.title = "TONG_QUAN_LANH_DAO"
    title.append(["Thông tin", "Giá trị / ghi chú"])
    title.append(["Ngày báo cáo (VN)", report["created_at_vn"]])
    title.append(["Loại báo cáo", "Nguồn thật trong phiên; không DEMO"])
    for key, value in report["summary"].items():
        title.append([key, value])
    title.append(["Trạng thái pháp lý", "CHUA XAC MINH — KHONG KET LUAN DU DIEU KIEN"])
    compare = wb.create_sheet("DOI_CHIEU_EHSMT_NSX")
    rows = report["comparison"]
    cols = list(rows[0]) if rows else [
        "MA_THUOC_EHSMT", "TRANG_THAI_DOI_CHIEU", "DAP_UNG_THAU"
    ]
    compare.append(cols)
    for row in rows:
        compare.append([_safe(row.get(k)) for k in cols])
    trace = report.get("traceability") or {}
    party = wb.create_sheet("TRUY_XUAT_DOANH_NGHIEP")
    party.append(["Mã thuốc NSX", "Tên thuốc", "SĐK", "Vai trò", "Đơn vị",
                  "Trạng thái", "Tệp", "Vị trí nguồn", "SHA256"])
    for p in trace.get("products") or []:
        for detail in p.get("parties", {}).values():
            party.append([_safe(z) for z in (
                p["product_id"], p["product_name"], p["registration_no"],
                detail["label"], detail["value"], detail["status"],
                detail["file"], detail["source"], detail["sha256"],
            )])
    proof = wb.create_sheet("CHUNG_CU_THEO_SAN_PHAM")
    proof.append(["Mã NSX", "Tên thuốc", "Giấy tờ cần rà soát",
                  "Có nguồn trực tiếp?", "Yêu cầu theo E-HSMT",
                  "Tệp", "Vị trí nguồn", "SHA256", "Trích dẫn"])
    for x in trace.get("checklist", []) or []:
        proof.append([_safe(z) for z in (
            x["product_id"], x["product_name"], x["label"], x["status"],
            x["requirement_status"], x["file"], x["location"], x["sha256"],
            x["excerpt"],
        )])
    # Leadership must see both ORIGINAL data values and source locators for
    # each criterion; row colours are visual summaries, not legal approval.
    evidence_sheet = wb.create_sheet("HAI_NGUON_THEO_TIEU_CHI")
    evidence_sheet.append([
        "Mã E-HSMT", "Thuốc NSX", "Tiêu chí", "E-HSMT — chuẩn",
        "NSX — nội dung", "Kết quả sơ bộ", "Tệp E-HSMT", "Trang / dòng E-HSMT",
        "Tệp NSX", "Trang / dòng NSX", "SHA256 NSX", "Lý do", "Kết luận pháp lý",
    ])
    for row in rows:
        for crit in row.get("criteria", []):
            left, right = crit["ehsmt_evidence"], crit["nsx_evidence"]
            evidence_sheet.append([_safe(v) for v in (
                row["MA_THUOC_EHSMT"], row["TEN_THUOC_NSX"],
                crit["label"], crit["ehsmt"], crit["nsx"],
                crit["status"], left.get("file", ""), left.get("location", ""),
                right.get("file", ""), right.get("location", ""),
                right.get("sha256", ""), crit["note"], "CHUA_XAC_MINH",
            )])
    sources = wb.create_sheet("DANH_MUC_TAI_LIEU")
    sources.append(["Hệ", "Tệp", "Loại", "SHA256", "Đọc đủ", "Xác minh"])
    for doc in report.get("ehsmt_documents", []):
        sources.append([_safe(z) for z in ("E-HSMT", doc.get("Tệp"),
                        doc.get("Loại nghiệp vụ"), doc.get("SHA256"),
                        doc.get("Đọc toàn bộ"), "CHUA")])
    for doc in report.get("nsx_documents", []):
        sources.append([_safe(z) for z in ("NSX", doc.get("filename"),
                        doc.get("kind"), doc.get("sha256"),
                        doc.get("complete_read"), "CHUA")])
    alerts = wb.create_sheet("CANH_BAO_KHUYEN_NGHI")
    alerts.append(["STT", "Nội dung", "Xử lý / giới hạn"])
    for i, warn in enumerate(report.get("warnings", []), 1):
        alerts.append([i, _safe(warn), "Tra hồ sơ gốc trước khi kết luận"])
    alerts.append([len(report.get("warnings", []))+1,
                   "Không có kết luận thuốc đủ điều kiện dự thầu.",
                   "Đối chiếu E-HSMT hiện hành và xác minh với người có thẩm quyền."])
    for sheet in wb.worksheets:
        _format(sheet, status_field="TRANG_THAI_DOI_CHIEU"
                if sheet.title == "DOI_CHIEU_EHSMT_NSX" else None)
    from openpyxl.styles import PatternFill, Font
    visual = {
        GREEN: ("E7F7ED", "176B41"),
        YELLOW: ("FFF4D5", "90620C"),
        RED: ("FDE9E9", "AF3030"),
    }
    for record in evidence_sheet.iter_rows(min_row=2):
        status = str(record[5].value)
        if status not in visual:
            continue
        bg, fg = visual[status]
        for item in record[:6]:
            item.fill = PatternFill("solid", fgColor=bg)
            item.font = Font(name="Calibri", size=10, color=fg, bold=True)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def export_leadership_docx(report: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.add_heading("BÁO CÁO LÃNH ĐẠO — TDV QC7 ONLINE", 0)
    doc.add_paragraph("Đối chiếu E-HSMT – NSX và truy xuất chứng cứ theo hồ sơ thật")
    doc.add_paragraph("Lập lúc: " + report["created_at_vn"])
    doc.add_heading("1. Tổng quan dữ liệu", level=1)
    labels = {
        "ehsmt_lots": "Mã E-HSMT nhận diện",
        "nsx_products": "Thuốc NSX có định danh nguồn",
        "nsx_files": "Tài liệu NSX đã nạp",
        "nsx_review": "Mục NSX cần kiểm tra",
        "green_criteria_preliminary": "Tiêu chí khớp SƠ BỘ (chưa xác minh)",
        "yellow_criteria": "Tiêu chí thiếu dữ liệu/chứng cứ",
        "red_mismatch_preliminary": "Tiêu chí không khớp SƠ BỘ, cần xác minh",
        "matching_percent_preliminary": "Tỷ lệ khớp sơ bộ (%), KHÔNG phản ánh đủ điều kiện",
        "fully_verified": "Mã được xác minh đáp ứng hoàn toàn",
    }
    for key, label in labels.items():
        doc.add_paragraph(f'{label}: {report["summary"].get(key,0)}',style="List Bullet")
    doc.add_heading("2. Kết quả đối chiếu và vấn đề tồn tại",level=1)
    doc.add_paragraph(report["note"])
    doc.add_paragraph("Không tự kết luận ĐẠT/ĐỦ ĐIỀU KIỆN từ các phép so sánh sơ bộ.")
    for row in report["comparison"][:40]:
        doc.add_paragraph(
            f'{row["MA_THUOC_EHSMT"]} — {row["HOAT_CHAT_EHSMT"]}: '
            f'{row["TRANG_THAI_DOI_CHIEU"]}',
            style="List Bullet",
        )
    if len(report["comparison"])>40:
        doc.add_paragraph("Các mã còn lại và chứng cứ đầy đủ nằm trong file Excel kèm theo.")
    doc.add_heading("3. Đối chiếu song song E-HSMT và tài liệu NSX", level=1)
    doc.add_paragraph(
        "Màu của mỗi dòng tiêu chí là trạng thái SƠ BỘ, không phải kết luận đủ "
        "điều kiện dự thầu. Giới hạn 20 mã hiển thị trực tiếp trong Word; "
        "Excel kèm theo lưu toàn bộ mã và cả hai nguồn chứng cứ."
    )
    from docx.shared import RGBColor
    colour = {
        GREEN: RGBColor(23, 107, 65),
        YELLOW: RGBColor(144, 98, 12),
        RED: RGBColor(175, 48, 48),
    }
    for drug in report["comparison"][:20]:
        doc.add_heading(drug["MA_THUOC_EHSMT"], level=2)
        t = doc.add_table(rows=1, cols=4)
        t.style = "Light Shading Accent 1"
        for cell, heading in zip(t.rows[0].cells, (
            "Tiêu chí", "E-HSMT — yêu cầu", "NSX — đối chiếu", "Kết quả",
        )):
            cell.text = heading
        for detail in drug.get("criteria", []):
            cells = t.add_row().cells
            for cell, value in zip(cells, (
                detail["label"], detail["ehsmt"] or "Chưa đọc được",
                detail["nsx"] or "Chưa có", detail["status"],
            )):
                cell.text = value
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = colour[detail["status"]]
            refs = t.add_row().cells
            left, right = detail["ehsmt_evidence"], detail["nsx_evidence"]
            refs[0].text = "Chứng cứ gốc"
            refs[1].text = f'{left.get("file","")} {left.get("location","")}'
            refs[2].text = f'{right.get("file","")} {right.get("location","")}'
            refs[3].text = "CHƯA XÁC MINH"
    for message in report["warnings"]:
        doc.add_paragraph(message, style="List Bullet")
    doc.add_heading("4. Kiến nghị cho Lãnh đạo",level=1)
    for message in (
        "Không phê duyệt đáp ứng E-HSMT chỉ dựa trên phép so trùng hoạt chất/hàm lượng.",
        "Đọc đầy đủ E-HSMT, phiên bản sửa đổi/làm rõ và yêu cầu hồ sơ của đúng gói.",
        "Đối chiếu gốc SĐK, GMP, nhà sản xuất, nhà máy, ủy quyền và báo giá theo từng thuốc.",
        "Tiếp tục hoàn thiện những trường chưa có nguồn, rồi xuất lại báo cáo đã kiểm chứng.",
    ):
        doc.add_paragraph(message, style="List Number")
    doc.add_paragraph("TRẠNG THÁI: BÁO CÁO SƠ BỘ — CHƯA XÁC MINH, KHÔNG LÀ KẾT LUẬN THẦU.")
    for p in doc.paragraphs:
        p.style.font.size = Pt(10)
    output=BytesIO()
    doc.save(output)
    return output.getvalue()
