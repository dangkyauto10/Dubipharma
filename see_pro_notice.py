"""Read E-TBMT notices as labeled fields, not fake drug-lot tables.

Supported: XLSX and DOCX. Safe read-only classification and provenance.
Notice data is NOT a drug line list; a declared lot count is NOT parsed lots.
"""
from __future__ import annotations

from io import BytesIO
import re
import unicodedata
from typing import Any, Iterable

MAX_NOTICE_SIZE = 30 * 1024 * 1024
MAX_PREVIEW_ROWS = 160
MAX_PREVIEW_COLUMNS = 12

LABEL_ALIASES: dict[str, tuple[str, ...]] = {
    "MA_TBMT": ("ma tbmt", "so tbmt", "ma thong bao moi thau"),
    "MA_KHLCNT": ("ma khlcnt", "ma ke hoach lua chon nha thau"),
    "MA_GOI_THAU": ("ma goi thau",),
    "TEN_GOI_THAU": ("ten goi thau",),
    "CHU_DAU_TU": ("chu dau tu",),
    "SO_PHAN_LO": ("so luong phan lo", "so phan lo", "so luong phan", "so luong mat hang"),
    "THOI_DIEM_DONG_THAU": ("thoi diem dong thau",),
    "THOI_DIEM_MO_THAU": ("thoi diem mo thau",),
    "NGAY_DANG_TAI": ("ngay dang tai",),
    "PHIEN_BAN": ("phien ban thay doi", "phien ban"),
    "SO_TIEN_BAO_DAM_DU_THAU": ("so tien dam bao du thau", "so tien bao dam du thau"),
    "HINH_THUC_DU_THAU": ("hinh thuc du thau",),
}
FIELD_BY_LABEL = {alias: field for field, aliases in LABEL_ALIASES.items() for alias in aliases}
STRONG_FIELDS = {"MA_TBMT", "MA_GOI_THAU", "MA_KHLCNT"}


def _normalize(text: Any) -> str:
    t = unicodedata.normalize("NFKD", str(text or "").replace("Đ", "D").replace("đ", "d"))
    t = "".join(c for c in t if not unicodedata.combining(c)).lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", t)).strip()


def _notice_title(cell: Any) -> bool:
    text = _normalize(cell)
    return text.startswith("thong bao moi thau") and len(text) < 110


def _rows(filename: str, content: bytes) -> Iterable[tuple[list[tuple[str, str]], str]]:
    """Yield (text, exact source) cells, preserving original row boundaries."""
    extension = filename.lower().rsplit(".", 1)[-1]
    if extension == "xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets[:8]:
                for row_no, values in enumerate(sheet.iter_rows(max_row=MAX_PREVIEW_ROWS,
                                                                 max_col=MAX_PREVIEW_COLUMNS,
                                                                 values_only=True), 1):
                    cells = [(str(value).strip(), f"{filename} · {sheet.title}!{_col(i)}{row_no}")
                             for i, value in enumerate(values, 1)
                             if value is not None and str(value).strip()]
                    if cells:
                        yield cells, f"{sheet.title}!{row_no}"
        finally:
            workbook.close()
    elif extension == "docx":
        from docx import Document
        doc = Document(BytesIO(content))
        for line_no, paragraph in enumerate(doc.paragraphs[:MAX_PREVIEW_ROWS], 1):
            text = paragraph.text.strip()
            if text:
                yield [(text, f"{filename} · đoạn {line_no}")], f"đoạn {line_no}"
        for t, table in enumerate(doc.tables[:15], 1):
            for row_no, row in enumerate(table.rows[:MAX_PREVIEW_ROWS], 1):
                cells = [(cell.text.strip(),
                          f"{filename} · bảng Word {t}, dòng {row_no}, ô {i}")
                         for i, cell in enumerate(row.cells, 1)
                         if cell.text.strip()]
                if cells:
                    yield cells, f"bảng Word {t}, dòng {row_no}"


def _col(n: int) -> str:
    result = ""
    while n:
        n, r = divmod(n - 1, 26)
        result = chr(65 + r) + result
    return result


def parse_notice(filename: str, content: bytes) -> dict | None:
    """Return notice metadata only if source has a strong notice signature.

    No regulatory or compliance conclusion is made. Non-notices return None,
    allowing existing lot parsers to continue unchanged.
    """
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext not in ("xlsx", "docx"):
        return None
    if not content or len(content) > MAX_NOTICE_SIZE:
        raise ValueError("Thông báo rỗng hoặc vượt giới hạn 30 MB")
    found: dict[str, dict] = {}
    saw_title = False
    inspected = 0
    for cells, where in _rows(filename, content):
        inspected += 1
        for cell, loc in cells:
            if inspected < 22 and _notice_title(cell):
                saw_title = True
        # Typical E-GP export: first column is label, second is value.
        first, first_loc = cells[0]
        key = FIELD_BY_LABEL.get(_normalize(first.strip(" \t\r\n:")))
        value, value_loc = "", ""
        if key and len(cells) > 1:
            value, value_loc = cells[1]
        elif not key:
            normalized = _normalize(first)
            # Handle a single Word cell containing "Mã TBMT: IB...".
            for alias, candidate in FIELD_BY_LABEL.items():
                if normalized.startswith(alias + " ") and len(normalized) > len(alias) + 1:
                    key = candidate
                    value = first[len(first.split(":")[0]) + 1:].strip() if ":" in first else ""
                    if not value:
                        pattern = {
                            "MA_TBMT": r"IB\d{8,14}(?:-\d{2})?",
                            "MA_GOI_THAU": r"BP\d{8,14}",
                            "MA_KHLCNT": r"PL\d{8,14}",
                        }.get(key)
                        if pattern:
                            match = re.search(pattern, first, re.I)
                            value = match.group(0) if match else ""
                    value_loc = first_loc
                    if value:
                        break
        if not key or not value or key in found:
            continue
        value = value.strip()
        if key == "MA_TBMT" and not re.fullmatch(r"IB\d{8,14}(?:-\d{2})?", value, re.I):
            continue
        if key == "MA_GOI_THAU" and not re.fullmatch(r"BP\d{8,14}", value, re.I):
            continue
        if key == "MA_KHLCNT" and not re.fullmatch(r"PL\d{8,14}", value, re.I):
            continue
        if key == "SO_PHAN_LO":
            digits = re.search(r"\d+", value)
            if not digits:
                continue
            value = digits.group(0)
        found[key] = {
            "value": value,
            "source": value_loc or first_loc or where,
            "verified": False,
        }
    # Prefer the notice heading, but allow an unmistakable metadata signature
    # if Word has omitted the visual heading in paragraph extraction.
    strong = sum(bool(found.get(key)) for key in STRONG_FIELDS)
    if not (saw_title and strong >= 1 or strong >= 2 and "SO_PHAN_LO" in found
            and "TEN_GOI_THAU" in found):
        return None
    bp = found.get("MA_GOI_THAU", {}).get("value", "")
    return {
        "type": "THONG_BAO_MOI_THAU",
        "package_id": bp,
        "evidence": found,
        "warnings": [],
        "verified": False,
        "declared_lots_not_extracted": True,
    }
