"""SEE PRO: read-only intake for full E-HSMT workbooks exported by e-GP.

Conservative recognition of the *bid-guarantee summary* is distinct from
technical rows. No eligibility decisions, file writes, network or manual tags.
The original workbook is processed in memory; no content is persisted here.
"""
from __future__ import annotations

from collections import defaultdict
from io import BytesIO
import hashlib
import re

_LOT = re.compile(r"(?<!\w)G1\s*[.．]\s*(\d{3})(?!\d)", re.I)
_SUMMARY = re.compile(
    r"^(\d{1,4})\s+(PP\d{10})\s+G1\s*[.．]\s*(\d{3})\s+"
    r"(.+?)\s+(\d{1,3}(?:\.\d{3})+|\d{1,15})\s+"
    r"(\d{1,3}(?:\.\d{3})+|\d{1,15})$", re.I)


def _summary_candidate(cells: list[tuple[int, str]]) -> dict | None:
    """Detect an exact line; do not merge unrelated rows or guess missing data."""
    # Actual bid-guarantee tables contain STT, PP code, drug code, ingredient,
    # estimated value and guarantee. Wider technical tables are NOT summaries.
    if not 1 <= len(cells) <= 6:
        return None
    text = re.sub(r"\s+", " ", " ".join(value for _, value in cells)).strip()
    matched = _SUMMARY.fullmatch(text)
    if not matched:
        return None
    stt, part, suffix, ingredient, value, guarantee = matched.groups()
    if int(stt) != int(suffix) or not re.search(r"[A-Za-zÀ-ỹ]", ingredient):
        return None
    return {
        "MA_THUOC": f"G1.{suffix}", "MA_PHAN_LO": part.upper(),
        "STT": int(stt), "HOAT_CHAT_CHUAN": ingredient.strip(),
        "GIA_TRI_UOC_TINH": int(value.replace(".", "")),
        "BAO_DAM_DU_THAU": int(guarantee.replace(".", "")),
    }


def inspect_ehsmt_xlsx(data: bytes, *, max_bytes: int = 40 * 1024 * 1024,
                       max_rows: int = 20000, max_columns: int = 120) -> dict:
    """Extract source-anchored candidates from an authentic workbook, fail closed.

    A row containing a G1 code is not itself proof of a valid lot. A summary
    only qualifies when all six fields appear coherently in that single row.
    A repeated, conflicting summary is excluded, rather than guessed.
    """
    if not isinstance(data, bytes) or not data or len(data) > max_bytes:
        raise ValueError("Chỉ chấp nhận XLSX không rỗng, tối đa giới hạn tệp")
    if not 1 <= max_rows <= 100000 or not 1 <= max_columns <= 200:
        raise ValueError("Giới hạn quét không hợp lệ")
    from openpyxl import load_workbook
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Tệp XLSX không đọc được hoặc không hợp lệ") from exc
    sha = hashlib.sha256(data).hexdigest()
    positions = defaultdict(list)
    summaries = {}
    conflicting = set()
    warnings = []
    scanned = 0
    truncated = False
    try:
        for sheet in wb.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                scanned += 1
                if scanned > max_rows:
                    truncated = True
                    break
                if len(row) > max_columns:
                    warnings.append(f"{sheet.title}, dòng {row_number}: vượt giới hạn cột; không phân tích dòng này")
                    continue
                cells = [
                    (index + 1, str(value).strip())
                    for index, value in enumerate(row)
                    if value is not None and str(value).strip()
                ]
                if not cells:
                    continue
                full_text = " ".join(value for _, value in cells)
                for match in _LOT.finditer(full_text):
                    code = f"G1.{match.group(1)}"
                    if len(positions[code]) < 8:
                        positions[code].append({"sheet": sheet.title, "row": row_number})
                candidate = _summary_candidate(cells)
                if candidate is None:
                    continue
                code = candidate["MA_THUOC"]
                candidate["EVIDENCE"] = {
                    "file_sha256": sha, "sheet": sheet.title, "row": row_number,
                    "cells": [{"column": i, "raw": value} for i, value in cells],
                    "verified": False,
                }
                if code in summaries:
                    fields = ("MA_PHAN_LO", "STT", "HOAT_CHAT_CHUAN",
                              "GIA_TRI_UOC_TINH", "BAO_DAM_DU_THAU")
                    if any(candidate[f] != summaries[code][f] for f in fields):
                        conflicting.add(code)
                else:
                    summaries[code] = candidate
            if truncated:
                break
    finally:
        wb.close()
    if truncated:
        warnings.append(f"Giới hạn {max_rows} dòng: CHỈ ĐỌC MỘT PHẦN tệp")
    if conflicting:
        warnings.append(f"{len(conflicting)} mã có bảng tổng hợp MÂU THUẪN; không tự chọn dòng")
    missing = set(positions) - set(summaries) | conflicting
    if missing:
        warnings.append(f"{len(missing)} mã chỉ có tham chiếu hoặc tổng hợp thiếu/mâu thuẫn")
    if not positions:
        warnings.append("Không thấy mã G1.xxx; chưa thể kết luận về danh mục")
    lots = []
    for code in sorted(positions, key=lambda x: int(x.split(".")[1])):
        if code in summaries and code not in conflicting:
            record = dict(summaries[code])
            record["TRANG_THAI"] = "TRICH_XUAT_SO_BO_CHUA_XAC_MINH"
        else:
            record = {"MA_THUOC": code, "TRANG_THAI": "CAN_KIEM_TRA_DONG_GOC",
                      "MA_PHAN_LO": ""}
        record["VI_TRI_THAM_CHIEU"] = positions[code]
        record["verified"] = False
        lots.append(record)
    return {
        "sha256": sha, "sheets": len(wb.worksheets), "rows_scanned": min(scanned, max_rows),
        "complete_read": not truncated and not any("vượt giới hạn cột" in w for w in warnings),
        "lot_count": len(positions),
        "summary_rows": sum(k not in conflicting for k in summaries),
        "missing_summary": sorted(missing), "lots": lots,
        "warnings": warnings, "verified": False,
        "note": "Chỉ là trích xuất bảng tổng hợp từ E-HSMT; không phê duyệt NSX/thuốc/GMP/SĐK.",
    }


def inspect_ehsmt_document(data: bytes) -> dict:
    """Automatically route authentic PDF/XLSX E-HSMT by file signature."""
    if not isinstance(data, bytes) or not data:
        raise ValueError("Hồ sơ E-HSMT rỗng hoặc không hợp lệ")
    if data.startswith(b"%PDF-"):
        from tdv_ehsmt_intake import inspect_ehsmt_pdf
        return inspect_ehsmt_pdf(data)
    if data.startswith(b"PK\x03\x04"):
        from zipfile import BadZipFile, ZipFile
        try:
            with ZipFile(BytesIO(data)) as archive:
                if "xl/workbook.xml" in archive.namelist():
                    return inspect_ehsmt_xlsx(data)
        except BadZipFile:
            pass
    raise ValueError("Chưa nhận dạng được E-HSMT PDF/XLSX; không tự suy đoán loại hồ sơ")
