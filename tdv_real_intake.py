"""Read-only, session-only intake for real tender and manufacturer files.
No data persistence, no legal approval, no automatic mapping.
"""
from __future__ import annotations
from io import BytesIO
import hashlib
import re

MAX_UPLOAD = 30 * 1024 * 1024
MAX_PDF_PAGES = 800
MAX_ROWS = 10000
EXTENSIONS = {"pdf", "docx", "xlsx", "xls", "csv", "png", "jpg", "jpeg", "tif", "tiff"}

def extract_document(filename: str, data: bytes) -> dict:
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext not in EXTENSIONS:
        raise ValueError("Định dạng chưa được hỗ trợ")
    if not data or len(data) > MAX_UPLOAD:
        raise ValueError("Tệp rỗng hoặc lớn hơn 30 MB; cần chia nhỏ an toàn")
    result = {"name": filename, "sha256": hashlib.sha256(data).hexdigest(),
              "type": ext, "lines": [], "warnings": [], "pages": 0,
              "complete": True, "verified": False}
    def add(source, value):
        if value is not None and str(value).strip():
            result["lines"].append({"source": source, "text": str(value).strip()[:2500]})
    if ext == "pdf":
        import pymupdf
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            if doc.needs_pass:
                raise ValueError("PDF cần mật khẩu")
            result["pages"] = doc.page_count
            if doc.page_count > MAX_PDF_PAGES:
                result["complete"] = False
                result["warnings"].append("Quá 800 trang; chỉ xem trước 800 trang")
            for i in range(min(doc.page_count, MAX_PDF_PAGES)):
                page = doc[i]
                raw = page.get_text("text", sort=True) or ""
                if not raw.strip():
                    result["complete"] = False
                    result["warnings"].append(f"Trang {i+1}: không có lớp chữ; cần OCR và đối chiếu ảnh gốc")
                for line in raw.splitlines():
                    add(f"Trang {i+1}", line)
    elif ext == "docx":
        from docx import Document
        doc = Document(BytesIO(data))
        for i, p in enumerate(doc.paragraphs, 1):
            add(f"Đoạn {i}", p.text)
        for t, table in enumerate(doc.tables, 1):
            for rowno, row in enumerate(table.rows, 1):
                add(f"Bảng {t}, dòng {rowno}", " | ".join(c.text for c in row.cells))
    elif ext in ("xlsx", "xls"):
        if ext == "xls":
            raise ValueError("Excel .xls cũ: vui lòng lưu thành .xlsx; không tự chuyển đổi")
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        count = 0
        for sheet in wb.worksheets:
            for ri, row in enumerate(sheet.iter_rows(), 1):
                if count >= MAX_ROWS:
                    result["complete"] = False
                    result["warnings"].append("Vượt 10.000 dòng; kết quả chỉ là bản xem trước")
                    break
                # read_only iter_rows yields ReadOnlyCell objects; read their values,
                # not their object representations. Preserve numeric zero.
                values = [f"{col}:{str(cell.value).strip()[:350]}" for col, cell in enumerate(row, 1)
                          if cell.value is not None and str(cell.value).strip()]
                if values:
                    add(f"Sheet {sheet.title}, dòng {ri}", " | ".join(values))
                count += 1
            if count >= MAX_ROWS:
                break
        wb.close()
    elif ext == "csv":
        import csv
        import io
        decoded = data.decode("utf-8-sig", errors="replace")
        for ri, row in enumerate(csv.reader(io.StringIO(decoded)), 1):
            if ri > MAX_ROWS:
                result["complete"] = False
                result["warnings"].append("CSV quá 10.000 dòng")
                break
            add(f"Dòng {ri}", " | ".join(row))
    else:
        result["complete"] = False
        result["warnings"].append("Ảnh cần OCR cục bộ và kiểm tra trực quan; chưa tự trích xuất")
    if not result["lines"]:
        result["complete"] = False
    return result

def candidate_lots(document: dict) -> list[dict]:
    """Candidate IDs only, never infer fields or claim verified rows."""
    pattern = re.compile(r"\bG1[.]\d{3}\b", re.IGNORECASE)
    found = {}
    for line in document["lines"]:
        for match in pattern.findall(line["text"]):
            code = match.upper()
            found.setdefault(code, {"Mã lô": code, "Nguồn": line["source"],
                                     "Trích dẫn": line["text"][:240],
                                     "Xác minh": "CHƯA"})
    return list(found.values())
