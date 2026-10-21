"""Synthetic multi-worksheet evidence preservation and conservative merging.

Do NOT commit actual tender or supplier files to public test repositories.
"""
from io import BytesIO
import unittest

from openpyxl import Workbook, load_workbook

from tdv_nsx_multisheet import reconcile_product_rows
from tdv_nsx_auto import analyze_auto, export_auto_xlsx
from tdv_leadership_report import compare_sources


def item(name, factory, strength, ingredient="", sheet="Sheet One", no=5,
         registration="", product_code=""):
    f = {"PRODUCT_NAME": name, "FACTORY": factory, "STRENGTH": strength}
    if ingredient:
        f["INGREDIENT"] = ingredient
    if registration:
        f["REGISTRATION_NO"] = registration
    if product_code:
        f["PRODUCT_ID"] = product_code
    return {"uid": "P0001", "fields": f,
        "evidence": {k: {"file": "synthetic.xlsx", "sha256": "s"*64,
                         "source": f"{sheet}, hàng {no}, cột {i+1}",
                         "raw": v, "value": v, "verified": False}
                     for i, (k, v) in enumerate(f.items())},
        "conflicts": {}, "file": "synthetic.xlsx", "sha256": "s"*64,
        "origin": f"{sheet}, dòng {no}",
        "complete_read": True, "verified": False}


def synthetic_two_sheet_catalogue():
    wb = Workbook()
    sh = wb.active
    sh.title = "Bang gia"
    sh.append(["Tên thuốc", "Hàm lượng", "Nhà máy sản xuất", "Dạng bào chế"])
    sh.append(["4", "6", "13", "8"])
    sh.append(["Medicine A", "50 mg", "Factory A", "Viên nén"])
    sh = wb.create_sheet("Bang ky thuat")
    sh.append(["Tên thuốc", "Hoạt chất", "Hàm lượng", "Nhà máy sản xuất"])
    sh.append(["Medicine A", "Active A", "100 mg", "Factory A"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ConservativeMultiSheetQA(unittest.TestCase):
    def test_exclude_numeric_header_and_merge_only_source_corroborated(self):
        products, info = reconcile_product_rows([
            item("4","13","6","",sheet="Sheet A",no=2),
            item("Medicine A","Factory A","50 mg",sheet="Sheet A",no=3),
            item("Medicine A","Factory A","100 mg",ingredient="Active A",
                 sheet="Sheet B",no=2),
        ])
        self.assertEqual(info["source_rows"],3)
        self.assertEqual(len(info["excluded_ghost_rows"]),1)
        self.assertEqual(info["merged_secondary_rows"],1)
        self.assertEqual(len(products),1)
        self.assertEqual(products[0]["fields"]["INGREDIENT"],"Active A")
        self.assertEqual(products[0]["evidence"]["INGREDIENT"]["source"],"Sheet B, hàng 2, cột 4")
        self.assertIn("STRENGTH",products[0]["conflicts"])
        self.assertEqual(len(products[0]["source_rows"]),2)
        self.assertEqual({r["fields"]["STRENGTH"] for r in products[0]["source_rows"]},
                         {"50 mg","100 mg"})
        self.assertFalse(products[0]["verified"])

    def test_same_name_without_shared_identity_stays_unmerged(self):
        a=item("Medicine A","Factory A","50 mg",sheet="Sheet A")
        b=item("Medicine A","Factory B","50 mg",sheet="Sheet B")
        products, info = reconcile_product_rows([a,b])
        self.assertEqual(len(products),2)
        self.assertEqual(info["ambiguous_name_groups"],1)

    def test_duplicated_name_in_one_sheet_is_not_arbitrarily_paired(self):
        products, info = reconcile_product_rows([
            item("Medicine A","Factory A","50 mg",sheet="Sheet A",no=2),
            item("Medicine A","Factory A","100 mg",sheet="Sheet A",no=3),
            item("Medicine A","Factory A","50 mg",sheet="Sheet B",no=4),
        ])
        self.assertEqual(len(products),3)
        self.assertEqual(info["merged_secondary_rows"],0)

    def test_cross_sheet_line_break_and_proven_identity(self):
        a=item("Long\nmedicine","Factory A","50 mg",sheet="Sheet A")
        b=item("Longmedicine","Factory A","50 mg",sheet="Sheet B")
        out,_=reconcile_product_rows([a,b])
        self.assertEqual(len(out),1)
        self.assertEqual(len(out[0]["source_rows"]),2)

    def test_registration_conflict_blocks_all_merging(self):
        a=item("Medicine A","Factory A","50 mg",sheet="Sheet A",registration="REG-001")
        b=item("Medicine A","Factory A","50 mg",sheet="Sheet B",registration="REG-002")
        out,info=reconcile_product_rows([a,b])
        self.assertEqual(len(out),2)
        self.assertEqual(info["ambiguous_name_groups"],1)

    def test_header_with_real_registration_identifier_is_not_silently_removed(self):
        a=item("4","Factory A","50 mg",sheet="Sheet A",registration="REG-123")
        out,info=reconcile_product_rows([a])
        self.assertEqual(len(out),1)
        self.assertEqual(len(info["excluded_ghost_rows"]),0)

    def test_end_to_end_ghost_filtered_technical_conflict_yellow_and_export(self):
        raw=synthetic_two_sheet_catalogue()
        report=analyze_auto([("synthetic.xlsx",raw)],enable_ocr=False)
        self.assertEqual(report["summary"]["files"],1)
        self.assertEqual(report["summary"]["source_product_rows"],3)
        self.assertEqual(report["summary"]["products"],1)
        self.assertEqual(report["summary"]["reconciled_rows"],1)
        self.assertEqual(report["summary"]["excluded_ghost_rows"],1)
        p=report["products"][0]
        self.assertEqual(p["fields"]["INGREDIENT"],"Active A")
        self.assertEqual(p["fields"]["PRODUCT_NAME"],"Medicine A")
        self.assertIn("STRENGTH",p["conflicts"])
        self.assertFalse(report["bid_eligible"])
        ehsmt={"rows":[{"MA_THUOC":"G1.001","HOAT_CHAT_HSMT":"Active A",
                        "NONG_DO_HAM_LUONG":"50 mg",
                        "DANG_BAO_CHE":"Viên nén",
                        "DUONG_DUNG":"Uống",
                        "FILE_NGUON":"synthetic_ehsmt.pdf",
                        "EVIDENCE_KY_THUAT":{"HOAT_CHAT_HSMT":{
                            "file":"synthetic_ehsmt.pdf","raw":"Active A","source":"p 1"}
                        }}]}
        result=compare_sources(ehsmt,report)
        row=result["comparison"][0]
        self.assertEqual(row["SO_UNG_VIEN_NSX"],1)
        self.assertEqual(row["criteria"][1]["status"],"CHUA_DU_CAN_CU")
        self.assertNotEqual(row["criteria"][1]["status"],"KHOP_KY_THUAT_SO_BO")
        self.assertEqual(row["SO_TIEU_CHI_KHONG_KHOP"],0)
        self.assertEqual(result["summary"]["fully_verified"],0)
        result_xlsx=load_workbook(BytesIO(export_auto_xlsx(report)),read_only=True)
        try:
            self.assertIn("GOP_DONG_NHIEU_SHEET",result_xlsx.sheetnames)
            names=list(result_xlsx["TRUY_XUAT_TUNG_TRUONG"].values)
            strength_sources=[r[3] for r in names if r[1]=="PRODUCT" and r[2]=="STRENGTH"]
            self.assertIn("50 mg",strength_sources)
            self.assertIn("100 mg",strength_sources)
        finally:
            result_xlsx.close()

if __name__=="__main__":
    unittest.main()
