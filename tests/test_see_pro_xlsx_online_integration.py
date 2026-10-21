"""Synthetic regression of real-source READ-ONLY XLSX wiring to existing online SEE PRO.

This is a code QA fixture, never evidence of correct extraction of a real tender.
"""
from __future__ import annotations

from io import BytesIO
import unittest
from openpyxl import Workbook, load_workbook
from see_pro_workbench import analyze_uploads
from see_pro_traceability import build_audit_workbook


def ehsmt_xlsx(*, incomplete=False) -> bytes:
    book=Workbook()
    ws=book.active
    ws.title="Tong hop"
    ws.append(("STT", "Ma phan", "Ma thuoc", "Hoat chat", "Gia tri", "Bao dam"))
    ws.append((1, "PP1234567890", "G1.001", "Ingredient X", "10.000", "1.000"))
    ws.append((2, "PP1234567891", "G1.002", "Ingredient Y", "20.000", None if incomplete else "2.000"))
    out=BytesIO()
    book.save(out)
    return out.getvalue()


class SeeProOnlineXlsxIntegration(unittest.TestCase):
    def test_summary_rows_are_visible_in_existing_online_workbench(self):
        result=analyze_uploads([("ehsmt.xlsx",ehsmt_xlsx())])
        self.assertEqual(len(result["documents"]),1)
        self.assertEqual(len(result["structured_tables"]),1)
        self.assertEqual(len(result["rows"]),2)
        self.assertEqual({r["MA_THUOC"] for r in result["rows"]},{"G1.001","G1.002"})
        self.assertEqual(result["structured_tables"][0]["Dòng bảng tổng hợp rõ"],2)
        for row in result["rows"]:
            self.assertEqual(row["TRANG_THAI"],"TRICH_XUAT_SO_BO_CHUA_XAC_MINH")
            self.assertFalse(row["verified"])
            self.assertTrue(row["EVIDENCE"]["DONG_BANG_TONG_HOP"]["row"])
        evidence=load_workbook(BytesIO(build_audit_workbook(result)),read_only=True)
        self.assertEqual(evidence.sheetnames,[
            "DANH_MUC_PHAN_LO","NGUON_TAI_LIEU","TRUY_XUAT_CHUNG_CU",
            "CANH_BAO_CAN_XEM","PHAM_VI_KIEM_TOAN"])
        self.assertEqual(evidence["DANH_MUC_PHAN_LO"].max_row,3)
        self.assertEqual(evidence["TRUY_XUAT_CHUNG_CU"].max_row,3)
        evidence.close()

    def test_incomplete_summary_stays_unverified_in_download(self):
        result=analyze_uploads([("ehsmt.xlsx",ehsmt_xlsx(incomplete=True))])
        self.assertEqual(len(result["rows"]),2)
        self.assertEqual(result["structured_tables"][0]["Dòng bảng tổng hợp rõ"],1)
        self.assertFalse(result["structured_tables"][0]["Nhận diện cấu trúc"])
        self.assertTrue(result["warnings"])
        missing=next(r for r in result["rows"] if r["MA_THUOC"]=="G1.002")
        self.assertEqual(missing["TRANG_THAI"],"CAN_KIEM_TRA_DONG_GOC")
        self.assertFalse(missing["verified"])

if __name__=="__main__":
    unittest.main()
