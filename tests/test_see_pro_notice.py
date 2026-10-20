"""SEE PRO notice parser regression tests; no actual customer files committed."""
from __future__ import annotations

from io import BytesIO
import unittest

from docx import Document
from openpyxl import Workbook

from see_pro_notice import parse_notice
from see_pro_structured import identify_document, parse_excel_lots
from see_pro_workbench import analyze_uploads


def notice_xlsx() -> bytes:
    wb = Workbook()
    sh = wb.active
    sh.title = "Sheet1"
    for label, value in (
        ("THÔNG BÁO MỜI THẦU", None),
        (None, None),
        ("Thông tin cơ bản", None),
        ("Mã TBMT", "IB2600581793"),
        ("Ngày đăng tải", "20/10/2026 13:57"),
        ("Mã KHLCNT", "PL2600355990"),
        ("Thông tin gói thầu", None),
        ("Tên gói thầu", "Gói thầu số 1: Gói thầu thuốc generic"),
        ("Mã gói thầu", "BP2600899572"),
        ("Chủ đầu tư", "Bệnh viện Công an tỉnh Tuyên Quang"),
        ("Số lượng phần (lô)", 341),
        ("Thời điểm đóng thầu", "29/10/2026 10:00"),
        ("Số tiền đảm bảo dự thầu", "99.585.600 VND"),
    ):
        sh.append([label, value])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def notice_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("THÔNG BÁO MỜI THẦU")
    table = doc.add_table(rows=0, cols=2)
    for label, value in (
        ("Mã TBMT", "IB2600581793"),
        ("Mã KHLCNT", "PL2600355990"),
        ("Tên gói thầu", "Gói thầu thuốc generic"),
        ("Mã gói thầu", "BP2600899572"),
        ("Số lượng phần (lô)", "341"),
        ("Thời điểm đóng thầu", "29/10/2026 10:00"),
    ):
        r = table.add_row()
        r.cells[0].text = label
        r.cells[1].text = value
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def matching_lot_excel() -> bytes:
    wb = Workbook()
    sh = wb.active
    sh.append(["Mã phần/lô", "Mã thuốc", "Tên hoạt chất",
               "Số lượng", "Đơn giá (VND)", "Tổng giá trị (VND)"])
    sh.append(["PP2600443936", "GE841.001", "Acenocoumarol", 100, 250, 25000])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class NoticeClassificationTests(unittest.TestCase):
    def test_xlsx_metadata_and_original_cell_provenance(self):
        raw = notice_xlsx()
        notice = parse_notice("Generic Thông báo mời thầu.xlsx", raw)
        self.assertIsNotNone(notice)
        self.assertEqual(notice["type"], "THONG_BAO_MOI_THAU")
        self.assertEqual(notice["package_id"], "BP2600899572")
        self.assertEqual(notice["evidence"]["MA_TBMT"]["value"], "IB2600581793")
        self.assertEqual(notice["evidence"]["MA_KHLCNT"]["value"], "PL2600355990")
        self.assertEqual(notice["evidence"]["SO_PHAN_LO"]["value"], "341")
        self.assertIn("Sheet1!B4", notice["evidence"]["MA_TBMT"]["source"])
        self.assertIn("Sheet1!B11", notice["evidence"]["SO_PHAN_LO"]["source"])
        self.assertFalse(notice["verified"])

    def test_docx_metadata_and_cell_provenance(self):
        notice = identify_document("Generic Thông báo mời thầu.docx", notice_docx())
        self.assertEqual(notice["type"], "THONG_BAO_MOI_THAU")
        self.assertEqual(notice["evidence"]["SO_PHAN_LO"]["value"], "341")
        self.assertIn("bảng Word 1", notice["evidence"]["MA_TBMT"]["source"])
        self.assertEqual(notice["package_id"], "BP2600899572")

    def test_notice_no_false_lot_or_merge_with_unrelated_bp(self):
        raw = notice_docx()
        report = analyze_uploads([("Generic Thông báo mời thầu.docx", raw)])
        self.assertEqual(len(report["notices"]), 1)
        self.assertEqual(report["rows"], [])
        self.assertEqual(report["structured_tables"], [])
        self.assertEqual(report["warnings"], [])
        self.assertIn("Chưa có bảng danh mục", report["notices"][0]["lot_list_status"])

        report2 = analyze_uploads([
            ("Generic Thông báo mời thầu.docx", raw),
            ("parts_BP2600899590.xlsx", matching_lot_excel()),
        ])
        self.assertEqual(len(report2["rows"]), 1)
        self.assertEqual(report2["rows"][0]["MA_GOI_THAU"], "BP2600899590")
        self.assertIn("Chưa có bảng danh mục", report2["notices"][0]["lot_list_status"])
        self.assertFalse(any("khớp gói" in warning for warning in report2["warnings"]))

    def test_xlsx_notice_by_content_is_not_lot_table(self):
        report = analyze_uploads([("any_filename.xlsx", notice_xlsx())])
        self.assertEqual(report["documents"][0]["Loại nghiệp vụ"], "THONG_BAO_MOI_THAU")
        self.assertEqual(len(report["notices"]), 1)
        self.assertEqual(report["structured_tables"], [])
        self.assertEqual(report["warnings"], [])
        self.assertEqual(report["notices"][0]["evidence"]["SO_PHAN_LO"]["value"], "341")

    def test_same_package_link_and_count_mismatch_still_warns(self):
        report = analyze_uploads([
            ("Generic Thông báo mời thầu.xlsx", notice_xlsx()),
            ("parts_BP2600899572.xlsx", matching_lot_excel()),
        ])
        self.assertEqual(len(report["notices"]), 1)
        self.assertEqual(len(report["rows"]), 1)
        self.assertIn("CẦN KIỂM TRA", report["notices"][0]["lot_list_status"])
        self.assertTrue(any("341, bảng khớp 1" in x for x in report["warnings"]))

    def test_non_notice_lot_sheet_preserves_old_parser(self):
        lot = matching_lot_excel()
        self.assertIsNone(parse_notice("parts_BP2600899572.xlsx", lot))
        doc = parse_excel_lots("parts_BP2600899572.xlsx", lot)
        self.assertEqual(doc["total_lots"], 1)
        self.assertTrue(doc["complete"])

    def test_narrative_docx_with_no_strong_notice_marker_not_classified(self):
        doc = Document()
        doc.add_paragraph("Các lưu ý đối với THÔNG BÁO MỜI THẦU không liên quan tới lô.")
        doc.add_paragraph("Không có mã IB hoặc BP.")
        stream = BytesIO()
        doc.save(stream)
        self.assertIsNone(parse_notice("narrative.docx", stream.getvalue()))

    def test_single_cell_tbmt_docx_and_missing_heading(self):
        doc = Document()
        tbl = doc.add_table(rows=0, cols=1)
        for value in ("Mã TBMT: IB2600581793", "Mã KHLCNT: PL2600355990",
                      "Tên gói thầu: Generic", "Mã gói thầu: BP2600899572",
                      "Số lượng phần (lô): 341"):
            tbl.add_row().cells[0].text = value
        stream = BytesIO()
        doc.save(stream)
        notice = parse_notice("no_title.docx", stream.getvalue())
        self.assertIsNotNone(notice)
        self.assertEqual(notice["evidence"]["MA_GOI_THAU"]["value"], "BP2600899572")


if __name__ == "__main__":
    unittest.main()
