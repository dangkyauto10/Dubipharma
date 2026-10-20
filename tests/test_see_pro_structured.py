"""SEE PRO structured mapping tests; no customer documents committed."""
from __future__ import annotations
from io import BytesIO
import unittest

from openpyxl import Workbook
import pymupdf
from see_pro_structured import (
    parse_excel_lots, identify_document, reconcile_notice, normalize)
from see_pro_workbench import analyze_uploads


HEADERS = [
    "STT", "Mã phần/lô", "Mã thuốc", "Tên hoạt chất", "Nhóm TCKT",
    "Nồng độ hoặc hàm lượng", "Đường dùng", "Dạng bào chế", "Đơn vị tính",
    "Số lượng", "Đơn giá (VND)", "Tổng giá trị (VND)",
    "Số lượng tùy chọn mua thêm", "Tỷ lệ phần trăm (%)",
    "Giá trị tùy chọn mua thêm", "Giá dự toán ước tính từng lô",
]


def make_book(n: int = 2, code_prefix: str = "GE841", herbal: bool = False,
              corrupt: bool = False, optional_bad: bool = False):
    book = Workbook()
    sheet = book.active
    sheet.title = "Danh sách phần lô"
    sheet.append([x if x != "Tên hoạt chất" or not herbal
                  else "Tên thành phần của thuốc" for x in HEADERS])
    for idx in range(1, n + 1):
        code = f"{code_prefix}.{idx:02d}" if herbal else f"{code_prefix}.{idx:03d}"
        amount = 25000 + (100 if corrupt and idx == 2 else 0)
        optional_value = 7500 + (3 if optional_bad and idx == 2 else 0)
        sheet.append([idx, f"PP{2600443935 + idx:010d}", code,
                      "Acenocoumarol" if not herbal else "Actiso",
                      "N4" if not herbal else "3",
                      "4 mg", "Uống", "Viên", "Viên", 100, 250, amount,
                      30, 30, optional_value, "23" if not herbal else "20"])
    sheet.append([None]*11 + [25000 * n] + [None]*4)
    out = BytesIO()
    book.save(out)
    return out.getvalue()


def notice_dict():
    return {"package_id": "BP2600899572",
            "evidence": {"SO_PHAN_LO": {"value": "341", "verified": False}}}


class SeeStructuredTests(unittest.TestCase):
    def test_header_matching_generic(self):
        doc = parse_excel_lots("Danh_sach_phan_lo_BP2600899572.xlsx", make_book(341))
        self.assertEqual(doc["total_lots"], 341)
        self.assertTrue(doc["complete"])
        self.assertFalse(doc["verified"])
        self.assertEqual(doc["rows"][0]["MA_THUOC"], "GE841.001")
        self.assertEqual(doc["rows"][-1]["MA_THUOC"], "GE841.341")
        self.assertEqual(doc["rows"][0]["MA_GOI_THAU"], "BP2600899572")
        self.assertEqual(doc["rows"][0]["EVIDENCE"]["SO_LUONG_CHINH"]["row"], 2)
        self.assertEqual(doc["rows"][0]["EVIDENCE"]["SO_LUONG_CHINH"]["column"], 10)
        self.assertFalse(doc["issues"])
        self.assertEqual(doc["rows"][0]["GIA_DU_TOAN_COT_GOC"], "23")

    def test_alternate_herbal_headers(self):
        doc = parse_excel_lots("Danh_sach_phan_lo_BP2600899590.xlsx",
                               make_book(39, "CP841", herbal=True))
        self.assertEqual(doc["total_lots"], 39)
        self.assertEqual(doc["rows"][0]["MA_THUOC"], "CP841.01")
        self.assertEqual(doc["rows"][0]["HOAT_CHAT_HSMT"], "Actiso")
        self.assertEqual(doc["rows"][-1]["MA_THUOC"], "CP841.39")

    def test_exact_notice_package_join(self):
        a = parse_excel_lots("Danh_sach_phan_lo_BP2600899572.xlsx", make_book(341))
        b = parse_excel_lots("Danh_sach_phan_lo_BP2600899590.xlsx",
                             make_book(39, "CP841", herbal=True))
        self.assertEqual(reconcile_notice(notice_dict(), [a, b]), [])
        self.assertEqual(reconcile_notice(notice_dict(), [b])[0]["type"],
                         "NO_MATCHING_LOT_LIST")
        self.assertEqual(reconcile_notice(notice_dict(), [parse_excel_lots(
            "Danh_sach_phan_lo_BP2600899572.xlsx", make_book(340))])[0]["type"],
                         "NOTICE_LOT_COUNT_MISMATCH")

    def test_unaltered_source_prices_on_mismatch(self):
        doc = parse_excel_lots("parts_BP2600899572.xlsx",
                               make_book(corrupt=True, optional_bad=True))
        kinds = {issue["type"] for issue in doc["issues"]}
        self.assertIn("MATH_TOTAL_MISMATCH", kinds)
        self.assertIn("OPTIONAL_TOTAL_MISMATCH", kinds)
        self.assertIn("FOOTER_SUM_MISMATCH", kinds)
        self.assertEqual(doc["rows"][1]["THANH_TIEN_KE_HOACH"], 25100)
        self.assertEqual(doc["rows"][1]["DON_GIA_KE_HOACH"], 250)

    def test_duplicate_part_or_drug_is_flagged(self):
        book = Workbook()
        sheet = book.active
        sheet.append(HEADERS)
        row = [1, "PP2600443936", "GE841.001", "Acenocoumarol",
               "N4", "4mg", "Uống", "Viên", "Viên", 100,
               250, 25000, 30, 30, 7500, "23"]
        sheet.append(row)
        sheet.append([2] + row[1:])
        out = BytesIO()
        book.save(out)
        doc = parse_excel_lots("data_BP2600899572.xlsx", out.getvalue())
        types = {i["type"] for i in doc["issues"]}
        self.assertIn("DUPLICATE_PART", types)
        self.assertIn("DUPLICATE_DRUG", types)

    def test_notice_pdf_is_not_forced_into_g1(self):
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "THONG BAO MOI THAU")
        raw = pdf.tobytes()
        pdf.close()
        info = identify_document("invitation.pdf", raw)
        self.assertEqual(info["type"], "THONG_BAO_MOI_THAU")
        result = analyze_uploads([("invitation.pdf", raw)])
        self.assertEqual(result["rows"], [])
        self.assertEqual(len(result["notices"]), 1)
        self.assertFalse(any("Không nhận diện được mã G1" in w
                             for w in result["warnings"]))

    def test_workbench_count_and_distinct_packages(self):
        doc_a = make_book(2)
        doc_b = make_book(1, "CP841", herbal=True)
        result = analyze_uploads([
            ("ds_BP2600899572.xlsx", doc_a),
            ("ds_BP2600899590.xlsx", doc_b),
        ])
        self.assertEqual(len(result["rows"]), 3)
        self.assertEqual(len(result["structured_tables"]), 2)
        self.assertEqual({x["MA_GOI_THAU"] for x in result["rows"]},
                         {"BP2600899572", "BP2600899590"})
        self.assertEqual(result["rows"][0]["verified"], False)

    def test_unknown_header_is_not_silently_claimed(self):
        book = Workbook()
        book.active.append(["Unknown", "Quantity", "Price"])
        book.active.append(["GE841.001", 1, 2])
        out = BytesIO()
        book.save(out)
        result = parse_excel_lots("ds_BP2600899572.xlsx", out.getvalue())
        self.assertFalse(result["complete"])
        self.assertEqual(result["total_lots"], 0)

    def test_filename_underscore_does_not_hide_bp(self):
        self.assertEqual(identify_document(
            "Danh_sach_phan_lo_BP2600899572.xlsx", make_book())["package_id"],
            "BP2600899572")

    def test_vietnamese_header_normalization(self):
        self.assertEqual(normalize("Nồng độ hoặc hàm lượng"),
                         "nong do hoac ham luong")
        self.assertEqual(normalize("Đường dùng"), "duong dung")


if __name__ == "__main__":
    unittest.main()
