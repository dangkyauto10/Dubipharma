"""SEE PRO authentic PDF source viewing and safe professional audit XLSX."""
from __future__ import annotations

import hashlib
from io import BytesIO
import unittest

import pymupdf
from openpyxl import load_workbook

from see_pro_traceability import build_audit_workbook, safe_cell
from see_pro_pdf_viewer import read_original_pdf_page, _integer_page


class TraceableExportTests(unittest.TestCase):
    def _analysis(self):
        return {
            "documents": [{"Tệp": "test_BP2600899572.xlsx", "Loại": "XLSX",
                            "SHA256": "c" * 64, "Xác minh": "CHƯA"},
                           {"Tệp": "test.pdf", "Loại": "PDF",
                            "SHA256": "b" * 64, "Xác minh": "CHƯA"}],
            "rows": [
                {
                    "FILE_NGUON": "test_BP2600899572.xlsx",
                    "MA_GOI_THAU": "BP2600899572",
                    "MA_PHAN_LO": "PP2600443936",
                    "MA_THUOC": "GE841.001",
                    "HOAT_CHAT_HSMT": "Acenocoumarol",
                    "SO_LUONG_CHINH": 100,
                    "DON_GIA_KE_HOACH": 1250,
                    "THANH_TIEN_KE_HOACH": 125000,
                    "TRANG_THAI": "TRICH_XUAT_CHUA_XAC_MINH",
                    "TRICH_DAN": '=HYPERLINK("https://example.invalid/evil")',
                    "EVIDENCE": {
                        "SO_LUONG_CHINH": {"source": "test_BP2600899572.xlsx",
                                          "sheet": "Sheet1", "row": 2,
                                          "column": 10, "raw": "100"},
                        "DON_GIA_KE_HOACH": {"source": "test_BP2600899572.xlsx",
                                            "sheet": "Sheet1", "row": 2,
                                            "column": 11, "raw": "1250"},
                    }, "verified": False,
                }, {
                    "FILE_NGUON": "test.pdf", "MA_THUOC": "G1.362",
                    "MA_PHAN_LO": "PP2600428741",
                    "TRANG_VI_TRI": 273,
                    "TRANG_BANG_KY_THUAT": 118,
                    "TRANG_THAI": "CAN_XAC_MINH",
                    "EVIDENCE_KY_THUAT": {
                        "SO_LUONG_CHINH": {"pages": [118, 119],
                                          "x_range": [499, 550],
                                          "raw_tokens": ["17.25", "0"]}
                    },
                },
            ],
            "warnings": ['test.pdf: \t=HYPERLINK("https://bad.invalid")'],
        }

    def test_export_contains_original_and_field_level_source(self):
        book = load_workbook(BytesIO(build_audit_workbook(self._analysis())))
        self.assertEqual(book.sheetnames, [
            "DANH_MUC_PHAN_LO", "NGUON_TAI_LIEU",
            "TRUY_XUAT_CHUNG_CU", "CANH_BAO_CAN_XEM",
            "PHAM_VI_KIEM_TOAN",
        ])
        ws = book["DANH_MUC_PHAN_LO"]
        headers = [cell.value for cell in ws[1]]
        self.assertEqual(ws.cell(2, headers.index("MA_GOI_THAU") + 1).value,
                         "BP2600899572")
        self.assertEqual(ws.cell(2, headers.index("DON_GIA_KE_HOACH") + 1).value,
                         1250)
        injection = ws.cell(2, headers.index("TRICH_DAN") + 1)
        self.assertTrue(injection.value.startswith("'="))
        self.assertNotEqual(injection.data_type, "f")
        self.assertTrue(ws.auto_filter.ref)
        self.assertEqual(ws.freeze_panes, "D2")
        self.assertEqual(ws.cell(1, headers.index("TRANG_THAI") + 1).fill.fgColor.rgb[-6:],
                         "163B67")
        self.assertEqual(len(list(book["TRUY_XUAT_CHUNG_CU"].values)), 4)
        e = list(book["TRUY_XUAT_CHUNG_CU"].values)
        self.assertEqual(e[1][4], "test_BP2600899572.xlsx")
        self.assertIn("118, 119", str(e[3][8]))
        book.close()

    def test_formula_like_warning_is_not_executed(self):
        book = load_workbook(BytesIO(build_audit_workbook(self._analysis())))
        cell = book["CANH_BAO_CAN_XEM"]["B2"]
        self.assertNotEqual(cell.data_type, "f")
        self.assertTrue(book["NGUON_TAI_LIEU"]["C2"].value.startswith("c"))
        book.close()

    def test_empty_report_creates_explained_workbook(self):
        book = load_workbook(BytesIO(build_audit_workbook({
            "rows": [], "documents": [], "warnings": [],
        })))
        self.assertEqual(book["PHAM_VI_KIEM_TOAN"]["A2"].value, "PHẠM VI")
        self.assertEqual(book["DANH_MUC_PHAN_LO"].max_row, 1)
        book.close()

    def test_safe_cell_formula_string_and_numbers(self):
        self.assertEqual(safe_cell("  =2+2"), "'  =2+2")
        self.assertEqual(safe_cell("+SUM(A1:B3)"), "'+SUM(A1:B3)")
        self.assertEqual(safe_cell(125000), 125000)
        self.assertEqual(safe_cell(None), None)


class OriginalPDFViewerTests(unittest.TestCase):
    @staticmethod
    def _pdf():
        document = pymupdf.open()
        document.new_page().insert_text((50, 60), "PAGE FIRST 111")
        document.new_page().insert_text((50, 60), "PAGE SECOND 222")
        data = document.tobytes()
        document.close()
        return data

    def test_displayed_png_is_original_page(self):
        pdf = self._pdf()
        digest = hashlib.sha256(pdf).hexdigest()
        result = read_original_pdf_page(pdf, 2, expected_sha256=digest, dpi=110)
        self.assertEqual(result["page"], 2)
        self.assertEqual(result["total_pages"], 2)
        self.assertEqual(result["sha256"], digest)
        self.assertTrue(result["png"].startswith(b"\x89PNG\r\n\x1a\n"))
        self.assertFalse(result["verified_technical"])

    def test_invalid_page_and_fake_digest_refused(self):
        pdf = self._pdf()
        with self.assertRaises(ValueError):
            read_original_pdf_page(pdf, 0)
        with self.assertRaises(ValueError):
            read_original_pdf_page(pdf, 3)
        with self.assertRaises(ValueError):
            read_original_pdf_page(pdf, 1, expected_sha256="fake")
        with self.assertRaises(ValueError):
            read_original_pdf_page(pdf, 1, dpi=300)

    def test_non_pdf_rejected_and_page_normalization(self):
        with self.assertRaises(ValueError):
            read_original_pdf_page(b"not a PDF", 1)
        self.assertEqual(_integer_page("123"), 123)
        self.assertIsNone(_integer_page("Sheet 1"))
        self.assertIsNone(_integer_page(True))
        self.assertIsNone(_integer_page(-1))


if __name__ == "__main__":
    unittest.main()
