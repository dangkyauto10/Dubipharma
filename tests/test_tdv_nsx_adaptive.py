"""Real NSX adaptive table recognition; fixtures are synthetic, never approvals."""
from __future__ import annotations

from io import BytesIO
import unittest

from docx import Document
from openpyxl import Workbook

from tdv_nsx_adaptive import extract_product_rows
from tdv_nsx_intake import analyze_nsx_documents


def multi_xlsx() -> bytes:
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Danh muc NSX"
    sheet.append(["BẢNG DANH MỤC THUỐC CỦA NHÀ SẢN XUẤT"])
    sheet.append(["STT", "Tên thuốc", "Hàm lượng", "Hoạt chất",
                  "Mã sản phẩm", "Nhà máy sản xuất", "Tên nhà sản xuất",
                  "Số đăng ký", "Dạng bào chế"])
    sheet.append([1, "Thuốc một", "1 mg", "Acenocoumarol",
                  "P001", "Nhà máy A", "NSX Alpha", "VN-001-26", "Viên nén"])
    sheet.append([2, "Thuốc hai", "4 mg", "Acenocoumarol",
                  "P002", "Nhà máy B", "NSX Beta", "VN-002-26", "Viên nén"])
    f = BytesIO()
    wb.save(f)
    return f.getvalue()


def single_quote_csv(code: str = "P002") -> bytes:
    return ("Mã sản phẩm;Đơn giá báo giá\n" + code + ";5000\n").encode("utf-8")


def proof_docx(kind="GMP", factory="Nhà máy B") -> bytes:
    doc = Document()
    tbl = doc.add_table(rows=0, cols=2)
    for a, b in [("Nhà máy sản xuất", factory),
                 ("Số chứng nhận GMP", "GMP-12345")]:
        cells = tbl.add_row().cells
        cells[0].text, cells[1].text = a, b
    f = BytesIO()
    doc.save(f)
    return f.getvalue()


class AdaptiveNSXTests(unittest.TestCase):
    def test_two_row_catalogue_requires_explicit_selection(self):
        raw = multi_xlsx()
        res = extract_product_rows("cat.xlsx", raw)
        self.assertEqual(res["total_candidates"], 2)
        self.assertTrue(res["complete"])
        self.assertEqual([x["fields"]["PRODUCT_ID"] for x in res["candidates"]],
                         ["P001", "P002"])
        self.assertEqual(res["candidates"][0]["row"], 3)
        self.assertIn("hàng 3, cột 5",
                      res["candidates"][0]["evidence"]["PRODUCT_ID"]["source"])
        intake = analyze_nsx_documents([
            ("PRODUCT", "cat.xlsx", raw),
            ("GMP_CERT", "gmp.docx", proof_docx()),
        ])
        self.assertTrue(intake["product"]["catalogue_selection_required"])
        self.assertEqual(intake["product"]["catalogue_selected_index"], None)
        self.assertEqual(intake["links"][0]["status"], "CHUA_CHON_SAN_PHAM")
        self.assertFalse(intake["gmp_confirmed"])

    def test_selected_row_uses_correct_factory(self):
        files = [
            ("PRODUCT", "cat.xlsx", multi_xlsx()),
            ("GMP_CERT", "gmp.docx", proof_docx()),
        ]
        selected = analyze_nsx_documents(files, selected_product_index=1)
        self.assertEqual(selected["product"]["fields"]["PRODUCT_ID"], "P002")
        self.assertEqual(selected["product"]["fields"]["FACTORY"], "Nhà máy B")
        self.assertEqual(selected["links"][0]["status"], "LIEN_KET_SO_BO_CAN_XAC_MINH")
        self.assertFalse(selected["links"][0]["gmp_confirmed"])
        wrong = analyze_nsx_documents(files, selected_product_index=0)
        self.assertEqual(wrong["links"][0]["status"], "XUNG_DOT_DINH_DANH")

    def test_bad_index_never_silently_selects_first_product(self):
        res = analyze_nsx_documents(
            [("PRODUCT", "cat.xlsx", multi_xlsx()),
             ("GMP_CERT", "gmp.docx", proof_docx())],
            selected_product_index=90,
        )
        self.assertIn("error", res["product"])
        self.assertEqual(res["links"][0]["status"], "CHUA_DU_CHUNG_CU")

    def test_docx_multicolumn_detected(self):
        doc = Document()
        t = doc.add_table(rows=0, cols=4)
        for values in [
            ("Mã sản phẩm", "Tên thuốc", "Hoạt chất", "Hàm lượng"),
            ("X1", "Viên A", "Acenocoumarol", "1 mg"),
            ("X2", "Viên B", "Acenocoumarol", "2 mg"),
        ]:
            cells = t.add_row().cells
            for c, value in zip(cells, values):
                c.text = value
        f = BytesIO()
        doc.save(f)
        res = extract_product_rows("catalogue.docx", f.getvalue())
        self.assertEqual(res["total_candidates"], 2)
        self.assertIn("Bảng Word 1", res["candidates"][1]["evidence"]["STRENGTH"]["source"])

    def test_csv_semicolon_multicolumn_and_bom(self):
        payload = ("\ufeffMã sản phẩm;Tên thuốc;Hoạt chất;Hàm lượng\n"
                   "A10;Thuốc A;Ciprofloxacin;500 mg\n"
                   "A20;Thuốc B;Amoxicillin;250 mg\n").encode("utf-8")
        res = extract_product_rows("multiple.csv", payload)
        self.assertEqual(res["total_candidates"], 2)
        self.assertEqual(res["candidates"][1]["fields"]["INGREDIENT"], "Amoxicillin")

    def test_single_candidate_promoted_with_trace(self):
        wb = Workbook()
        sh = wb.active
        sh.append(["Tên thuốc", "Hoạt chất", "Số đăng ký", "Nhà máy sản xuất"])
        sh.append(["Thuốc A", "Paracetamol", "VN-101-26", "Nhà máy A"])
        dest = BytesIO()
        wb.save(dest)
        report = analyze_nsx_documents([
            ("PRODUCT", "product.xlsx", dest.getvalue()),
            ("GMP_CERT", "gmp.docx", proof_docx(factory="Nhà máy A")),
        ])
        self.assertFalse(report["product"]["catalogue_selection_required"])
        self.assertEqual(report["product"]["fields"]["PRODUCT_NAME"], "Thuốc A")
        self.assertIn("hàng 2", report["product"]["evidence"]["INGREDIENT"]["source"])
        self.assertEqual(report["links"][0]["status"], "LIEN_KET_SO_BO_CAN_XAC_MINH")

    def test_csv_quote_table_cannot_fake_proven_match(self):
        report = analyze_nsx_documents([
            ("PRODUCT", "cat.xlsx", multi_xlsx()),
            ("QUOTE", "quote.csv", single_quote_csv()),
        ], selected_product_index=1)
        self.assertEqual(report["links"][0]["status"], "CHUA_DU_CHUNG_CU")
        self.assertFalse(report["quote_approved"])
        # Multi-product quote requires an independent per-product extraction step.


if __name__ == "__main__":
    unittest.main()
