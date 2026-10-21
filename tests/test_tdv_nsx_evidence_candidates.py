"""Synthetic legal fail-closed validation for SĐK and GMP source discovery.

No original manufacturer documents, tender data, company secrets or personal
registration identifiers are included in this PUBLIC repo.
"""
from datetime import date
from io import BytesIO
import unittest

from openpyxl import load_workbook
from tdv_nsx_evidence_candidates import discover_evidence_candidates
from tdv_nsx_auto import export_auto_xlsx


def record(fields, *, uid="", kind="", filename="source.pdf", sha="a"*64,
           complete=True, conflicts=None):
    return {
        "uid":uid, "kind":kind, "fields":dict(fields), "file":filename,
        "filename":filename, "sha256":sha, "complete_read":complete,
        "conflicts":dict(conflicts or {}),
        "evidence":{key:{"file":filename,"sha256":sha,
                         "source":f"Trang 1, dòng {i+1}", "raw":value,
                         "value":value, "verified":False}
                    for i,(key,value) in enumerate(fields.items())},
        "verified":False
    }


def classified(record):
    return {"kind":record["kind"],"filename":record["filename"],
            "sha256":record["sha256"],"complete_read":record["complete_read"]}


def run(products, sources, *, today=date(2026, 10, 22)):
    return discover_evidence_candidates(products, [classified(r) for r in sources],
                                        sources, reference_date=today)


class EvidenceDiscoveryQA(unittest.TestCase):
    def setUp(self):
        self.product = record({"PRODUCT_NAME":"Medicine A", "STRENGTH":"50 mg",
                               "FACTORY":"Site A", "MANUFACTURER":"Factory Group A"},
                              uid="P0001",filename="catalogue.xlsx",sha="b"*64)

    def test_registration_supported_and_cited_but_never_approved(self):
        s=record({"PRODUCT_NAME":"Medicine A", "STRENGTH":"50 mg",
                  "REGISTRATION_NO":"REG-2026-X"},
                 kind="REGISTRATION")
        result=run([self.product],[s])
        self.assertEqual(result["summary"]["registration_hints"],1)
        c=result["registrations"][0]
        self.assertEqual(c["status"],"GOI_Y_CO_DINH_DANH_PHU_CHUA_XAC_MINH")
        self.assertEqual(c["registration_no_in_proof"],"REG-2026-X")
        self.assertEqual(c["supporting_fields"],["STRENGTH"])
        self.assertEqual(c["product_source"]["location"],"Trang 1, dòng 1")
        self.assertEqual(c["registration_source"]["location"],"Trang 1, dòng 3")
        self.assertEqual(result["independent_proof_links_created"],0)
        self.assertFalse(result["legal_approval"])
        self.assertFalse(c["registration_confirmed"])
        self.assertFalse(c["bid_eligible"])

    def test_name_only_is_weak_hint_and_no_sdk_guess(self):
        s=record({"PRODUCT_NAME":"Medicine A"},kind="REGISTRATION")
        c=run([self.product],[s])["registrations"][0]
        self.assertEqual(c["status"],"GOI_Y_CHI_TRUNG_TEN_CHUA_XAC_MINH")
        self.assertEqual(c["registration_no_in_proof"],"")
        self.assertEqual(c["supporting_fields"],[])

    def test_similar_name_is_not_exact_product(self):
        s=record({"PRODUCT_NAME":"Medicine A Plus"},kind="REGISTRATION")
        data=run([self.product],[s])
        self.assertEqual(data["summary"]["registration_hints"],0)
        self.assertEqual(data["unresolved"][0]["status"],"CHUA_TIM_THAY_TEN_TRUNG_CHINH_XAC")

    def test_conflicting_strength_cannot_be_promoted(self):
        s=record({"PRODUCT_NAME":"Medicine A","STRENGTH":"100 mg"},kind="REGISTRATION")
        c=run([self.product],[s])["registrations"][0]
        self.assertEqual(c["status"],"MAU_THUAN_DINH_DANH")
        self.assertEqual(c["conflict_fields"],["STRENGTH"])
        self.assertFalse(c["verified"])

    def test_two_same_product_names_must_remain_ambiguous(self):
        other=record({"PRODUCT_NAME":"Medicine A","STRENGTH":"100 mg"},
                     uid="P0002",filename="other.xlsx",sha="c"*64)
        s=record({"PRODUCT_NAME":"Medicine A"},kind="REGISTRATION")
        hints=run([self.product,other],[s])["registrations"]
        self.assertEqual(len(hints),2)
        self.assertTrue(all(r["status"]=="TRUNG_TEN_NHIEU_UNG_VIEN" for r in hints))

    def test_uncited_fields_do_not_form_identity(self):
        s=record({"PRODUCT_NAME":"Medicine A"},kind="REGISTRATION")
        s["evidence"].pop("PRODUCT_NAME")
        data=run([self.product],[s])
        self.assertEqual(data["summary"]["registration_hints"],0)
        self.assertEqual(data["unresolved"][0]["status"],"THIEU_TEN_THUOC_CO_NGUON")

    def test_self_sourced_registration_document_never_cross_validates(self):
        s=record({"PRODUCT_NAME":"Medicine A","REGISTRATION_NO":"REG-X"},
                 uid="P0001",kind="REGISTRATION",filename="reg.pdf",sha="d"*64)
        s["registration_document_source"]=True
        data=run([s],[s])
        self.assertEqual(data["summary"]["registration_hints"],0)
        self.assertEqual(data["summary"]["unresolved_evidence"],1)

    def test_factory_level_ref_shared_by_multiple_drugs_not_per_drug_approval(self):
        second=record({"PRODUCT_NAME":"Medicine B","FACTORY":"Site A"},
                      uid="P0002",filename="catalogue.xlsx",sha="b"*64)
        cert=record({"FACTORY":"Site A","GMP_NUMBER":"GMP-X",
                     "GMP_SCOPE":"Tablets", "GMP_EXPIRY":"30/09/2026"},
                    kind="GMP_CERT")
        data=run([self.product,second],[cert])
        self.assertEqual(data["summary"]["factory_source_groups"],1)
        g=data["factories"][0]
        self.assertEqual(g["related_product_count"],2)
        self.assertEqual(g["related_product_ids"],["P0001","P0002"])
        self.assertEqual(g["expiry_observation"],"QUA_NGAY_GHI_TREN_GIAY_CAN_XAC_MINH_GIA_HAN")
        self.assertFalse(g["gmp_confirmed"])
        self.assertFalse(g["bid_eligible"])
        self.assertEqual(data["independent_proof_links_created"],0)

    def test_unclear_expiry_never_turns_into_valid_gmp(self):
        cert=record({"FACTORY":"Site A","GMP_EXPIRY":"unknown"},kind="GMP_CERT")
        g=run([self.product],[cert])["factories"][0]
        self.assertEqual(g["expiry_observation"],"NGAY_GHI_TREN_GIAY_CAN_DOC_BAN_GOC")
        self.assertFalse(g["gmp_confirmed"])

    def test_future_expiry_only_calendar_observation_not_validity(self):
        cert=record({"FACTORY":"Site A","GMP_EXPIRY":"2027-10-22"},kind="GMP_CERT")
        g=run([self.product],[cert])["factories"][0]
        self.assertIn("KHONG_PHAI_DA_XAC_NHAN_GMP",g["expiry_observation"])
        self.assertFalse(g["gmp_confirmed"])

    def test_manufacturer_name_cannot_replace_missing_factory(self):
        cert=record({"MANUFACTURER":"Factory Group A"},kind="GMP_CERT")
        data=run([self.product],[cert])
        self.assertEqual(data["summary"]["factory_source_groups"],0)
        self.assertEqual(data["unresolved"][0]["status"],"THIEU_NHA_MAY_CO_TRICH_DAN")

    def test_incomplete_document_is_not_suggested(self):
        cert=record({"FACTORY":"Site A"},kind="GMP_CERT",complete=False)
        data=run([self.product],[cert])
        self.assertEqual(data["summary"]["factory_source_groups"],0)
        self.assertEqual(data["unresolved"][0]["status"],"CAN_KIEM_TRA_BAN_GOC")

    def test_export_separate_review_tabs_without_changing_link_count(self):
        reg=record({"PRODUCT_NAME":"Medicine A","STRENGTH":"50 mg"},kind="REGISTRATION")
        cert=record({"FACTORY":"Site A"},kind="GMP_CERT",filename="gmp.pdf",sha="e"*64)
        proposal=run([self.product],[reg,cert])
        report={
            "files":[classified(reg),classified(cert)],"products":[self.product],
            "sources":[reg,cert],"links":[],"review_queue":[],
            "traceability":{"products":[],"checklist":[],"unassigned_files":[]},
            "tender_documents":[],"catalogue_audit":{},
            "warnings":[],"evidence_candidates":proposal,
        }
        data=export_auto_xlsx(report)
        wb=load_workbook(BytesIO(data),read_only=True)
        try:
            self.assertIn("GOI_Y_SDK_CHUA_XAC_MINH",wb.sheetnames)
            self.assertIn("GMP_THEO_NHA_MAY",wb.sheetnames)
            self.assertIn("CHUNG_CU_CHUA_GHEP",wb.sheetnames)
            self.assertEqual(wb["LIEN_KET_CHUNG_CU"].max_row,1)
            self.assertEqual(wb["GMP_THEO_NHA_MAY"].max_row,2)
        finally:
            wb.close()

if __name__ == "__main__":
    unittest.main()
