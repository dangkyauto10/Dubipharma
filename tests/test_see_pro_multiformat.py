"""Read-only regression checks for SEE PRO unified intake."""
from __future__ import annotations
from io import BytesIO
import unittest
from unittest.mock import patch
import pymupdf
from see_pro_multiformat import extract_file, identify_codes
from see_pro_workbench import analyze_uploads


class SeeProMultiFormatTests(unittest.TestCase):
    @staticmethod
    def summary_pdf():
        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((50, 70),
            "362 PP2600428741 G1.362 Hydroxypropylmethylcellulose 565.800.000 8.487.000")
        data = pdf.tobytes()
        pdf.close()
        return data

    def test_pdf_summary_evidence(self):
        data = self.summary_pdf()
        result = analyze_uploads([("hsmt.pdf", data)])
        self.assertEqual(len(result["documents"]), 1)
        self.assertEqual(result["pdf_reports"][0]["Dòng tổng hợp"], 1)
        self.assertEqual(result["rows"][0]["MA_PHAN_LO"], "PP2600428741")
        self.assertEqual(result["rows"][0]["GIA_UOC_TINH_PHAN"], 565800000)
        self.assertNotEqual(result["rows"][0]["TRANG_THAI"], "DA_XAC_MINH")

    def test_docx_paragraphs_and_table(self):
        from docx import Document
        doc = Document()
        doc.add_paragraph("Mã thuốc: G1.003")
        t = doc.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "G1.004"
        t.cell(0, 1).text = "4 mg"
        out = BytesIO()
        doc.save(out)
        result = extract_file("ho_so.docx", out.getvalue())
        self.assertEqual({x["Mã lô"] for x in identify_codes(result)}, {"G1.003", "G1.004"})
        self.assertFalse(result["verified"])

    def test_xlsx_multiple_sheets(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active["A1"] = "G1.011"
        ws = book.create_sheet("Technical")
        ws["A2"] = "G1.012"
        out = BytesIO()
        book.save(out)
        result = extract_file("bang.xlsx", out.getvalue())
        self.assertEqual({x["Mã lô"] for x in identify_codes(result)}, {"G1.011", "G1.012"})

    def test_xlsx_blank_numeric_text_and_code_cells(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet.title = "Lots"
        sheet["A1"] = "G1.036"
        sheet["B1"] = 0
        sheet["C1"] = 1250
        sheet["D1"] = "   "
        sheet["E1"] = "Acenocoumarol"
        sheet["A3"] = "G1.037"
        output = BytesIO()
        book.save(output)
        result = extract_file("lots.xlsx", output.getvalue())
        self.assertEqual({x["Mã lô"] for x in identify_codes(result)}, {"G1.036", "G1.037"})
        self.assertTrue(result["complete"])
        self.assertFalse(result["verified"])
        first = next(x["text"] for x in result["lines"] if x["source"] == "Sheet Lots, dòng 1")
        self.assertIn("1:G1.036", first)
        self.assertIn("2:0", first)
        self.assertIn("3:1250", first)
        self.assertIn("5:Acenocoumarol", first)
        self.assertNotIn("ReadOnlyCell", first)
        self.assertNotIn("4:   ", first)
        self.assertFalse(any(x["source"] == "Sheet Lots, dòng 2" for x in result["lines"]))

    def test_xlsx_formula_without_cached_result_is_not_invented(self):
        from openpyxl import Workbook
        book = Workbook()
        sheet = book.active
        sheet["A1"] = "G1.038"
        sheet["B1"] = "=100+200"
        output = BytesIO()
        book.save(output)
        result = extract_file("formula.xlsx", output.getvalue())
        self.assertEqual({x["Mã lô"] for x in identify_codes(result)}, {"G1.038"})
        combined = " ".join(x["text"] for x in result["lines"])
        self.assertNotIn("300", combined)
        self.assertNotIn("=100+200", combined)

    def test_csv(self):
        result = extract_file("items.csv", "Mã,Hoạt chất\nG1.021,Acenocoumarol".encode("utf-8"))
        self.assertEqual(identify_codes(result)[0]["Mã lô"], "G1.021")

    def test_photo_without_ocr_remains_unverified(self):
        from PIL import Image
        picture = Image.new("RGB", (100, 100), "white")
        out = BytesIO()
        picture.save(out, format="PNG")
        result = extract_file("label.png", out.getvalue(), ocr=False)
        self.assertEqual(result["lines"], [])
        self.assertFalse(result["complete"])
        self.assertFalse(result["verified"])

    def test_photo_ocr_does_not_claim_verification(self):
        with patch("see_pro_multiformat._image_ocr", return_value=("G1.362", 180)):
            result = extract_file("label.png", b"placeholder", ocr=True)
        self.assertIn("G1.362", result["lines"][0]["text"])
        self.assertIn("CHƯA XÁC MINH", " ".join(result["warnings"]))
        self.assertFalse(result["verified"])

    def test_untrusted_extension_rejected(self):
        with self.assertRaises(ValueError):
            extract_file("run.exe", b"some data")

    def test_no_bid_price_inference(self):
        result = analyze_uploads([("hsmt.pdf", self.summary_pdf())])
        self.assertNotIn("DON_GIA_DU_THAU", result["rows"][0])
        self.assertFalse(result["documents"][0]["Xác minh"] == "ĐẠT")


if __name__ == "__main__":
    unittest.main()
