"""Synthetic only: leadership work ordering never upgrades SĐK/GMP or bid state."""
import unittest
from io import BytesIO

from openpyxl import load_workbook
from docx import Document

from tdv_review_priorities import build_review_priorities
from tdv_leadership_report import (
    compare_sources, export_leadership_xlsx, export_leadership_docx,
    GREEN, YELLOW, RED
)


def proof(file_name, pos="Trang 1", raw="Source text"):
    return {"file": file_name, "location": pos,
            "sha256": "a"*64, "raw": raw}


def criterion(status, *, label="Hàm lượng", lhs="50 mg", rhs="100 mg"):
    return {"label": label, "ehsmt": lhs, "nsx": rhs, "status": status,
            "ehsmt_evidence": proof("E-HSMT.pdf","trang 8"),
            "nsx_evidence": proof("NSX.xlsx","sheet 1 hàng 4"),
            "note": "CHUA_DUYET"}


def row(code="G1.001", *, candidates=0, hints=0, yellow=4,
        criteria=None):
    return {
        "MA_THUOC_EHSMT": code, "NGUON_EHSMT": "E-HSMT.pdf",
        "VI_TRI_EHSMT": "Trang 8", "NGUON_THUOC_NSX": "NSX.xlsx",
        "SO_UNG_VIEN_NSX": candidates, "SO_GOI_Y_TEN_NSX": hints,
        "SO_TIEU_CHI_CHUA_DAT": yellow,
        "GOI_Y_TEN_NSX": [{
            "source_file":"NSX.xlsx","source_location":"Sheet 1 hàng 2"
        }] if hints else [],
        "criteria":criteria or [],
    }


def summary():
    return {
        "ehsmt_lots":2, "nsx_products":1, "nsx_files":2,
        "nsx_review":3, "green_criteria_preliminary":0,
        "yellow_criteria":7, "red_mismatch_preliminary":1,
        "total_criteria":8, "matching_percent_preliminary":None,
        "criteria_ehsmt_available":8, "lots_with_candidate":1,
        "lots_uniquely_linked":1, "lots_with_title_hints":0,
        "linked_requirements":4, "fully_verified":0,
    }


def example_report():
    return {
        "created_at_vn":"22/10/2026 07:00:00 +07:00",
        "mode":"REAL_SOURCES_ONLY_NO_DEMO",
        "summary":summary(),
        "comparison":[row("G1.001", candidates=1, yellow=3,
                           criteria=[criterion(RED)]),
                      row("G1.002", candidates=0, yellow=4)],
        "nsx_evidence_candidates":{
            "registrations":[{
                "product_id":"P0001", "status":"MAU_THUAN_DINH_DANH",
                "document_file":"sdk.pdf",
                "document_source":proof("sdk.pdf", "trang 2"),
                "conflict_fields":["STRENGTH"],
            }],
            "factories":[{
                "factory":"Site A","related_product_count":3,
                "expiry_observation":"QUA_NGAY_GHI_TREN_GIAY_CAN_XAC_MINH_GIA_HAN",
                "document_file":"gmp.pdf",
                "factory_source":proof("gmp.pdf","trang 3"),
            }],
            "unresolved":[{
                "file":"unknown.pdf", "reason":"Không trích được SĐK có nhãn"
            }]
        },
        "traceability": {"products":[],"checklist":[]},
        "ehsmt_documents":[], "nsx_documents":[],
        "warnings":[], "note":"Tất cả chưa xác minh",
        "verified":False, "bid_eligible":False,
        "gmp_approved":False, "requirements_checked":False,
    }


class LeadershipTriageQA(unittest.TestCase):
    def test_technical_disagreement_before_unmatched_drugs(self):
        got=build_review_priorities(example_report())
        self.assertEqual(got["by_priority"]["1"],2)
        self.assertEqual(got["by_priority"]["2"],3)
        self.assertEqual(got["by_priority"]["3"],1)
        self.assertEqual(got["total_issues"],6)
        self.assertEqual(got["rows"][0]["muc_uu_tien"],1)
        self.assertEqual(got["rows"][-1]["nhom"],"CHUA_CO_UNG_VIEN")
        self.assertTrue(all(x["trang_thai"]=="CAN_XAC_MINH" for x in got["rows"]))
        self.assertFalse(got["bid_eligible"])
        self.assertFalse(got["verified"])

    def test_name_hint_only_never_true_product_link(self):
        got=build_review_priorities({"comparison":[row(hints=2)]})
        self.assertEqual(got["total_issues"],1)
        self.assertEqual(got["rows"][0]["nhom"],"GOI_Y_TEN")
        self.assertEqual(got["rows"][0]["tai_lieu"],"NSX.xlsx")
        self.assertEqual(got["rows"][0]["muc_uu_tien"],3)

    def test_multiple_candidates_never_auto_choose_one(self):
        got=build_review_priorities({"comparison":[row(candidates=3,yellow=4)]})
        self.assertEqual(got["rows"][0]["nhom"],"NHIEU_UNG_VIEN")
        self.assertIn("không tự chọn",got["rows"][0]["viec_can_lam"])

    def test_many_medicines_keep_full_counts_with_bounded_view(self):
        records=[row(f"G1.{i:03d}") for i in range(1,771)]
        got=build_review_priorities({"comparison":records},max_rows=5)
        self.assertEqual(got["total_issues"],770)
        self.assertEqual(got["shown_issues"],5)
        self.assertTrue(got["truncated"])
        self.assertEqual(got["by_priority"]["3"],770)
        self.assertFalse(got["bid_eligible"])

    def test_export_xlsx_includes_source_aware_action_queue(self):
        report=example_report()
        book=load_workbook(BytesIO(export_leadership_xlsx(report)),read_only=True)
        try:
            self.assertIn("UU_TIEN_XU_LY",book.sheetnames)
            actions=list(book["UU_TIEN_XU_LY"].values)
            self.assertEqual(len(actions),8) # header, 6 issues, totals
            self.assertEqual(actions[0][0],"Ưu tiên")
            self.assertEqual(actions[1][8],"CAN_XAC_MINH")
            self.assertEqual(actions[1][9],"CHUA_XAC_MINH")
            self.assertIn("GMP",str(actions))
        finally:
            book.close()

    def test_word_mentions_actionable_review_and_no_legality(self):
        report=example_report()
        document=Document(BytesIO(export_leadership_docx(report)))
        content="\n".join(p.text for p in document.paragraphs)
        self.assertIn("Thứ tự xử lý chứng cứ",content)
        self.assertIn("UU_TIEN_XU_LY",content)
        self.assertIn("không phải kết luận pháp lý",content.lower())
        self.assertFalse(report["bid_eligible"])

    def test_real_comparison_retains_original_color_and_review_hints_separately(self):
        product={
            "uid":"P0001", "file":"nsx.xlsx", "sha256":"b"*64,
            "fields":{"PRODUCT_NAME":"Drug A", "INGREDIENT":"Acarbose",
                      "STRENGTH":"50 mg"},
            "evidence":{"INGREDIENT":{"file":"nsx.xlsx","source":"Hàng 3",
                                     "raw":"Acarbose","sha256":"b"*64},
                        "STRENGTH":{"file":"nsx.xlsx","source":"Hàng 3",
                                    "raw":"50 mg","sha256":"b"*64}},
            "conflicts":{},
        }
        lot={
            "MA_THUOC":"G1.001", "HOAT_CHAT_HSMT":"Acarbose",
            "NONG_DO_HAM_LUONG":"50 mg", "DANG_BAO_CHE":"Viên nén",
            "DUONG_DUNG":"Uống", "FILE_NGUON":"E-HSMT.pdf",
        }
        source={"products":[product],
                "files":[], "traceability":{},
                "evidence_candidates":{"registrations":[{
                    "product_id":"P0001", "status":"GOI_Y_CHI_TRUNG_TEN_CHUA_XAC_MINH",
                    "document_source":proof("sdk.pdf"),"conflict_fields":[],
                }]}}
        report=compare_sources({"rows":[lot]}, source)
        self.assertEqual(report["summary"]["fully_verified"],0)
        self.assertEqual(report["comparison"][0]["SO_TIEU_CHI_DAT"],2)
        self.assertTrue(report["nsx_evidence_candidates"]["registrations"])
        got=build_review_priorities(report)
        self.assertIn("SDK_CHUA_XAC_MINH",got["by_category"])
        self.assertFalse(got["bid_eligible"])

if __name__=="__main__":
    unittest.main()
