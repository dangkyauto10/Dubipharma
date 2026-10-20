"""Read-only NSX intake regression: source evidence and no false GMP approval."""
from __future__ import annotations
from io import BytesIO
import hashlib
import unittest
import pymupdf
from docx import Document
from openpyxl import Workbook, load_workbook

from tdv_nsx_intake import analyze_nsx_documents, export_nsx_xlsx, _evidence_pair


def docx_pairs(pairs: list[tuple[str, str]]) -> bytes:
    doc = Document()
    table = doc.add_table(rows=0, cols=2)
    for a, b in pairs:
        row = table.add_row()
        row.cells[0].text, row.cells[1].text = a, b
    dest = BytesIO()
    doc.save(dest)
    return dest.getvalue()


def product_file() -> bytes:
    return docx_pairs([
        ("Mã sản phẩm", "SP-100"),
        ("Số đăng ký", "VN-123-26"),
        ("Tên thuốc", "Acenoral"),
        ("Hoạt chất", "Acenocoumarol"),
        ("Hàm lượng", "4 mg"),
        ("Dạng bào chế", "Viên nén"),
        ("Đường dùng", "Uống"),
        ("Tên nhà sản xuất", "NSX A"),
        ("Nhà máy sản xuất", "Nhà máy số 1"),
    ])


def xlsx_pairs(pairs: list[tuple[str, str]]) -> bytes:
    wb = Workbook()
    sheet = wb.active
    for a, b in pairs:
        sheet.append([a, b])
    dest = BytesIO()
    wb.save(dest)
    return dest.getvalue()


def gmp_file(factory="Nhà máy số 1") -> bytes:
    return xlsx_pairs([
        ("Nhà máy sản xuất", factory),
        ("Số giấy chứng nhận GMP", "GMP-7788"),
        ("Phạm vi GMP", "Thuốc viên nén"),
        ("Ngày hết hạn GMP", "31/12/2027"),
    ])


def pdf_lines(lines: list[str]) -> bytes:
    doc = pymupdf.open()
    p = doc.new_page()
    for idx, line in enumerate(lines):
        p.insert_text((52, 52 + idx * 23), line, fontsize=11)
    result = doc.tobytes()
    doc.close()
    return result


class NSXSourceIntakeTests(unittest.TestCase):
    def test_mixed_docx_xlsx_pdf_quote_no_gmp_auto_approval(self):
        quote = pdf_lines([
            "Product ID: SP-100",
            "Registration Number: VN-123-26",
            "Quoted Unit Price: 12345",
            "Currency: VND",
        ])
        report = analyze_nsx_documents([
            ("PRODUCT", "product.docx", product_file()),
            ("GMP_CERT", "gmp.xlsx", gmp_file()),
            ("QUOTE", "bao_gia.pdf", quote),
        ])
        self.assertEqual(report["summary"]["files_received"], 3)
        self.assertEqual(report["summary"]["files_parsed"], 3)
        self.assertEqual(report["summary"]["links_provisional"], 2)
        self.assertFalse(report["gmp_confirmed"])
        self.assertFalse(report["quote_approved"])
        self.assertFalse(report["bid_eligible"])
        p = report["product"]
        self.assertEqual(p["fields"]["REGISTRATION_NO"], "VN-123-26")
        self.assertEqual(p["fields"]["MANUFACTURER"], "NSX A")
        self.assertEqual(p["fields"]["FACTORY"], "Nhà máy số 1")
        ev = p["evidence"]["PRODUCT_ID"]
        self.assertEqual(ev["sha256"], hashlib.sha256(product_file()).hexdigest())
        self.assertIn("Bảng", ev["source"])
        self.assertFalse(ev["verified"])
        quote_report = report["documents"][-1]
        self.assertEqual(quote_report["fields"]["QUOTE_UNIT_PRICE"], "12345")
        self.assertFalse(quote_report["quote_approved"])

    def test_different_factory_is_blocked_not_auto_reconciled(self):
        report = analyze_nsx_documents([
            ("PRODUCT", "product.docx", product_file()),
            ("GMP_CERT", "other.xlsx", gmp_file("Nhà máy khác")),
        ])
        self.assertEqual(report["links"][0]["status"], "XUNG_DOT_DINH_DANH")
        self.assertEqual(report["links"][0]["conflicting_fields"], ["FACTORY"])
        self.assertFalse(report["links"][0]["gmp_confirmed"])

    def test_quote_without_product_identifier_remains_unlinked(self):
        raw = pdf_lines(["Quoted Unit Price: 1200", "Currency: USD"])
        report = analyze_nsx_documents([
            ("PRODUCT", "product.docx", product_file()),
            ("QUOTE", "quote.pdf", raw),
        ])
        self.assertEqual(report["links"][0]["status"], "CHUA_DU_CHUNG_CU")
        self.assertEqual(report["links"][0]["matching_fields"], [])
        self.assertFalse(report["quote_approved"])

    def test_product_registration_can_anchor_quote_if_no_internal_id(self):
        raw = docx_pairs([
            ("Số đăng ký", "VN-123-26"),
            ("Đơn giá báo giá", "18.500 VND"),
        ])
        report = analyze_nsx_documents([
            ("PRODUCT", "product.docx", product_file()),
            ("QUOTE", "quote.docx", raw),
        ])
        self.assertEqual(report["links"][0]["status"], "LIEN_KET_SO_BO_CAN_XAC_MINH")
        self.assertFalse(report["links"][0]["quote_approved"])

    def test_conflicting_values_not_silently_overwritten(self):
        raw = docx_pairs([
            ("Mã sản phẩm", "SP-100"),
            ("Mã sản phẩm", "SP-200"),
            ("Nhà máy sản xuất", "Nhà máy số 1"),
        ])
        report = analyze_nsx_documents([
            ("PRODUCT", "p.docx", raw),
            ("QUOTE", "q.pdf", pdf_lines(["Product ID: SP-100"])),
        ])
        self.assertIn("PRODUCT_ID", report["product"]["conflicts"])
        self.assertEqual(report["links"][0]["status"], "CHUA_DU_CHUNG_CU")
        self.assertTrue(report["warnings"])

    def test_word_header_is_not_confused_for_field_value(self):
        self.assertEqual(_evidence_pair(
            "Mã sản phẩm | Tên thuốc | Hoạt chất | Hàm lượng"), [])
        self.assertEqual(_evidence_pair(
            "1:Mã sản phẩm | 2:SP-100"), [("PRODUCT_ID", "SP-100")])
        self.assertEqual(_evidence_pair(
            "Tên nhà sản xuất: NSX A"), [("MANUFACTURER", "NSX A")])

    def test_spreadsheet_headers_do_not_fake_product_data(self):
        sheet = Workbook().active
        sheet.append(["Mã sản phẩm", "Tên thuốc", "Hoạt chất"])
        sheet.append(["SP-200", "Drug 2", "Acenocoumarol"])
        dest = BytesIO()
        sheet.parent.save(dest)
        report = analyze_nsx_documents([("PRODUCT", "listing.xlsx", dest.getvalue())])
        self.assertNotIn("PRODUCT_ID", report["product"]["fields"])
        self.assertFalse(report["verified"])

    def test_authorization_from_other_manufacturer_not_accepted(self):
        proof = docx_pairs([
            ("Mã sản phẩm", "SP-100"),
            ("Nhà sản xuất", "NSX B"),
            ("Hiệu lực ủy quyền đến", "31/12/2028"),
        ])
        report = analyze_nsx_documents([
            ("PRODUCT", "product.docx", product_file()),
            ("AUTHORIZATION", "auth.docx", proof),
        ])
        self.assertEqual(report["links"][0]["status"], "XUNG_DOT_DINH_DANH")
        self.assertIn("MANUFACTURER", report["links"][0]["conflicting_fields"])

    def test_explanation_is_not_certificate_and_registration_not_verified(self):
        report = analyze_nsx_documents([
            ("PRODUCT", "product.docx", product_file()),
            ("GMP_EXPLANATION", "letter.xlsx", gmp_file()),
            ("REGISTRATION", "reg.docx", docx_pairs(
                [("Số đăng ký", "VN-123-26")]))
        ])
        self.assertFalse(report["gmp_confirmed"])
        self.assertFalse(report["registration_confirmed"])
        self.assertIn("Công văn giải trình không thay thế", " ".join(report["links"][0]["reasons"]))

    def test_file_type_and_product_count_guard(self):
        with self.assertRaises(ValueError):
            analyze_nsx_documents([("QUOTE", "q.pdf", b"abc")])
        with self.assertRaises(ValueError):
            analyze_nsx_documents([
                ("PRODUCT", "p.docx", product_file()),
                ("PRODUCT", "p2.docx", product_file()),
            ])
        with self.assertRaises(ValueError):
            analyze_nsx_documents([])

    def test_corrupt_attachment_is_a_warning_not_a_success(self):
        report = analyze_nsx_documents([
            ("PRODUCT", "product.docx", product_file()),
            ("GMP_CERT", "fake.xlsx", b"not-an-excel"),
        ])
        self.assertEqual(report["summary"]["files_parsed"], 1)
        self.assertEqual(report["links"][0]["status"], "CHUA_DU_CHUNG_CU")
        self.assertFalse(report["gmp_confirmed"])
        self.assertTrue(any("KHÔNG ĐỌC ĐƯỢC" in x for x in report["warnings"]))

    def test_export_source_ledger_no_excel_formula(self):
        raw = docx_pairs([
            ("Mã sản phẩm", "SP-100"),
            ("Tên thuốc", '=HYPERLINK("https://invalid.example")'),
            ("Nhà máy sản xuất", "Nhà máy số 1"),
        ])
        report = analyze_nsx_documents([("PRODUCT", "src.docx", raw)])
        excel = export_nsx_xlsx(report)
        wb = load_workbook(BytesIO(excel), data_only=False)
        self.assertEqual(wb.sheetnames,
                         ["DANH_SACH_TAI_LIEU", "DOI_CHIEU_DINH_DANH",
                          "CHUNG_CU_TUNG_TRUONG", "GIOI_HAN_KET_LUAN"])
        records = list(wb["CHUNG_CU_TUNG_TRUONG"].iter_rows(min_row=2))
        suspect = next(row for row in records if row[2].value == "PRODUCT_NAME")
        self.assertNotEqual(suspect[3].data_type, "f")
        self.assertTrue(str(suspect[3].value).startswith("'="))
        wb.close()

    def test_duplicate_sha_is_flagged(self):
        proof = gmp_file()
        report = analyze_nsx_documents([
            ("PRODUCT", "product.docx", product_file()),
            ("GMP_CERT", "first.xlsx", proof),
            ("GMP_CERT", "second.xlsx", proof),
        ])
        self.assertTrue(any("trùng nội dung SHA256" in x for x in report["warnings"]))


if __name__ == "__main__":
    unittest.main()
