"""NSX original-value conflict comparison; synthetic inputs only."""
import unittest
from io import BytesIO
from openpyxl import Workbook, load_workbook

from tdv_nsx_source_diff import compare_conflict_sources
from tdv_nsx_auto import analyze_auto, export_auto_xlsx


def synthetic_catalogue():
    wb = Workbook()
    first = wb.active
    first.title = "Biểu 05"
    first.append(["Tên thuốc", "Hàm lượng", "Nhà máy sản xuất", "Dạng bào chế"])
    first.append(["Drug A", "50 mg", "Nhà máy A", "Viên nén"])
    first.append(["Drug B", "500 mg/100 ml", "Nhà máy B", "Dung dịch tiêm"])
    second = wb.create_sheet("ÁP")
    second.append(["Tên thuốc", "Hàm lượng", "Nhà máy sản xuất", "Dạng bào chế"])
    second.append(["Drug A", "100 mg", "Nhà máy A", "Viên nén"])
    second.append(["Drug B", "500 mg/100 ml", "Nhà máy B", "Dung dịch truyền"])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class NsXSourceDiffQA(unittest.TestCase):
    def test_per_field_two_sided_values_keep_exact_original_sheet(self):
        report=analyze_auto([("fixture.xlsx",synthetic_catalogue())],enable_ocr=False)
        pairs=compare_conflict_sources(report["products"])
        self.assertEqual(len(pairs),2)
        checks={x["Tên thuốc"]: x for x in pairs}
        a=checks["Drug A"]
        self.assertEqual(a["Trường cần đối chiếu"],"STRENGTH")
        self.assertEqual({a["Giá trị nguồn 1"],a["Giá trị nguồn 2"]},{"50 mg","100 mg"})
        self.assertIn("Sheet Biểu 05",a["Ô nguồn 1"])
        self.assertIn("Sheet ÁP",a["Ô nguồn 2"])
        self.assertEqual(a["SHA256 nguồn 1"],a["SHA256 nguồn 2"])
        self.assertEqual(len(a["SHA256 nguồn 1"]),64)
        self.assertIn("CHƯA XÁC MINH",a["Kết luận"])
        b=checks["Drug B"]
        self.assertEqual(b["Trường cần đối chiếu"],"DOSAGE_FORM")
        self.assertEqual({b["Giá trị nguồn 1"],b["Giá trị nguồn 2"]},
                         {"Dung dịch tiêm","Dung dịch truyền"})

    def test_new_excel_sheet_preserves_values_and_sources(self):
        report=analyze_auto([("fixture.xlsx",synthetic_catalogue())],enable_ocr=False)
        data=export_auto_xlsx(report)
        book=load_workbook(BytesIO(data),read_only=True)
        try:
            self.assertIn("MAU_THUAN_GIA_TRI_GOC",book.sheetnames)
            ws=list(book["MAU_THUAN_GIA_TRI_GOC"].values)
            headers=ws[0]
            field=headers.index("Trường cần đối chiếu")
            original1=headers.index("Giá trị nguồn 1")
            original2=headers.index("Giá trị nguồn 2")
            source1=headers.index("Ô nguồn 1")
            source2=headers.index("Ô nguồn 2")
            self.assertEqual(len(ws),3)
            for row in ws[1:]:
                self.assertIn(row[field],{"STRENGTH","DOSAGE_FORM"})
                self.assertTrue(row[original1])
                self.assertTrue(row[original2])
                self.assertIn("Sheet Biểu 05",row[source1])
                self.assertIn("Sheet ÁP",row[source2])
            self.assertIn("GOP_DONG_NHIEU_SHEET",book.sheetnames)
            self.assertIn("TRUY_XUAT_TUNG_TRUONG",book.sheetnames)
        finally:
            book.close()

    def test_diagnostics_are_read_only_and_never_link_proofs(self):
        report=analyze_auto([("fixture.xlsx",synthetic_catalogue())],enable_ocr=False)
        before=[
            (p["uid"],dict(p["fields"]),dict(p["conflicts"]),len(p["source_rows"]))
            for p in report["products"]
        ]
        pairs=compare_conflict_sources(report["products"])
        after=[
            (p["uid"],dict(p["fields"]),dict(p["conflicts"]),len(p["source_rows"]))
            for p in report["products"]
        ]
        self.assertEqual(before,after)
        self.assertEqual(report["summary"]["linked"],0)
        self.assertFalse(report["registration_confirmed"])
        self.assertFalse(report["gmp_confirmed"])
        self.assertEqual(len(pairs),2)

    def test_multiple_sheets_dont_drop_third_source(self):
        p={"uid":"P0010","fields":{"PRODUCT_NAME":"Drug C"},
           "conflicts":{"STRENGTH":[{"raw":"B"}]},
           "source_rows":[
              {"file":"x.xlsx","sha256":"a"*64,"evidence":{"STRENGTH":{"raw":"A","source":"Sheet A, ô B1"}}},
              {"file":"x.xlsx","sha256":"a"*64,"evidence":{"STRENGTH":{"raw":"B","source":"Sheet B, ô B2"}}},
              {"file":"x.xlsx","sha256":"a"*64,"evidence":{"STRENGTH":{"raw":"C","source":"Sheet C, ô B3"}}},
           ]}
        pairs=compare_conflict_sources([p])
        self.assertEqual(len(pairs),2)
        self.assertEqual({x["Giá trị nguồn 2"] for x in pairs},{"B","C"})
        self.assertEqual({x["Giá trị nguồn 1"] for x in pairs},{"A"})

    def test_no_disputed_fields_means_no_spurious_diff(self):
        self.assertEqual(compare_conflict_sources([{
            "uid":"P0001","fields":{"PRODUCT_NAME":"No conflict"},
            "source_rows":[],"conflicts":{},
        }]),[])


if __name__=="__main__":
    unittest.main()
