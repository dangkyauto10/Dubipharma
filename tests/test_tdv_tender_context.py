"""TDV tender companion sources and leadership reverse-provenance QA.

Entirely synthetic fixtures; NOT real documents nor legal compliance verdicts.
"""
from __future__ import annotations
from io import BytesIO
import unittest

from docx import Document
from openpyxl import Workbook, load_workbook

from tdv_nsx_auto import analyze_auto, export_auto_xlsx
from tdv_tender_context import detect_procurement_document, index_tender_documents


def word(*lines, pairs=()):
    doc=Document()
    for line in lines:
        doc.add_paragraph(line)
    if pairs:
        table=doc.add_table(rows=0,cols=2)
        for label,value in pairs:
            cells=table.add_row().cells
            cells[0].text=label
            cells[1].text=value
    out=BytesIO()
    doc.save(out)
    return out.getvalue()


def catalogue():
    wb=Workbook()
    ws=wb.active
    ws.append(["Tên thuốc","Số đăng ký","Mã sản phẩm","Hoạt chất",
               "Nhà sản xuất","Nhà máy sản xuất","Đơn vị cung cấp",
               "Đơn vị đăng ký", "Tên nhà thầu"])
    ws.append(["Medicine C","VN-X-42","PRD-42","Ingredient Q",
               "Manufacturer X","Factory X","Supplier X",
               "Registration Holder X","Bidder X"])
    out=BytesIO()
    wb.save(out)
    return out.getvalue()


class TenderContextQA(unittest.TestCase):
    def test_detect_amendment_with_source_id_date_and_action(self):
        lines=[{"source":"Trang 1","text":"SỬA ĐỔI E-HSMT"} ,
               {"source":"Trang 1","text":"Mã gói thầu BP1234567890"},
               {"source":"Trang 1","text":"Ngày 12/10/2026"}]
        found=detect_procurement_document(lines)
        self.assertEqual(found["kind"],"SUA_DOI_EHSMT")
        self.assertEqual(found["package_ids"],["BP1234567890"])
        self.assertEqual(found["date_candidates"][0]["value"],"12/10/2026")
        self.assertEqual(found["revision_status"],"CHUA_XAC_MINH_PHIEN_BAN_HIEU_LUC")
        self.assertIn("phiên bản",found["suggested_action"])
        self.assertFalse(found["verified"])

    def test_do_not_classify_proof_from_filename_or_text_incidental_mention(self):
        self.assertIsNone(detect_procurement_document([{
            "source":"Trang 1","text":"Manufacturer authorization and GMP certificate",
        }]))
        lines=[{"source":f"Trang {i}","text":"Thông tin về nhà sản xuất"}
               for i in range(85)]
        lines.append({"source":"Trang 86","text":"Một ghi chú trích dẫn sửa đổi E-HSMT"})
        self.assertIsNone(detect_procurement_document(lines))

    def test_separate_packages_and_flag_unassigned(self):
        ids=lambda kind,names: {"kind":kind,"filename":kind,
                               "package_ids":names}
        x=index_tender_documents([
            ids("SUA_DOI_EHSMT",["BP1111111111"]),
            ids("LAM_RO_EHSMT",["BP2222222222"]),
            ids("PHU_LUC_KY_THUAT",[]),
            ids("THONG_BAO_MOI_THAU",["BP1111111111","BP2222222222"]),
        ])
        self.assertEqual(len(x["by_package"]),2)
        self.assertEqual(len(x["unassigned"]),2)
        self.assertEqual(len(x["by_package"]["BP1111111111"]),1)
        self.assertFalse(x["verified"])

    def test_bid_dossier_tracks_supplier_registration_holder_manufacturer_and_gmp(self):
        pack=analyze_auto([
            ("catalogue.xlsx", catalogue()),
            ("factory_proof.docx",word("GIẤY CHỨNG NHẬN GMP",
                 pairs=[("Nhà máy sản xuất","Factory X"),
                        ("Số giấy chứng nhận GMP","FAKE-GMP")]))
        ],enable_ocr=False)
        self.assertEqual(pack["summary"]["products"],1)
        self.assertEqual(pack["summary"]["linked"],1)
        dossier=pack["traceability"]["products"][0]
        parties=dossier["parties"]
        self.assertEqual(parties["SUPPLIER"]["value"],"Supplier X")
        self.assertEqual(parties["BIDDER"]["value"],"Bidder X")
        self.assertEqual(parties["REGISTRATION_HOLDER"]["value"],
                         "Registration Holder X")
        self.assertEqual(parties["MANUFACTURER"]["value"],"Manufacturer X")
        self.assertEqual(parties["FACTORY"]["value"],"Factory X")
        for party in parties.values():
            self.assertEqual(party["status"],"CO_TRICH_DAN_CHUA_XAC_MINH")
            self.assertIn("source",party)
            self.assertTrue(party["sha256"])
        documents={p["kind"]:p for p in dossier["proofs"]}
        self.assertEqual(documents["GMP_CERT"]["status"],
                         "CO_LIEN_KET_SO_BO_CHUA_XAC_MINH")
        self.assertEqual(documents["GMP_CERT"]["sources"][0]["file"],"factory_proof.docx")
        self.assertEqual(documents["QUOTE"]["status"],"CHUA_CO_LIEN_KET_CHUNG_CU")
        self.assertTrue(all(x["requirement_status"] ==
                            "CHUA_DOI_CHIEU_EHSMT_CUA_GOI_THAU"
                            for x in dossier["proofs"]))
        self.assertFalse(pack["traceability"]["complete"])
        self.assertFalse(pack["gmp_confirmed"])
        book=load_workbook(BytesIO(export_auto_xlsx(pack)),read_only=True)
        self.assertEqual(len(book.sheetnames),11)
        self.assertEqual(book["TRUY_CUU_LANH_DAO"].max_row,2)
        self.assertEqual(book["GIAY_TO_THEO_THUOC"].max_row,5)
        self.assertEqual(book["CHU_THE_VA_CHUNG_CU"].max_row,6)
        self.assertGreater(book["NGUON_TRA_CUU_CONG_KHAI"].max_row,2)
        book.close()

    def test_tender_documents_in_same_batch_not_misidentified_as_drug(self):
        report=analyze_auto([
            ("unknown.docx",word("THÔNG BÁO MỜI THẦU",
                 "Mã TBMT IB9876543210",
                 "Mã gói thầu BP1234567890")),
            ("amendment.docx",word("SỬA ĐỔI E-HSMT",
                 "Mã gói thầu BP1234567890")),
            ("second.docx",word("LÀM RÕ E-HSMT",
                 "Mã gói thầu BP9999999999")),
        ],enable_ocr=False)
        self.assertEqual(report["summary"]["products"],0)
        self.assertEqual(len(report["tender_documents"]),3)
        self.assertEqual(len(report["tender_index"]["by_package"]),2)
        self.assertFalse(report["verified"])
        for f in report["files"]:
            self.assertEqual(f["role_confidence"],"TAI_LIEU_GOI_THAU")
            self.assertEqual(f["kind"],"OTHER")
        book=load_workbook(BytesIO(export_auto_xlsx(report)),read_only=True)
        self.assertEqual(book["TAI_LIEU_GOI_THAU"].max_row,4)
        book.close()

if __name__=="__main__":
    unittest.main()
