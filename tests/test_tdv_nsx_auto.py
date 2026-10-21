"""NSX auto: end-to-end batch classification and source-bound linking tests.

Tests use generated fictitious products only. No real manufacturer data, Drive
uploads, Gemini calls or automatic compliance decisions.
"""
from __future__ import annotations

from io import BytesIO
import hashlib
import inspect
import unittest

from docx import Document
from openpyxl import Workbook, load_workbook

from tdv_nsx_auto import (
    analyze_auto, classify_document, export_auto_xlsx, STATUS_MATCH,
    STATUS_AMBIGUOUS, STATUS_UNKNOWN,
)
from tdv_nsx_workbench import render_nsx_workbench


def document(lines: list[str], pairs: list[tuple[str, str]] | None = None) -> bytes:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    if pairs:
        tbl = doc.add_table(rows=0, cols=2)
        for key, val in pairs:
            cells = tbl.add_row().cells
            cells[0].text = key
            cells[1].text = val
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def workbook(headers: list[str], data: list[list]) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Nguon NSX"
    sheet.append(headers)
    for row in data:
        sheet.append(row)
    target = BytesIO()
    book.save(target)
    return target.getvalue()


def products(*, shared_factory: bool = False) -> bytes:
    return workbook(
        ["Tên thuốc", "Mã sản phẩm", "Số đăng ký", "Hoạt chất", "Hàm lượng",
         "Nhà máy sản xuất", "Tên nhà sản xuất", "Dạng bào chế"],
        [
            ["Thuốc A", "P001", "VN-001", "Acenocoumarol", "1 mg",
             "Nhà máy B" if shared_factory else "Nhà máy A",
             "NSX Test", "Viên nén"],
            ["Thuốc B", "P002", "VN-002", "Acenocoumarol", "4 mg",
             "Nhà máy B", "NSX Test", "Viên nén"],
        ],
    )


def gmp(factory="Nhà máy B") -> bytes:
    return document(
        ["GIẤY CHỨNG NHẬN GMP"],
        [("Nhà máy sản xuất", factory),
         ("Số giấy chứng nhận GMP", "GMP-FAKE-01"),
         ("Phạm vi GMP", "Sản xuất dạng rắn"),
         ("Ngày hết hạn GMP", "30/12/2030")],
    )


def quotation(*, quote_ids: list[str] | None = None) -> bytes:
    return workbook(
        ["Mã sản phẩm", "Đơn giá báo giá", "Đơn vị tiền tệ"],
        [[ident, (i + 1) * 11000, "VND"] for i, ident
         in enumerate(quote_ids or ["P001", "P002"])],
    )


def registration(number="VN-002") -> bytes:
    return document(["GIẤY ĐĂNG KÝ LƯU HÀNH"],
                    [("Số đăng ký", number)])


def authorization() -> bytes:
    return document(
        ["GIẤY ỦY QUYỀN"],
        [("Mã sản phẩm", "P002"),
         ("Nhà sản xuất", "NSX Test"),
         ("Hiệu lực ủy quyền đến", "31/12/2030")],
    )


class AutoBatchTests(unittest.TestCase):
    def test_one_batch_classifies_without_user_labels(self):
        files = [
            ("a.xlsx", products()),
            ("b.docx", gmp()),
            ("c.xlsx", quotation()),
            ("d.docx", registration()),
            ("e.docx", authorization()),
        ]
        report = analyze_auto(files)
        self.assertEqual(report["summary"]["files"], 5)
        self.assertEqual(report["summary"]["parsed"], 5)
        self.assertEqual(report["summary"]["products"], 2)
        kinds = [x["kind"] for x in report["files"]]
        self.assertEqual(kinds, [
            "PRODUCT", "GMP_CERT", "QUOTE", "REGISTRATION", "AUTHORIZATION",
        ])
        self.assertTrue(all(x["verified"] is False for x in report["files"]))
        self.assertFalse(report["gmp_confirmed"])
        self.assertFalse(report["registration_confirmed"])
        self.assertFalse(report["quote_approved"])
        self.assertFalse(report["bid_eligible"])
        self.assertFalse(report["cloud_write_enabled"])
        self.assertEqual(len(report["links"]), 5)

        # GMP/registration/authorization uniquely link only to drug B.
        for kind in ("GMP_CERT", "REGISTRATION", "AUTHORIZATION"):
            self.assertEqual([l["product_id"] for l in report["links"]
                              if l["evidence_kind"] == kind], ["P0002"])
        # Pricing table links each price to the correct medicine, source row.
        prices = [l for l in report["links"] if l["evidence_kind"] == "QUOTE"]
        self.assertEqual(len(prices), 2)
        self.assertEqual({x["product_id"] for x in prices}, {"P0001", "P0002"})
        self.assertTrue(all("hàng " in x["evidence_origin"] for x in prices))
        self.assertEqual({x["evidence"]["QUOTE_UNIT_PRICE"]["value"]
                          for x in prices}, {"11000", "22000"})
        self.assertEqual(report["review_queue"], [])

    def test_registration_documents_seed_product_without_catalogue_or_manual_tags(self):
        """Two independent labeled original-like documents produce source-bound,
        UNVERIFIED candidates; they never self-certify their SĐK."""
        reg_a = document(
            ["GIẤY ĐĂNG KÝ LƯU HÀNH"],
            [("Tên thuốc", "Medicine A"), ("Số đăng ký", "VN-TEST-001"),
             ("Hoạt chất", "Ingredient One")])
        reg_b = document(
            ["GIẤY ĐĂNG KÝ LƯU HÀNH"],
            [("Tên thuốc", "Medicine B"), ("Số đăng ký", "VN-TEST-002"),
             ("Hoạt chất", "Ingredient Two")])
        batch = analyze_auto([
            ("register_a.docx", reg_a), ("register_b.docx", reg_b),
            ("gmp.docx", gmp()),
        ])
        self.assertEqual(batch["summary"]["products"], 2)
        self.assertEqual(batch["summary"]["linked"], 0)
        self.assertEqual({p["fields"]["PRODUCT_NAME"] for p in batch["products"]},
                         {"Medicine A", "Medicine B"})
        self.assertTrue(all(p["registration_document_source"]
                            for p in batch["products"]))
        self.assertTrue(all(p["evidence"]["REGISTRATION_NO"]["sha256"] == p["sha256"]
                            for p in batch["products"]))
        self.assertTrue(all(p["verified"] is False for p in batch["products"]))
        self.assertFalse(batch["registration_confirmed"])
        self.assertFalse(batch["gmp_confirmed"])
        self.assertFalse(batch["bid_eligible"])
        self.assertFalse(any(q["status"] == "CHUA_NHAN_DIEN_DANH_MUC"
                             for q in batch["review_queue"]))

    def test_native_pdf_two_line_registration_is_detected_without_catalogue(self):
        """PDFs with separated line labels produce only provisional records."""
        import pymupdf
        pdf = pymupdf.open()
        page = pdf.new_page()
        for index, line in enumerate([
            "Registration Certificate", "Drug Name", "Medicine C",
            "Registration Number", "VN-TEST-003", "Active Ingredient",
            "Ingredient Three",
        ]):
            page.insert_text((72, 72 + 27 * index), line, fontsize=11)
        content = pdf.tobytes()
        pdf.close()
        report = analyze_auto([("unlabeled_document.pdf", content)], enable_ocr=False)
        self.assertEqual(report["files"][0]["kind"], "REGISTRATION")
        self.assertEqual(report["summary"]["products"], 1)
        self.assertEqual(report["summary"]["linked"], 0)
        product = report["products"][0]
        self.assertEqual(product["fields"]["PRODUCT_NAME"], "Medicine C")
        self.assertEqual(product["fields"]["REGISTRATION_NO"], "VN-TEST-003")
        self.assertIn("Trang", product["evidence"]["PRODUCT_NAME"]["source"])
        self.assertTrue(product["registration_document_source"])
        self.assertFalse(report["registration_confirmed"])

    def test_registration_number_without_product_name_stays_source_only(self):
        standalone = analyze_auto([("register.docx", registration())])
        self.assertEqual(standalone["summary"]["products"], 0)
        self.assertEqual(standalone["summary"]["linked"], 0)
        self.assertTrue(any(q["status"] == "CHUA_NHAN_DIEN_DANH_MUC"
                            for q in standalone["review_queue"]))

    def test_gmp_for_shared_factory_never_arbitrarily_attaches_one_drug(self):
        report = analyze_auto([
            ("cat.xlsx", products(shared_factory=True)),
            ("cert.docx", gmp()),
        ])
        self.assertEqual(len(report["products"]), 2)
        self.assertEqual(len(report["links"]), 0)
        self.assertEqual(report["review_queue"][0]["status"], STATUS_AMBIGUOUS)
        self.assertEqual(set(report["review_queue"][0]["candidate_products"]),
                         {"P0001", "P0002"})

    def test_wrong_gmp_factory_is_not_marked_as_valid(self):
        report = analyze_auto([
            ("product.xlsx", products()),
            ("doc.docx", gmp("Nhà máy không liên quan")),
        ])
        self.assertEqual(len(report["links"]), 0)
        self.assertEqual(len(report["review_queue"]), 1)
        self.assertEqual(report["review_queue"][0]["status"], STATUS_UNKNOWN)

    def test_quote_with_conflicting_identifiers_does_not_autolink(self):
        bad = workbook(["Mã sản phẩm", "Số đăng ký", "Đơn giá báo giá"],
                       [["P001", "VN-002", 3500]])
        report = analyze_auto([
            ("products.xlsx", products()), ("quote.xlsx", bad),
        ])
        self.assertEqual(len(report["links"]), 0)
        self.assertEqual(report["files"][1]["kind"], "QUOTE")
        self.assertTrue(report["review_queue"])

    def test_role_not_inferred_from_filename_alone(self):
        text = document(["Hướng dẫn vận chuyển và bảo quản"],
                        [("Tên người liên hệ", "ABC")])
        cls = classify_document("GMP_CERT_QUOTE.docx", text)
        self.assertEqual(cls["kind"], "OTHER")
        self.assertEqual(cls["role_confidence"], "CHƯA_PHÂN_LOẠI")
        report = analyze_auto([
            ("document_a.xlsx", products()),
            ("GMP_CERT_QUOTE.docx", text),
        ])
        self.assertEqual(report["summary"]["products"], 2)
        self.assertEqual(report["summary"]["linked"], 0)
        self.assertTrue(any(x["status"] == "CAN_PHAN_LOAI_LAI"
                            for x in report["review_queue"]))

    def test_corrupt_file_is_isolated_not_silently_approved(self):
        report = analyze_auto([
            ("source.xlsx", products()),
            ("invalid.xlsx", b"INVALID"),
        ])
        self.assertEqual(report["summary"]["products"], 2)
        self.assertEqual(report["summary"]["parsed"], 1)
        self.assertTrue(any(x["role_confidence"] == "KHÔNG_ĐỌC_ĐƯỢC"
                            for x in report["files"]))
        self.assertFalse(report["verified"])

    def test_productless_certificate_causes_review_not_dummy_product(self):
        report = analyze_auto([("cert.docx", gmp())])
        self.assertEqual(report["summary"]["products"], 0)
        self.assertEqual(report["summary"]["linked"], 0)
        self.assertTrue(any(x["status"] == "CHUA_NHAN_DIEN_DANH_MUC"
                            for x in report["review_queue"]))
        self.assertFalse(report["gmp_confirmed"])

    def test_source_hashes_and_workbook_trace_for_each_quote_row(self):
        data = products()
        price = quotation()
        report = analyze_auto([
            ("cat.xlsx", data), ("some_table.xlsx", price),
        ])
        self.assertEqual(report["files"][0]["sha256"],
                         hashlib.sha256(data).hexdigest())
        wb = load_workbook(BytesIO(export_auto_xlsx(report)), data_only=False)
        self.assertEqual(wb.sheetnames, [
            "TU_DONG_NHAN_DIEN", "DANH_MUC_THUOC", "LIEN_KET_CHUNG_CU",
            "CAN_XAC_MINH", "TRUY_XUAT_TUNG_TRUONG", "GIOI_HAN_KIEM_TOAN",
        ])
        rows = list(wb["TRUY_XUAT_TUNG_TRUONG"].values)
        quote_prices = [r for r in rows[1:] if r[2] == "QUOTE_UNIT_PRICE"]
        self.assertEqual(len(quote_prices), 2)
        self.assertTrue(all("hàng" in str(r[4]) for r in quote_prices))
        self.assertTrue(all(r[5] == hashlib.sha256(price).hexdigest()
                            for r in quote_prices))
        wb.close()

    def test_large_quotation_is_not_silently_treated_as_complete(self):
        # Hitting a 250-row limit must not claim every price has been read.
        many = workbook(
            ["Mã sản phẩm", "Đơn giá báo giá"],
            [["P001", i + 10] for i in range(251)],
        )
        report = analyze_auto([
            ("products.xlsx", products()),
            ("prices.xlsx", many),
        ])
        self.assertEqual(report["files"][1]["kind"], "QUOTE")
        self.assertFalse(report["files"][1]["complete_read"])
        self.assertFalse(report["links"])
        self.assertTrue(any(x["status"] == "CHUA_DOC_DU_BANG_BAO_GIA"
                            for x in report["review_queue"]))
        self.assertTrue(any("giới hạn" in x for x in report["warnings"]))

    def test_preview_nsx_has_no_visible_legacy_demo_workflow(self):
        from pathlib import Path
        root = Path(__file__).resolve().parents[1]
        code = (root / "app.py").read_text(encoding="utf-8")
        nsx = code.split("    elif page == PAGES[2]:", 1)[1].split(
            "    elif page == PAGES[3]:", 1
        )[0]
        self.assertIn("render_nsx_workbench()", nsx)
        self.assertNotIn("DEMO", nsx)
        self.assertNotIn("file_uploader", nsx)
        self.assertNotIn("st.selectbox", nsx)

    def test_no_manual_upload_fields_or_product_picker_in_natural_ui(self):
        source = inspect.getsource(render_nsx_workbench)
        self.assertEqual(source.count("st.file_uploader("), 1)
        self.assertNotIn("st.selectbox(", source)
        self.assertNotIn("st.radio(", source)
        self.assertNotIn("st.multiselect(", source)
        self.assertIn("accept_multiple_files=True", source)
        self.assertNotIn("save_batch(", source)
        self.assertNotIn("drive_service(", source)


if __name__ == "__main__":
    unittest.main()
