"""Quality gate tests use entirely synthetic source records, no private bid data."""
import unittest
from io import BytesIO

from openpyxl import load_workbook, Workbook
from see_pro_quality_gate import audit_ehsmt, audit_nsx, audit_comparison
from see_pro_traceability import build_audit_workbook
from see_pro_workbench import analyze_uploads


class SeeProQualityGateQA(unittest.TestCase):
    def test_unregistered_upload_and_failure_dont_disappear(self):
        d = {
            "input_files":[{"name":"bad.pdf","sha256":"a"*64}],
            "documents":[], "rows":[],
            "failed_inputs":[{"name":"bad.pdf","error":"ValueError"}],
        }
        v=audit_ehsmt(d)
        self.assertEqual(v["files"],1)
        self.assertEqual(v["entries"],0)
        self.assertIn("TEP_CHUA_DUOC_GHI_NHAN",{x["type"] for x in v["issues"]})
        self.assertIn("LOI_DOC",{x["type"] for x in v["issues"]})
        self.assertFalse(v["accuracy_validated"])

    def test_field_citation_coverage_separate_from_present_value(self):
        row={
            "MA_THUOC":"G1.001","FILE_NGUON":"E-HSMT.pdf",
            "HOAT_CHAT_HSMT":"Acarbose","NONG_DO_HAM_LUONG":"50 mg",
            "EVIDENCE_KY_THUAT":{
                "NONG_DO_HAM_LUONG":{"pages":[11],"raw_tokens":["50 mg"]}
            }
        }
        data={"documents":[{"Tệp":"E-HSMT.pdf","SHA256":"a"*64,
                            "Đọc toàn bộ":True}],
              "rows":[row]}
        v=audit_ehsmt(data)
        self.assertEqual(v["coverage"]["HOAT_CHAT_HSMT"]["values"],1)
        self.assertEqual(v["coverage"]["HOAT_CHAT_HSMT"]["field_cited"],0)
        self.assertEqual(v["coverage"]["NONG_DO_HAM_LUONG"]["field_cited"],1)
        self.assertIsNone(v["accuracy_rate"])

    def test_duplicate_lot_codes_are_flagged_even_if_fields_match(self):
        v=audit_ehsmt({"documents":[{"Tệp":"e.pdf","Đọc toàn bộ":True}],
                       "rows":[{"MA_THUOC":"G1.010"},{"MA_THUOC":"G1.010"}]})
        self.assertIn("TRUNG_MA_LO",{x["type"] for x in v["issues"]})

    def test_nsx_with_87_rows_has_zero_sdk_citation_not_zero_legal_value(self):
        ps=[{
            "file":"list.xlsx","fields":{"PRODUCT_NAME":f"Drug {i}",
                                          "FACTORY":"Plant A"},
            "evidence":{
                "PRODUCT_NAME":{"file":"list.xlsx","source":f"Sheet1!A{i}"},
                "FACTORY":{"file":"list.xlsx","source":f"Sheet1!B{i}"},
            }, "conflicts": {},
        } for i in range(87)]
        v=audit_nsx({"products":ps,"files":[{"filename":"list.xlsx",
                                         "complete_read":True}]})
        self.assertEqual(v["entries"],87)
        self.assertEqual(v["coverage"]["PRODUCT_NAME"]["field_cited"],87)
        self.assertEqual(v["coverage"]["REGISTRATION_NO"]["field_cited"],0)
        self.assertFalse(v["legal_verified"])
        self.assertFalse(v["accuracy_validated"])

    def test_nsx_conflicting_cell_not_counted_as_trustworthy(self):
        p={
            "file":"nsx.xlsx", "fields":{"PRODUCT_NAME":"Drug A",
                                          "STRENGTH":"100 mg"},
            "evidence":{"PRODUCT_NAME":{"file":"nsx.xlsx","source":"A2"},
                        "STRENGTH":{"file":"nsx.xlsx","source":"B2"}},
            "conflicts":{"STRENGTH":[{"raw":"50 mg"}]}
        }
        v=audit_nsx({"products":[p],"files":[{"complete_read":False,
                                                 "filename":"nsx.xlsx"}]})
        self.assertEqual(v["coverage"]["STRENGTH"]["values"],1)
        self.assertEqual(v["coverage"]["STRENGTH"]["field_cited"],0)
        self.assertEqual(v["issue_count"],2)

    def test_green_missing_two_sided_source_detected(self):
        c={"status":"DAT_SO_BO","label":"Hoạt chất",
           "ehsmt_evidence":{"file":"e.pdf"},
           "nsx_evidence":{"file":"n.xlsx","location":"B2","sha256":"1"*64}}
        v=audit_comparison({"comparison":[{
            "MA_THUOC_EHSMT":"G1.001","criteria":[c]}]})
        self.assertEqual(v["issue_count"],1)
        self.assertEqual(v["issues"][0]["type"],"XANH_THIEU_DAU_VET")
        self.assertFalse(v["legal_verified"])

    def test_green_cited_two_sided_stays_only_preliminary(self):
        c={"status":"DAT_SO_BO","label":"Hoạt chất",
           "ehsmt_evidence":{"file":"e.pdf","location":"p. 5"},
           "nsx_evidence":{"file":"n.xlsx","location":"B2","sha256":"1"*64}}
        v=audit_comparison({"comparison":[{
            "MA_THUOC_EHSMT":"G1.001","criteria":[c]}]})
        self.assertEqual(v["issue_count"],0)
        self.assertEqual(v["status"],"TRICH_XUAT_SO_BO")
        self.assertIsNone(v["accuracy_rate"])

    def test_quality_sheet_is_exported_with_original_source(self):
        analysis={
            "input_files":[{"name":"e.pdf","sha256":"a"*64}],
            "documents":[{"Tệp":"e.pdf","SHA256":"a"*64,"Đọc toàn bộ":True}],
            "rows":[{"MA_THUOC":"G1.001","HOAT_CHAT_HSMT":"Acarbose"}],
            "warnings":[]
        }
        book=load_workbook(BytesIO(build_audit_workbook(analysis)),read_only=True)
        try:
            self.assertIn("KIEM_TOAN_CHAT_LUONG",book.sheetnames)
            a=list(book["KIEM_TOAN_CHAT_LUONG"].values)
            self.assertTrue(any("99,9999%" in str(row) for row in a))
            notes=list(book["PHAM_VI_KIEM_TOAN"].values)
            self.assertTrue(any(row[0]=="SỐ TỆP TIẾP NHẬN" and row[1]==1
                                for row in notes))
        finally:
            book.close()

    def test_actual_ehsmt_parser_keeps_input_count_on_invalid_document(self):
        # Parser may reject this malformed file. In any case count the upload.
        result=analyze_uploads([("invalid.pdf",b"this is not a PDF")])
        self.assertEqual(len(result["input_files"]),1)
        self.assertEqual(result["quality_gate"]["files"],1)
        self.assertFalse(result["quality_gate"]["accuracy_validated"])

if __name__=="__main__":
    unittest.main()
