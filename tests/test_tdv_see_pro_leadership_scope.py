"""Synthetic scoped leadership safety regressions; no real data or approvals."""
import unittest
from io import BytesIO
from openpyxl import load_workbook
from docx import Document
from tdv_see_pro_scope import bind_report, check_scope, switch_workspace
from tdv_see_pro_leadership import build_register, export_xlsx, export_docx, DAT, CHUA, KHONG, CHO

S={"project_id":"DA-TQ-01","package_id":"BP2600899572","ehsmt_revision_id":"R01"}

def sample():
    t=bind_report({"rows":[{"MA_THUOC":"G1.001","MA_GOI_THAU":S["package_id"]}]},S,[("E.pdf",b"ehsmt")])
    n=bind_report({"products":[{"uid":"P0001"}]},S,[("N.xlsx",b"nsx")])
    return t,n

def comparison(status="DAT_SO_BO",linked=True):
    return {"comparison":[{
        "MA_THUOC_EHSMT":"G1.001", "MA_THUOC_NSX":"P0001" if linked else "",
        "SO_UNG_VIEN_NSX":1 if linked else 0,"TEN_THUOC_NSX":"A",
        "criteria":[{"label":"Hàm lượng","ehsmt":"50mg","nsx":"100mg" if linked else "",
                     "status":status,
                     "ehsmt_evidence":{"file":"E.pdf","location":"Trang 30"},
                     "nsx_evidence":{"file":"N.xlsx","location":"Sheet A, B9"}}]
    }]}

class ScopedLeadershipQA(unittest.TestCase):
    def test_ram_switch_restores_only_own_project(self):
        a,b=sample()
        s={"tdv_see_pro_active_scope":dict(S),
           "tdv_see_pro_analysis":a,"tdv_see_pro_signature":"one",
           "tdv_nsx_auto_report":b,"tdv_nsx_auto_signature":"two"}
        other={**S,"project_id":"DA-TQ-02"}
        switch_workspace(s,other)
        self.assertNotIn("tdv_see_pro_analysis",s)
        self.assertNotIn("tdv_nsx_auto_report",s)
        c=bind_report({"rows":[]},other,[("B.pdf",b"b")])
        s["tdv_see_pro_analysis"]=c
        s["tdv_see_pro_signature"]="three"
        switch_workspace(s,S)
        self.assertEqual(s["tdv_see_pro_analysis"]["_tdv_scope"],S)
        self.assertEqual(s["tdv_nsx_auto_report"]["_tdv_scope"],S)
        self.assertEqual(s["tdv_see_pro_signature"],"one")
        self.assertEqual(s["tdv_nsx_auto_signature"],"two")
        switch_workspace(s,other)
        self.assertEqual(s["tdv_see_pro_signature"],"three")
        self.assertNotIn("tdv_nsx_auto_report",s)
    def test_workspace_rejects_missing_scope(self):
        with self.assertRaises(ValueError):switch_workspace({},{"project_id":"A"})
    def test_same_scope(self):
        a,b=sample()
        self.assertTrue(check_scope(S,a,b)[0])
    def test_different_project(self):
        a,b=sample()
        b["_tdv_scope"]["project_id"]="DA-TQ-02"
        self.assertFalse(check_scope(S,a,b)[0])
    def test_different_package(self):
        a,b=sample()
        b["_tdv_scope"]["package_id"]="BP9999999999"
        self.assertFalse(check_scope(S,a,b)[0])
    def test_different_revision(self):
        a,b=sample()
        b["_tdv_scope"]["ehsmt_revision_id"]="R02"
        self.assertFalse(check_scope(S,a,b)[0])
    def test_missing_manifest(self):
        a,b=sample()
        b["_tdv_consumed_sha256"]=[]
        self.assertFalse(check_scope(S,a,b)[0])
    def test_explicit_tender_package_conflict(self):
        a,b=sample()
        a["rows"][0]["MA_GOI_THAU"]="BP9999999999"
        self.assertFalse(check_scope(S,a,b)[0])
    def test_preliminary_green_not_final_pass(self):
        result=build_register(comparison(),scope_ok=True)
        self.assertEqual(result["summary"][DAT],0)
        self.assertEqual(result["summary"][CHUA],1)
    def test_preliminary_red_not_final_reject(self):
        result=build_register(comparison("KHONG_KHOP_CAN_XAC_MINH"),scope_ok=True)
        self.assertEqual(result["summary"][KHONG],0)
        self.assertIn("50mg",result["rows"][0]["nguyen_nhan"])
        self.assertIn("100mg",result["rows"][0]["nguyen_nhan"])
    def test_no_product_waiting_not_rejected(self):
        result=build_register(comparison(linked=False),scope_ok=True)
        self.assertEqual(result["summary"][CHO],1)
        self.assertEqual(result["summary"][CHUA],0)
    def test_no_scope_blocks(self):
        result=build_register(comparison(),scope_ok=False)
        self.assertTrue(result["blocked"])
        self.assertEqual(sum(result["summary"].values()),0)
    def test_duplicate_lot_rejected(self):
        r=comparison()
        r["comparison"].append(dict(r["comparison"][0]))
        with self.assertRaises(ValueError):build_register(r,scope_ok=True)
    def test_word_lists_drugs_reasons_and_scope(self):
        r=build_register(comparison("KHONG_KHOP_CAN_XAC_MINH"),scope_ok=True)
        doc=Document(BytesIO(export_docx(r,S)))
        whole=" ".join(x.text for x in doc.paragraphs)
        self.assertIn("G1.001",whole)
        self.assertIn("50mg",whole)
        self.assertIn("100mg",whole)
        self.assertIn(S["package_id"],whole)
        self.assertIn("CHƯA ĐẠT",whole)
    def test_excel_contains_reasons_and_proof(self):
        r=build_register(comparison("KHONG_KHOP_CAN_XAC_MINH"),scope_ok=True)
        book=load_workbook(BytesIO(export_xlsx(r,S)),read_only=True)
        self.assertEqual(book["03_CHUA_DAT"]["A2"].value,"G1.001")
        self.assertEqual(book["06_HAI_NGUON_LY_DO"]["D2"].value,"50mg")
        self.assertEqual(book["06_HAI_NGUON_LY_DO"]["E2"].value,"100mg")
        self.assertEqual(book["01_TONG_QUAN"]["B7"].value,1)
if __name__=="__main__":unittest.main()
