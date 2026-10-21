"""Regression for real-source comparison; all numbers below are synthetic QA fixtures.

NO business result/real supplier tender is inferred from these tests.
"""
import io
import inspect
import unittest

from docx import Document
from openpyxl import load_workbook

from tdv_leadership_report import compare_sources, export_leadership_xlsx, export_leadership_docx
from tdv_comparison_ui import render_comparison_view, render_leadership_view


def lot(code, ingredient, strength="", route="", form=""):
    return {"MA_THUOC":code,"FILE_NGUON":"EHSMT_TEST.xlsx",
        "HOAT_CHAT_HSMT":ingredient,"NONG_DO_HAM_LUONG":strength,
        "DUONG_DUNG":route,"DANG_BAO_CHE":form,
        "TRANG_VI_TRI":"Sheet X dong 5",
        "EVIDENCE":{key:{"source":"EHSMT_TEST.xlsx","sheet":"Sheet X",
                "row":5,"column":col,"raw":v} for col,(key,v) in enumerate([
            ("HOAT_CHAT_HSMT",ingredient),
            ("NONG_DO_HAM_LUONG",strength),
            ("DUONG_DUNG",route),("DANG_BAO_CHE",form)],1) if v},
    }


def product(uid="P0001",ingredient="Ingredient One",strength="0.5 g",
            form="Tablet",route="Oral"):
    fields={"PRODUCT_NAME":"Medicine Q","INGREDIENT":ingredient,"STRENGTH":strength,
            "DOSAGE_FORM":form,"ROUTE":route,"REGISTRATION_NO":"VN-FAKE-000"}
    return {"uid":uid,"file":"NSX_TEST.docx","sha256":"0"*64,
            "fields":fields,
            "evidence":{k:{"source":"Paragraph 2","raw":f"{k}: {v}",
                "sha256":"0"*64} for k,v in fields.items()},
            "verified":False}


def nsx(products):
    return {"products":products,"files":[{"filename":"NSX_TEST.docx","kind":"PRODUCT"}],
            "review_queue":[],"warnings":[],
            "traceability":{"products":[{
                "product_id":p["uid"],"product_name":p["fields"]["PRODUCT_NAME"],
                "registration_no":p["fields"]["REGISTRATION_NO"],
                "parties":{},"proofs":[], "origin_file":"NSX_TEST.docx",
                "origin":"Paragraph 2"} for p in products],
                "checklist":[],"provenance":[] }}


class RealComparisonQA(unittest.TestCase):
    def test_green_yellow_red_both_original_text_and_37_5_percent(self):
        tender={"rows":[
            lot("G1.001","Ingredient One","500 mg","Injection","Tablet"),
            lot("G1.002","Ingredient Two","20 mg","Oral","Tablet"),
        ],"documents":[{"Tệp":"EHSMT_TEST.xlsx"}],"warnings":[]}
        report=compare_sources(tender,nsx([product()]))
        summary=report["summary"]
        self.assertEqual(summary["ehsmt_lots"],2)
        self.assertEqual(summary["green_criteria_preliminary"],3)
        self.assertEqual(summary["red_mismatch_preliminary"],1)
        self.assertEqual(summary["yellow_criteria"],4)
        self.assertEqual(summary["matching_percent_preliminary"],37.5)
        self.assertEqual(summary["fully_verified"],0)
        r=report["comparison"][0]
        self.assertEqual(r["SO_UNG_VIEN_NSX"],1)
        self.assertEqual(r["TY_LE_KHOP_SO_BO"],75.0)
        self.assertEqual(r["criteria"][1]["ehsmt"],"500 mg")
        self.assertEqual(r["criteria"][1]["nsx"],"0.5 g")
        self.assertEqual(r["criteria"][1]["status"],"DAT_SO_BO")
        self.assertEqual(r["criteria"][2]["status"],"KHONG_KHOP_CAN_XAC_MINH")
        self.assertEqual(r["criteria"][2]["ehsmt_evidence"]["file"],"EHSMT_TEST.xlsx")
        self.assertEqual(r["criteria"][2]["nsx_evidence"]["file"],"NSX_TEST.docx")
        self.assertFalse(report["bid_eligible"])
        self.assertFalse(report["verified"])
        wb=load_workbook(io.BytesIO(export_leadership_xlsx(report)),read_only=True)
        self.assertIn("DOI_CHIEU_EHSMT_NSX",wb.sheetnames)
        self.assertEqual(wb["DOI_CHIEU_EHSMT_NSX"].max_row,3)
        wb.close()
        word=Document(io.BytesIO(export_leadership_docx(report)))
        self.assertTrue(any("BÁO CÁO LÃNH ĐẠO" in p.text for p in word.paragraphs))

    def test_ambiguous_sourcing_is_never_automatically_green(self):
        a=product("P0001")
        b=product("P0002")
        t={"rows":[lot("G1.001","Ingredient One","500 mg","Oral","Tablet")]}
        result=compare_sources(t,nsx([a,b]))
        row=result["comparison"][0]
        self.assertEqual(row["SO_UNG_VIEN_NSX"],2)
        self.assertEqual(row["SO_TIEU_CHI_DAT"],0)
        self.assertTrue(all(c["status"]=="CHUA_DU_CAN_CU" for c in row["criteria"]))

    def test_no_product_does_not_become_noncompliance(self):
        t={"rows":[lot("G1.001","Ingredient One","500 mg","Oral","Tablet")]}
        report=compare_sources(t,nsx([]))
        self.assertEqual(report["summary"]["red_mismatch_preliminary"],0)
        self.assertEqual(report["summary"]["yellow_criteria"],4)
        self.assertFalse(report["requirements_checked"])

    def test_preview_no_demo_and_no_fake_leader_edit(self):
        import pathlib
        app=(pathlib.Path(__file__).resolve().parents[1]/"app.py").read_text(encoding="utf-8")
        compare=app.split("    elif page == PAGES[5]:",1)[1].split("    elif page == PAGES[4]:",1)[0]
        leader=app.split("    elif page == PAGES[6]:",1)[1].split("    elif page == PAGES[-1]:",1)[0]
        self.assertIn("render_comparison_view()",compare)
        self.assertIn("render_leadership_view()",leader)
        self.assertNotIn("make_demo",compare+leader)
        self.assertNotIn("st.selectbox",compare+leader)
        self.assertNotIn("tdv_ehsmt_read",compare+leader)
        self.assertNotIn("tdv_nsx_read",compare+leader)
        combined=inspect.getsource(render_comparison_view)+inspect.getsource(render_leadership_view)
        self.assertNotIn("approve_bid",combined)
        self.assertNotIn("save_batch",combined)


if __name__=="__main__":
    unittest.main()
