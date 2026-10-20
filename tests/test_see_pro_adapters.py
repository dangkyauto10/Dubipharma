"""Regression tests for flexible SEE PRO input adapters."""
from __future__ import annotations

from io import BytesIO
from unittest.mock import patch
import unittest

from docx import Document
from see_pro_adapters import parse_tabular_document
from see_pro_workbench import analyze_uploads


HEAD = ("Mã phần/lô", "Mã thuốc", "Tên hoạt chất",
        "Số lượng", "Đơn giá (VND)", "Tổng giá trị (VND)")


class AdapterTests(unittest.TestCase):
    def test_csv_semicolon_and_vietnamese(self):
        data = (";".join(HEAD) + "\n" +
                "PP2600443936;GE841.001;Acenocoumarol;10;120;1200\n").encode("utf-8-sig")
        parsed = parse_tabular_document("parts_BP2600899572.csv", data)
        self.assertEqual(parsed["total_lots"], 1)
        self.assertEqual(parsed["rows"][0]["MA_THUOC"], "GE841.001")
        self.assertEqual(parsed["rows"][0]["MA_GOI_THAU"], "BP2600899572")
        self.assertEqual(parsed["rows"][0]["DON_GIA_KE_HOACH"], 120)
        self.assertEqual(parsed["rows"][0]["EVIDENCE"]["DON_GIA_KE_HOACH"]["source"],
                         "parts_BP2600899572.csv")
        self.assertFalse(parsed["verified"])

    def test_word_table_without_guessing(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=len(HEAD))
        for col, title in enumerate(HEAD):
            table.rows[0].cells[col].text = title
        row = table.add_row()
        for cell, value in zip(row.cells,
                               ("PP2600443937", "CP841.01", "Actiso", "100", "50", "5000")):
            cell.text = value
        out = BytesIO()
        doc.save(out)
        parsed = parse_tabular_document("parts_BP2600899590.docx", out.getvalue())
        self.assertEqual(parsed["total_lots"], 1)
        self.assertEqual(parsed["rows"][0]["MA_THUOC"], "CP841.01")
        self.assertEqual(parsed["rows"][0]["MA_GOI_THAU"], "BP2600899590")
        self.assertTrue(any("Word" in notice for notice in parsed["warnings"]))
        self.assertFalse(parsed["verified"])

    def test_docx_without_table_is_unclassified(self):
        doc = Document()
        doc.add_paragraph("Chỉ là thông báo - không có bảng.")
        out = BytesIO()
        doc.save(out)
        parsed = parse_tabular_document("notice.docx", out.getvalue())
        self.assertFalse(parsed["complete"])
        self.assertEqual(parsed["issues"][0]["type"], "NO_TABLE_FOUND")

    def test_xls_sheet_preserves_source_reference(self):
        class Sheet:
            name = "Danh sách phần lô"
            nrows = 2
            def row_values(self, i):
                return HEAD if i == 0 else (
                    "PP2600443936", "GE841.001", "Acenocoumarol", 100.0, 250.0, 25000.0)
        class Book:
            def sheets(self):
                return [Sheet()]
        with patch("xlrd.open_workbook", return_value=Book()):
            parsed = parse_tabular_document("legacy_BP2600899572.xls", b"mock-excel")
        self.assertEqual(parsed["total_lots"], 1)
        self.assertEqual(parsed["rows"][0]["DON_GIA_KE_HOACH"], 250)
        self.assertEqual(parsed["rows"][0]["FILE_NGUON"], "legacy_BP2600899572.xls")
        self.assertTrue(parsed["warnings"])

    def test_no_auto_values_or_cross_package_merge(self):
        csv = (";".join(HEAD) + "\n" +
               "PP2600443936;GE841.001;A;5;2;10\n").encode()
        report = analyze_uploads([("parts_BP2600899572.csv", csv)])
        self.assertEqual(len(report["rows"]), 1)
        self.assertEqual(report["rows"][0]["MA_GOI_THAU"], "BP2600899572")
        self.assertNotIn("DON_GIA_DU_THAU", report["rows"][0])


if __name__ == "__main__":
    unittest.main()
