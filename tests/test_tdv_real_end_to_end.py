"""Only synthetic source documents, exercise one shared SEE PRO pipeline, no DEMO output."""
from io import BytesIO
import unittest
from openpyxl import Workbook
from docx import Document

from see_pro_workbench import analyze_uploads
from tdv_nsx_auto import analyze_auto
from tdv_leadership_report import compare_sources


class EndToEndTDVIntegration(unittest.TestCase):
    def test_original_xlsx_and_registration_docx_to_comparison_without_manual_roles(self):
        book=Workbook()
        ws=book.active
        ws.title="summary"
        ws.append(["STT","Part","Code","Ingredient","Estimated","Guarantee"])
        ws.append([1,"PP1234567890","G1.001","Ingredient Q","10.000","1.000"])
        raw_xlsx=BytesIO()
        book.save(raw_xlsx)
        doc=Document()
        doc.add_paragraph("GIẤY ĐĂNG KÝ LƯU HÀNH")
        doc.add_paragraph("Tên thuốc: Medicine Q")
        doc.add_paragraph("Số đăng ký: VN-FAKE-Q")
        doc.add_paragraph("Hoạt chất: Ingredient Q")
        doc.add_paragraph("Hàm lượng: 500 mg")
        raw_docx=BytesIO()
        doc.save(raw_docx)
        tender=analyze_uploads([("example_ehsmt.xlsx",raw_xlsx.getvalue())])
        supplier=analyze_auto([("registration.docx",raw_docx.getvalue())],enable_ocr=False)
        report=compare_sources(tender,supplier)
        self.assertEqual(len(tender["rows"]),1)
        self.assertEqual(supplier["summary"]["products"],1)
        self.assertEqual(len(report["comparison"]),1)
        entry=report["comparison"][0]
        self.assertEqual(entry["MA_THUOC_EHSMT"],"G1.001")
        self.assertEqual(entry["SO_UNG_VIEN_NSX"],1)
        # The summary only states ingredient, NOT strength/dosage/route.
        # Other 3 checks must remain UNKNOWN rather than auto-approved.
        self.assertEqual(entry["SO_TIEU_CHI_DAT"],1)
        self.assertEqual(entry["SO_TIEU_CHI_CHUA_DAT"],3)
        self.assertEqual(entry["SO_TIEU_CHI_KHONG_KHOP"],0)
        self.assertEqual(entry["TY_LE_KHOP_SO_BO"],100.0)
        # One requirement actually present and source-linked: 1/1 preliminary.
        # The remaining 3 technical requirements were NOT in the summary.
        self.assertEqual(report["summary"]["linked_requirements"],1)
        self.assertEqual(report["summary"]["criteria_ehsmt_available"],1)
        self.assertEqual(report["summary"]["fully_verified"],0)
        self.assertEqual(entry["criteria"][0]["ehsmt"],"Ingredient Q")
        self.assertEqual(entry["criteria"][0]["nsx"],"Ingredient Q")
        self.assertEqual(entry["criteria"][0]["ehsmt_evidence"]["file"],"example_ehsmt.xlsx")
        self.assertEqual(entry["criteria"][0]["nsx_evidence"]["file"],"registration.docx")
        self.assertFalse(report["bid_eligible"])
        self.assertFalse(report["verified"])

if __name__=="__main__":
    unittest.main()
