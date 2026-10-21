"""Source-informed regression: official public notices vs direct NSX proof.

These are short structural *test fixtures* based on the public pages cited in
research/SEE_PRO_NSX_PUBLIC_EVIDENCE_20261021.md; they are NOT downloaded PDF
copies and are NOT proof of any actual manufacturer's valid GMP/registration.
"""
from __future__ import annotations

from io import BytesIO
import hashlib
import unittest

import pymupdf
from docx import Document
from openpyxl import Workbook

from tdv_nsx_auto import analyze_auto, classify_document
from tdv_nsx_public_guard import (
    detect_nonprimary_document, PUBLIC_GMP_LIST,
    PUBLIC_GMP_BLANK_TEMPLATE, PUBLIC_GMP_NONCOMPLIANCE, PUBLIC_GMP_GUIDANCE,
    PUBLIC_REGISTRATION_DECISION, QUOTE_INVITATION,
    QUOTE_BLANK_TEMPLATE,
)


def docx(lines: list[str], pairs: list[tuple[str, str]] = ()) -> bytes:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    if pairs:
        table = document.add_table(rows=0, cols=2)
        for a, b in pairs:
            cells = table.add_row().cells
            cells[0].text = a
            cells[1].text = b
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


def english_pdf(lines: list[str]) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    for idx, value in enumerate(lines):
        page.insert_text((50, 50 + 22 * idx), value, fontsize=11)
    raw = doc.tobytes()
    doc.close()
    return raw


def catalogue() -> bytes:
    wb = Workbook()
    sh = wb.active
    sh.title = "Danh muc"
    sh.append(["Mã sản phẩm", "Tên thuốc", "Hoạt chất", "Hàm lượng",
               "Nhà máy sản xuất", "Số đăng ký"])
    sh.append(["P001", "Thuốc A", "Ciprofloxacin", "500 mg", "Nhà máy A", "VN-101"])
    target = BytesIO()
    wb.save(target)
    return target.getvalue()


class PublicEvidenceSafetyTests(unittest.TestCase):
    def test_dav_official_pdf_heading_variant_is_not_factory_certificate(self):
        # The heading was observed in DAV's real public GMP list dated
        # 24/10/2025; this fixture is NOT the PDF binary or issuer certification.
        # URL: https://dav.gov.vn/danh-sach-cac-cssx-trong-nuoc-dat-tieu-chuan-nguyen-tac-gmp-gmp-bao-bi-duoc-pham-cap-nhat-toi-ngay-24102025-n5075.html
        official_heading = (
            "DANH SÁCH CÁC CƠ SỞ TRONG NƯỚC "
            "ĐẠT TIÊU CHUẨN NGUYÊN TẮC GMP"
        )
        detection = detect_nonprimary_document([
            {"source": "Trang 1", "text": official_heading},
            {"source": "Trang 1", "text": "Cập nhật tới ngày 24/10/2025"},
        ])
        self.assertIsNotNone(detection)
        self.assertEqual(detection["class"], PUBLIC_GMP_LIST)
        self.assertFalse(detection["automatic_primary_link_allowed"])
        classified = classify_document(
            "official_list_heading.docx",
            docx([official_heading, "Cập nhật tới ngày 24/10/2025"]),
        )
        self.assertEqual(classified["kind"], "OTHER")
        self.assertEqual(classified["public_document_class"], PUBLIC_GMP_LIST)

    def test_dav_public_gmp_register_not_an_individual_certificate(self):
        raw = docx([
            "PHỤ LỤC 1: DANH SÁCH CƠ SỞ SẢN XUẤT THUỐC NƯỚC NGOÀI ĐẠT TIÊU CHUẨN GMP",
            "Theo công văn của Cục Quản lý Dược",
        ], [
            ("Nhà máy sản xuất", "Nhà máy A"),
            ("Số giấy chứng nhận GMP", "GMP-TEST-01"),
        ])
        result = classify_document("foreign_gmp.docx", raw)
        self.assertEqual(result["kind"], "OTHER")
        self.assertEqual(result["public_document_class"], PUBLIC_GMP_LIST)
        self.assertFalse(result["verified"])
        report = analyze_auto([
            ("products.xlsx", catalogue()),
            ("foreign_gmp.docx", raw),
        ], enable_ocr=False)
        self.assertEqual(report["summary"]["products"], 1)
        self.assertEqual(report["summary"]["linked"], 0)
        self.assertEqual(report["summary"]["reference_documents"], 1)
        self.assertEqual(report["review_queue"][0]["status"], "TAI_LIEU_THAM_KHAO")
        self.assertEqual(
            report["files"][1]["sha256"], hashlib.sha256(raw).hexdigest()
        )

    def test_dav_registration_decision_not_a_single_drug_certificate(self):
        raw = docx([
            "Quyết định số 855/QĐ-QLD về việc ban hành danh mục "
            "180 thuốc sản xuất trong nước được cấp giấy đăng ký lưu hành tại Việt Nam",
            "Danh mục thuốc được cấp số đăng ký lưu hành."
        ], [("Số đăng ký", "VN-101")])
        record = classify_document("decision.docx", raw)
        self.assertEqual(record["public_document_class"], PUBLIC_REGISTRATION_DECISION)
        self.assertEqual(record["kind"], "OTHER")
        report = analyze_auto([("cat.xlsx", catalogue()),
                               ("decision.docx", raw)], enable_ocr=False)
        self.assertEqual(report["summary"]["linked"], 0)

    def test_public_quote_invitation_is_not_supplier_quote(self):
        raw = docx([
            "THƯ MỜI BÁO GIÁ VÀ CUNG CẤP HỢP ĐỒNG",
            "Kính gửi: Các công ty sản xuất, kinh doanh dược phẩm trên toàn quốc.",
            "Bệnh viện trân trọng kính mời Quý Công ty cung cấp báo giá các thuốc.",
        ], [
            ("Mã sản phẩm", "P001"),
            ("Đơn giá báo giá", "12345"),
        ])
        record = classify_document("quotation_request.docx", raw)
        self.assertEqual(record["public_document_class"], QUOTE_INVITATION)
        self.assertEqual(record["kind"], "OTHER")
        report = analyze_auto([("cat.xlsx", catalogue()),
                               ("quotation_request.docx", raw)], enable_ocr=False)
        self.assertEqual(report["summary"]["linked"], 0)

    def test_blank_quote_template_cannot_be_used_as_real_quote(self):
        raw = docx([
            "PHỤ LỤC II: MẪU BÁO GIÁ",
            "Kèm theo thư mời báo giá của cơ sở y tế",
            "BẢNG CHÀO GIÁ THUỐC",
            "Nhà thầu: ...................",
            "Đơn giá (có VAT): ..................",
        ])
        record = classify_document("quote_template.docx", raw)
        self.assertEqual(record["kind"], "OTHER")
        # Invitation appears early too, so conservative tie-breaking MUST
        # still prevent this document becoming a confirmed QUOTE.
        self.assertIn(record["public_document_class"],
                      (QUOTE_BLANK_TEMPLATE, QUOTE_INVITATION))
        self.assertEqual(analyze_auto([
            ("cat.xlsx", catalogue()), ("quote_template.docx", raw)
        ], enable_ocr=False)["summary"]["linked"], 0)

    def test_ema_noncompliance_statement_not_a_positive_gmp_cert(self):
        raw = english_pdf([
            "STATEMENTS OF NON-COMPLIANCE WITH GMP",
            "Regulatory authorities record GMP non-compliance findings",
            "Manufacturing site: Site A",
            "GMP Certificate Number: FAKE-01",
        ])
        record = classify_document("eudragmdp_notice.pdf", raw)
        self.assertEqual(record["public_document_class"],
                         PUBLIC_GMP_NONCOMPLIANCE)
        self.assertEqual(record["kind"], "OTHER")

    def test_ema_gmp_guidance_not_individual_certificate(self):
        raw = english_pdf([
            "The issue and update of GMP certificates",
            "European Medicines Agency: procedural guidance",
            "GMP certificates are issued following inspections",
        ])
        record = classify_document("guideline.pdf", raw)
        self.assertEqual(record["public_document_class"], PUBLIC_GMP_GUIDANCE)
        self.assertEqual(record["kind"], "OTHER")

    def test_dav_domestic_gmp_list_not_individual_factory_certificate(self):
        raw = docx([
            "Danh sách các CSSX trong nước đạt tiêu chuẩn nguyên tắc GMP, "
            "GMP bao bì dược phẩm (Cập nhật tới ngày 24/10/2025)",
            "Danh sách cơ sở sản xuất trong nước",
        ], [("Nhà máy sản xuất", "Nhà máy A")])
        result = classify_document("dav_domestic_list.docx", raw)
        self.assertEqual(result["public_document_class"], PUBLIC_GMP_LIST)
        self.assertEqual(result["kind"], "OTHER")

    def test_ema_blank_union_format_is_not_issued_gmp_certificate(self):
        raw = english_pdf([
            "Union format for a GMP certificate",
            "(LETTERHEAD OF COMPETENT AUTHORITY)",
            "Certificate No: ____ / ____ / ____",
            "CERTIFICATE OF GMP COMPLIANCE OF A MANUFACTURER",
        ])
        result = classify_document("union_format.pdf", raw)
        self.assertEqual(result["public_document_class"], PUBLIC_GMP_BLANK_TEMPLATE)
        self.assertEqual(result["kind"], "OTHER")

    def test_ema_interpretation_guidance_is_not_issued_certificate(self):
        raw = english_pdf([
            "Interpretation of the Union format for GMP certificate",
            "The purpose is to provide guidance for regulators",
            "This text is not an individual factory's GMP certificate",
        ])
        result = classify_document("interpretation.pdf", raw)
        self.assertEqual(result["public_document_class"], PUBLIC_GMP_GUIDANCE)

    def test_genuine_certificate_type_is_still_recognised_provisionally(self):
        raw = docx([
            "GIẤY CHỨNG NHẬN GMP",
        ], [
            ("Nhà máy sản xuất", "Nhà máy A"),
            ("Số giấy chứng nhận GMP", "GMP-TEST-01"),
            ("Phạm vi GMP", "Dạng rắn"),
        ])
        record = classify_document("site_proof.docx", raw)
        self.assertEqual(record["kind"], "GMP_CERT")
        self.assertNotIn("public_document_class", record)
        report = analyze_auto([("cat.xlsx", catalogue()),
                               ("site_proof.docx", raw)], enable_ocr=False)
        self.assertEqual(report["summary"]["linked"], 1)
        self.assertFalse(report["gmp_confirmed"])
        self.assertFalse(report["verified"])

    def test_real_quotation_type_is_not_suppressed(self):
        raw = docx(["BÁO GIÁ"],
                   [("Mã sản phẩm", "P001"), ("Đơn giá báo giá", "10000")])
        record = classify_document("supplier_letter.docx", raw)
        self.assertEqual(record["kind"], "QUOTE")
        self.assertNotIn("public_document_class", record)
        report = analyze_auto([("cat.xlsx", catalogue()),
                               ("supplier_letter.docx", raw)], enable_ocr=False)
        self.assertEqual(report["summary"]["linked"], 1)
        self.assertFalse(report["quote_approved"])

    def test_file_name_alone_cannot_turn_document_into_reference(self):
        lines = [{"source": "Trang 1", "text": "GIẤY CHỨNG NHẬN GMP"},
                 {"source": "Trang 1", "text": "Nhà máy sản xuất: Nhà máy A"}]
        self.assertIsNone(detect_nonprimary_document(lines))

    def test_no_primary_match_for_blended_invitation_and_price_table(self):
        raw = docx(["THƯ MỜI BÁO GIÁ", "PHỤ LỤC: MẪU BÁO GIÁ"],
                   [("Nhà máy sản xuất", "Nhà máy A"),
                    ("Mã sản phẩm", "P001"), ("Đơn giá báo giá", "9999")])
        out = analyze_auto([("cat.xlsx", catalogue()),
                            ("invitation.docx", raw)], enable_ocr=False)
        self.assertEqual(out["summary"]["reference_documents"], 1)
        self.assertEqual(out["summary"]["linked"], 0)


if __name__ == "__main__":
    unittest.main()
