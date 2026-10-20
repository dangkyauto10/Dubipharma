"""SEE PRO structured intake: source-first, no guessing, no external writes."""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO
import hashlib
import re
import unicodedata
from typing import Any

MAX_BYTES = 30 * 1024 * 1024
MAX_ROWS = 20_000
LOT_ID = re.compile(r"^PP\d{10}$", re.I)
PACKAGE_ID = re.compile(r"BP\d{10}(?!\d)", re.I)
FIELDS = {
    "STT": ("stt", "so thu tu"),
    "MA_PHAN_LO": ("ma phan/lo", "ma phan lo", "ma lo", "ma phan"),
    "MA_THUOC": ("ma thuoc", "ma mat hang", "ma danh muc"),
    "HOAT_CHAT_HSMT": ("ten hoat chat", "hoat chat", "ten thanh phan cua thuoc", "thanh phan thuoc"),
    "NHOM_THUOC": ("nhom tckt", "nhom thuoc", "nhom ky thuat"),
    "NONG_DO_HAM_LUONG": ("nong do hoac ham luong", "nong do ham luong", "nong do", "ham luong"),
    "DUONG_DUNG": ("duong dung",),
    "DANG_BAO_CHE": ("dang bao che",),
    "DON_VI_TINH": ("don vi tinh",),
    "SO_LUONG_CHINH": ("so luong", "so luong moi thau"),
    "DON_GIA_KE_HOACH": ("don gia vnd", "don gia", "don gia ke hoach"),
    "THANH_TIEN_KE_HOACH": ("tong gia tri vnd", "thanh tien", "tong gia tri"),
    "SO_LUONG_MUA_THEM": ("so luong tuy chon mua them",),
    "TY_LE_MUA_THEM": ("ty le phan tram", "ty le mua them"),
    "GIA_TRI_MUA_THEM": ("gia tri tuy chon mua them",),
    "GIA_DU_TOAN_COT_GOC": ("gia du toan uoc tinh tung lo",),
}
HEADER_LOOKUP = {alias: key for key, aliases in FIELDS.items() for alias in aliases}
REQUIRED = ("MA_PHAN_LO", "MA_THUOC", "HOAT_CHAT_HSMT", "SO_LUONG_CHINH",
            "DON_GIA_KE_HOACH", "THANH_TIEN_KE_HOACH")


def normalize(text: Any) -> str:
    s = unicodedata.normalize("NFKD", str(text or "").replace("Đ", "D").replace("đ", "d"))
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", s)).strip()


def _decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        if isinstance(value, float) and (value != value or value in (float("inf"), -float("inf"))):
            return None
        return Decimal(str(value))
    txt = str(value).strip().replace(" ", "")
    if re.fullmatch(r"-?\d{1,3}(?:\.\d{3})+(?:,\d+)?", txt):
        txt = txt.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?\d+(?:,\d+)?", txt):
        txt = txt.replace(",", ".")
    elif not re.fullmatch(r"-?\d+(?:\.\d+)?", txt):
        return None
    try:
        return Decimal(txt)
    except InvalidOperation:
        return None


def identify_document(filename: str, content: bytes) -> dict:
    """Classify before attempting lot parsing; a notice is not a lot table."""
    if not content or len(content) > MAX_BYTES:
        raise ValueError("Tệp rỗng hoặc vượt giới hạn 30 MB")
    package = PACKAGE_ID.search(filename)
    info = {"type": "UNKNOWN", "package_id": package.group(0).upper() if package else "",
            "evidence": {}, "warnings": []}
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext in ("xlsx", "xls", "csv"):
        info["type"] = "TABULAR_CANDIDATE"
    elif ext == "pdf":
        import pymupdf
        with pymupdf.open(stream=content, filetype="pdf") as doc:
            if doc.needs_pass:
                raise ValueError("PDF yêu cầu mật khẩu")
            text = "\n".join((doc[i].get_text("text", sort=True) or "") for i in range(min(3, len(doc))))
            if "THÔNG BÁO MỜI THẦU" in text.upper() or "THONG BAO MOI THAU" in normalize(text).upper():
                info["type"] = "THONG_BAO_MOI_THAU"
                patterns = {
                    "MA_TBMT": r"Mã TBMT\s+(IB\d+)",
                    "MA_GOI_THAU": r"Mã gói thầu\s+(BP\d+)",
                    "SO_PHAN_LO": r"Số lượng phần\s*\(lô\)\s+(\d+)",
                    "NGAY_DONG_THAU": r"Thời điểm đóng thầu\s+([^\n]+)",
                    "SO_TIEN_BAO_DAM_DU_THAU": r"Số tiền đảm bảo dự thầu\s+([^\n]+)",
                }
                for key, pattern in patterns.items():
                    m = re.search(pattern, text, flags=re.I)
                    if m:
                        info["evidence"][key] = {
                            "value": m.group(1).strip(), "source": f"{filename}:trang 1", "verified": False}
                if "MA_GOI_THAU" in info["evidence"]:
                    info["package_id"] = info["evidence"]["MA_GOI_THAU"]["value"]
            else:
                anchor = re.compile(r"\bPP\d{10}\b.{0,120}\bG\d+\.\d{2,4}\b", re.I | re.S)
                found = bool(anchor.search(text))
                if not found:
                    for index in range(3, min(doc.page_count, 80)):
                        if anchor.search(doc[index].get_text("text", sort=True) or ""):
                            found = True
                            break
                info["type"] = "BANG_EHSMT_PDF" if found else "PDF_KHAC_HOAC_CAN_OCR"
    elif ext in ("docx", "doc", "png", "jpg", "jpeg", "tif", "tiff"):
        info["type"] = "TAI_LIEU_HOAC_ANH_CAN_PHAN_LOAI"
    return info


def _header(row: tuple) -> dict[str, int]:
    found = {}
    for index, cell in enumerate(row):
        field = HEADER_LOOKUP.get(normalize(cell))
        if field and field not in found:
            found[field] = index
    return found


def _value(row: tuple, fields: dict[str, int], key: str) -> Any:
    idx = fields.get(key)
    return row[idx] if idx is not None and idx < len(row) else None


def parse_excel_lots(filename: str, content: bytes) -> dict:
    """Read identified header columns and preserve per-cell evidence."""
    if not filename.lower().endswith(".xlsx"):
        raise ValueError("Bộ phân tích bảng có cấu trúc yêu cầu XLSX")
    if not content or len(content) > MAX_BYTES:
        raise ValueError("Tệp rỗng hoặc vượt giới hạn 30 MB")
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    package = PACKAGE_ID.search(filename)
    bp = package.group(0).upper() if package else ""
    rows, issues, sheets = [], [], []
    total_visited = 0
    try:
        for sheet in workbook.worksheets:
            header_at, fields = None, {}
            footer_total, amount_sum = None, Decimal(0)
            matched_count = 0
            for line_no, cells in enumerate(sheet.iter_rows(values_only=True), 1):
                total_visited += 1
                if total_visited > MAX_ROWS:
                    issues.append({"type": "ROW_LIMIT", "file": filename, "sheet": sheet.title})
                    break
                if len(cells) > 80:
                    issues.append({"type": "COLUMN_LIMIT", "file": filename, "sheet": sheet.title,
                                   "row": line_no})
                    break
                if header_at is None:
                    candidates = _header(cells)
                    if len(set(REQUIRED) & set(candidates)) >= 5:
                        header_at, fields = line_no, candidates
                    elif line_no > 40:
                        break
                    continue
                part_raw = _value(cells, fields, "MA_PHAN_LO")
                drug_raw = _value(cells, fields, "MA_THUOC")
                if part_raw is None and drug_raw is None:
                    possible_total = _decimal(_value(cells, fields, "THANH_TIEN_KE_HOACH"))
                    if possible_total is not None:
                        footer_total = possible_total
                    continue
                part, drug = str(part_raw or "").strip(), str(drug_raw or "").strip()
                if not LOT_ID.fullmatch(part) or not drug:
                    issues.append({"type": "INVALID_LOT_ROW", "file": filename,
                                   "sheet": sheet.title, "row": line_no, "part": part[:70],
                                   "drug": drug[:70]})
                    continue
                values = {key: _value(cells, fields, key) for key in fields}
                numeric_keys = ("STT", "SO_LUONG_CHINH", "DON_GIA_KE_HOACH",
                                "THANH_TIEN_KE_HOACH", "SO_LUONG_MUA_THEM",
                                "TY_LE_MUA_THEM", "GIA_TRI_MUA_THEM")
                for key in numeric_keys:
                    if key in values:
                        values[key] = _decimal(values[key])
                lot = {"MA_GOI_THAU": bp, "FILE_NGUON": filename,
                       "SHEET": sheet.title, "DONG_NGUON": line_no,
                       "MA_PHAN_LO": part, "MA_THUOC": drug,
                       "TRANG_THAI": "TRICH_XUAT_CHUA_XAC_MINH", "verified": False}
                lot.update(values)
                lot["EVIDENCE"] = {
                    field: {"source": filename, "sheet": sheet.title, "row": line_no,
                            "column": fields[field] + 1,
                            "raw": str(_value(cells, fields, field)), "verified": False}
                    for field in fields if _value(cells, fields, field) is not None}
                missing = [k for k in REQUIRED if lot.get(k) is None or lot.get(k) == ""]
                if missing:
                    issues.append({"type": "MISSING_REQUIRED_FIELD", "file": filename,
                                   "row": line_no, "drug": drug, "fields": missing})
                qty, price, total = (lot.get(k) for k in (
                    "SO_LUONG_CHINH", "DON_GIA_KE_HOACH", "THANH_TIEN_KE_HOACH"))
                if all(isinstance(v, Decimal) for v in (qty, price, total)):
                    difference = qty * price - total
                    if abs(difference) > Decimal("0.51"):
                        issues.append({"type": "MATH_TOTAL_MISMATCH", "file": filename,
                                       "row": line_no, "drug": drug,
                                       "difference": str(difference)})
                more = lot.get("SO_LUONG_MUA_THEM")
                optional_value = lot.get("GIA_TRI_MUA_THEM")
                if all(isinstance(v, Decimal) for v in (more, price, optional_value)):
                    difference = more * price - optional_value
                    if abs(difference) > Decimal("0.51"):
                        issues.append({"type": "OPTIONAL_TOTAL_MISMATCH", "file": filename,
                                       "row": line_no, "drug": drug,
                                       "difference": str(difference)})
                if isinstance(total, Decimal):
                    amount_sum += total
                for key in numeric_keys:
                    value = lot.get(key)
                    if isinstance(value, Decimal):
                        lot[key] = int(value) if value == value.to_integral_value() else str(value)
                rows.append(lot)
                matched_count += 1
            if footer_total is not None and abs(amount_sum - footer_total) > Decimal("0.51"):
                issues.append({"type": "FOOTER_SUM_MISMATCH", "file": filename,
                               "sheet": sheet.title, "calculated": str(amount_sum),
                               "footer": str(footer_total)})
            sheets.append({"sheet": sheet.title, "header_row": header_at,
                           "columns": fields, "lot_rows": matched_count,
                           "source_footer_total": str(footer_total) if footer_total is not None else None,
                           "detail_sum": str(amount_sum)})
            if total_visited > MAX_ROWS:
                break
    finally:
        workbook.close()
    seen_part, seen_drug = set(), set()
    for item in rows:
        if item["MA_PHAN_LO"] in seen_part:
            issues.append({"type": "DUPLICATE_PART", "part": item["MA_PHAN_LO"]})
        if item["MA_THUOC"] in seen_drug:
            issues.append({"type": "DUPLICATE_DRUG", "drug": item["MA_THUOC"]})
        seen_part.add(item["MA_PHAN_LO"])
        seen_drug.add(item["MA_THUOC"])
    return {"file": filename, "package_id": bp, "total_lots": len(rows),
            "rows": rows, "issues": issues, "sheets": sheets,
            "sha256": hashlib.sha256(content).hexdigest(), "verified": False,
            "complete": total_visited <= MAX_ROWS and any(s["header_row"] for s in sheets)}


def reconcile_notice(notice: dict, excel_results: list[dict]) -> list[dict]:
    """Match by exact package code; never merge different BP packages."""
    code = notice.get("package_id", "")
    if not code:
        return [{"type": "NOTICE_PACKAGE_UNKNOWN"}]
    found = [r for r in excel_results if r.get("package_id") == code]
    if not found:
        return [{"type": "NO_MATCHING_LOT_LIST", "package_id": code}]
    raw = (notice.get("evidence", {}).get("SO_PHAN_LO") or {}).get("value")
    if raw and str(raw).isdigit():
        actual = sum(item["total_lots"] for item in found)
        if int(raw) != actual:
            return [{"type": "NOTICE_LOT_COUNT_MISMATCH", "package_id": code,
                     "expected": int(raw), "actual": actual}]
    return []
