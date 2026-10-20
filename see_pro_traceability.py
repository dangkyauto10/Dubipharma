"""SEE PRO professional, provenance-preserving audit workbook.

The workbook is a READ-ONLY extraction/review aid. Nothing in this module
approves bids, certifies GMP, creates an offer price or writes cloud records.
"""
from __future__ import annotations

from io import BytesIO
import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "163B67"
BLUE = "EAF3FC"
PALE = "F5F9FE"
RED = "FBEAEA"
AMBER = "FFF3D8"
GREEN = "E8F5E8"
WHITE = "FFFFFF"
GRAY = "667588"
THIN = Side(style="hair", color="D5E1EF")

FIRST_COLUMNS = (
    "MA_GOI_THAU", "MA_PHAN_LO", "MA_THUOC", "HOAT_CHAT_HSMT",
    "NHOM_THUOC", "NONG_DO_HAM_LUONG", "DUONG_DUNG",
    "DANG_BAO_CHE", "DON_VI_TINH", "SO_LUONG_CHINH",
    "DON_GIA_KE_HOACH", "THANH_TIEN_KE_HOACH",
    "GIA_UOC_TINH_PHAN", "SO_LUONG_MUA_THEM",
    "TY_LE_MUA_THEM", "GIA_TRI_MUA_THEM",
    "TRANG_THAI", "KHOP_SO_HOC", "KHOP_MA_PHAN_LO",
    "FILE_NGUON", "SHEET", "DONG_NGUON", "TRANG_VI_TRI",
    "TRANG_BANG_KY_THUAT", "TRICH_DAN",
)
PRIVATE_COLUMNS = {"EVIDENCE", "EVIDENCE_KY_THUAT", "verified"}


def safe_cell(value: Any) -> str | int | float | bool | None:
    """Preserve numeric values; never turn evidence into an Excel formula."""
    if value is None:
        return None
    if isinstance(value, (dict, list, tuple, set)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, (bool, int, float)):
        return value
    output = str(value)
    if len(output) > 32_000:
        output = output[:31_950] + " …[CẮT THEO GIỚI HẠN EXCEL]"
    # Leading whitespace + executable marker is also dangerous when opened.
    if output.lstrip().startswith(("=", "+", "-", "@")) or output.startswith(("\t", "\r", "\n")):
        return "'" + output
    return output


def _sheet(wb: Workbook, name: str, headings: list[str]):
    ws = wb.create_sheet(name)
    ws.append(headings)
    ws.freeze_panes = "D2" if len(headings) >= 4 else "A2"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 36
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    return ws


def _format(ws, *, status_field: str | None = None):
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    headers = {str(cell.value): i for i, cell in enumerate(ws[1], 1)}
    status_idx = headers.get(status_field) if status_field else None
    wide = ("HOAT_CHAT", "TRICH", "CANH_BAO", "CHUNG_CU", "GIẢI", "GIAI_THICH", "NGUYEN_VAN", "DỮ LIỆU")
    money = ("DON_GIA_KE_HOACH", "THANH_TIEN_KE_HOACH", "GIA_UOC_TINH_PHAN", "GIA_TRI_MUA_THEM")
    for i, cell in enumerate(ws[1], 1):
        key = str(cell.value or "")
        width = 37 if any(x in key.upper() for x in wide) else 23
        if key in ("MA_THUOC", "MA_PHAN_LO", "MA_GOI_THAU", "STT", "ROW", "COL"):
            width = 19
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        state = str(row[status_idx - 1].value or "").upper() if status_idx else ""
        color = (RED if any(x in state for x in ("FAIL", "KHONG_DUOC_KET_LUAN", "KHONG_KHOP"))
                 else AMBER if any(x in state for x in ("CHUA", "CAN_", "WARNING", "CAN_XAC_MINH"))
                 else GREEN if "DA_KIEM_DUYET" in state else None)
        for cell in row:
            cell.font = Font(name="Calibri", size=10, color="26374A")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE)
            if cell.column == status_idx and color:
                cell.fill = PatternFill("solid", fgColor=color)
                cell.font = Font(name="Calibri", size=10, color="8C381F" if color == RED else "71521A", bold=True)
            key = str(ws.cell(1, cell.column).value or "")
            if key in money and isinstance(cell.value, (float, int)) and not isinstance(cell.value, bool):
                cell.number_format = '#,##0.## "đ"'
    ws.print_options.horizontalCentered = True


def _source_lines(record: dict):
    """Yield evidence per field with exact filename/page/sheet/row."""
    file_name = record.get("FILE_NGUON", "")
    package = record.get("MA_GOI_THAU", "")
    drug = record.get("MA_THUOC", "")
    part = record.get("MA_PHAN_LO", "")
    seen_any = False
    for field, evidence in (record.get("EVIDENCE") or {}).items():
        if not isinstance(evidence, dict):
            continue
        seen_any = True
        yield (
            package, part, drug, field,
            evidence.get("source", file_name), evidence.get("sheet", ""),
            evidence.get("row", ""), evidence.get("column", ""),
            "", evidence.get("raw", ""), "CHUA_XAC_MINH",
        )
    for field, evidence in (record.get("EVIDENCE_KY_THUAT") or {}).items():
        if not isinstance(evidence, dict):
            continue
        seen_any = True
        yield (
            package, part, drug, field, file_name, "",
            "", "",
            ", ".join(map(str, evidence.get("pages") or [])),
            " ".join(map(str, evidence.get("raw_tokens") or [])),
            "CHUA_XAC_MINH",
        )
    if not seen_any:
        # Preserve the weaker, generic PDF/Word provenance without pretending
        # it is evidence for every possible field.
        yield (
            package, part, drug, "TRICH_DOAN_TONG_QUAT",
            file_name, record.get("SHEET", ""), record.get("DONG_NGUON", ""),
            "", record.get("TRANG_VI_TRI", ""),
            record.get("TRICH_DAN", ""), "CHUA_XAC_MINH",
        )


def build_audit_workbook(analysis: dict) -> bytes:
    """Structured multi-sheet provenance export from existing SEE PRO result."""
    wb = Workbook()
    wb.remove(wb.active)
    rows = analysis.get("rows") or []
    documents = analysis.get("documents") or []
    warnings = analysis.get("warnings") or []

    extras = sorted({key for row in rows for key in row if key not in FIRST_COLUMNS
                     and key not in PRIVATE_COLUMNS})
    columns = [key for key in FIRST_COLUMNS if any(key in row for row in rows)] + extras
    if not columns:
        columns = ["MA_THUOC", "MA_PHAN_LO", "TRANG_THAI", "FILE_NGUON"]

    summary = _sheet(wb, "DANH_MUC_PHAN_LO", columns)
    for record in rows:
        summary.append([safe_cell(record.get(k)) for k in columns])

    doc_fields = list(dict.fromkeys(k for entry in documents for k in entry))
    if not doc_fields:
        doc_fields = ["Tệp", "Loại", "SHA256", "Xác minh"]
    doc_sheet = _sheet(wb, "NGUON_TAI_LIEU", doc_fields)
    for doc in documents:
        doc_sheet.append([safe_cell(doc.get(k)) for k in doc_fields])

    trace = _sheet(wb, "TRUY_XUAT_CHUNG_CU", [
        "MA_GOI_THAU", "MA_PHAN_LO", "MA_THUOC", "TRUONG_DU_LIEU",
        "TÊN_TỆP", "SHEET", "DÒNG", "CỘT", "TRANG_PDF",
        "TRICH_DAN_GOC", "TRANG_THAI",
    ])
    for record in rows:
        for values in _source_lines(record):
            trace.append([safe_cell(value) for value in values])

    alerts = _sheet(wb, "CANH_BAO_CAN_XEM", ["STT", "CẢNH BÁO", "TRẠNG THÁI"])
    for i, warning in enumerate(warnings, 1):
        alerts.append([i, safe_cell(warning), "CAN_XAC_MINH"])

    notes = _sheet(wb, "PHAM_VI_KIEM_TOAN", ["NỘI DUNG", "GIỚI HẠN / GHI CHÚ"])
    for title, detail in (
        ("PHẠM VI", "Bảng kiểm kê và đối chiếu chứng cứ, CHƯA PHẢI kết quả thẩm định pháp lý/ kỹ thuật."),
        ("NGUỒN CHỨNG CỨ", "Các dòng lưu tệp, Sheet, ô, trang và trích đoạn để kiểm tra lại bản gốc."),
        ("XÁC MINH", "Khớp số học hoặc tồn tại trích dẫn KHÔNG tương đương đã chứng minh đáp ứng E-HSMT."),
        ("GIÁ", "Giá kế hoạch/ước tính trong E-HSMT KHÔNG phải giá dự thầu hoặc giá chào."),
        ("AI", "Nếu dùng AI, mọi đề xuất phải kiểm tra trích dẫn và được người có thẩm quyền duyệt."),
        ("SỐ BẢN GHI TRÍCH XUẤT", len(rows)),
        ("SỐ TỆP TIẾP NHẬN", len(documents)),
        ("SỐ CẢNH BÁO", len(warnings)),
    ):
        notes.append([safe_cell(title), safe_cell(detail)])

    for ws in wb:
        _format(ws, status_field="TRANG_THAI" if ws.title == "DANH_MUC_PHAN_LO"
                else "TRẠNG THÁI" if ws.title == "CANH_BAO_CAN_XEM" else None)
    output = BytesIO()
    wb.save(output)
    return output.getvalue()
