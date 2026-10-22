"""TDV QC7 ONLINE · read-only real NSX/product evidence intake.

First stage: extract labeled, source-addressable facts and triage *links*.
Never approve GMP, registration, authority, bid eligibility or quoted prices.
No uploads to third-party AI, Google Sheets, Drive, GitHub or disk.
"""
from __future__ import annotations

from collections import defaultdict
from io import BytesIO
import hashlib
import json
import re
import unicodedata
from typing import Any

from see_pro_multiformat import ACCEPTED, extract_file
from tdv_nsx_proof_hints import extract_labeled_hints, diagnose_proof_source, collect_proof_contexts

MAX_DOCUMENTS = 24
MAX_LINE_COUNT = 10_000
MAX_FIELD_LENGTH = 350

KINDS = {
    "PRODUCT": "Hồ sơ sản phẩm",
    "GMP_CERT": "Giấy chứng nhận GMP",
    "GMP_EXPLANATION": "Giải trình GMP (không phải chứng nhận)",
    "REGISTRATION": "Giấy đăng ký lưu hành",
    "QUOTE": "Báo giá của nhà cung cấp",
    "AUTHORIZATION": "Giấy ủy quyền",
    "OTHER": "Tài liệu khác",
}
LABELS = {
    "PRODUCT_ID": ("ma san pham", "product id", "ma thuoc noi bo", "ma hang hoa"),
    "REGISTRATION_NO": ("so dang ky", "so dang ky luu hanh", "sdk",
                        "registration number", "registration no"),
    "PRODUCT_NAME": ("ten thuoc", "ten san pham", "trade name", "drug name"),
    "SUPPLIER": ("don vi cung cap", "nha cung cap", "ten nha cung cap", "supplier",
                 "distributor", "distributing company"),
    "BIDDER": ("ten nha thau", "don vi du thau", "bidder", "bidder name"),
    "REGISTRATION_HOLDER": ("co so dang ky", "co so dang ky thuoc",
                            "don vi dang ky", "chu so huu so dang ky",
                            "marketing authorization holder", "mah"),
    "INGREDIENT": ("hoat chat", "ten hoat chat", "thanh phan hoat chat",
                   "active ingredient", "ingredient"),
    "STRENGTH": ("ham luong", "nong do ham luong", "nong do hoac ham luong",
                 "strength", "concentration"),
    "DOSAGE_FORM": ("dang bao che", "dosage form"),
    "ROUTE": ("duong dung", "route of administration", "route"),
    "MANUFACTURER": ("ten nha san xuat", "nha san xuat", "manufacturer",
                     "manufacturer name", "don vi san xuat"),
    "FACTORY": ("nha may san xuat", "ten nha may", "co so san xuat",
                "manufacturing site", "factory", "factory name"),
    "GMP_NUMBER": ("so giay chung nhan gmp", "so chung nhan gmp",
                   "gmp certificate number", "gmp number"),
    "GMP_SCOPE": ("pham vi gmp", "gmp scope"),
    "GMP_EXPIRY": ("ngay het han gmp", "gmp expiry", "gmp expiration date"),
    "QUOTE_UNIT_PRICE": ("don gia bao gia", "gia bao gia", "don gia chao",
                         "quoted unit price", "unit price"),
    "QUOTE_CURRENCY": ("dong tien bao gia", "don vi tien te", "currency"),
    "QUOTE_VALID_UNTIL": ("hieu luc bao gia den", "bao gia co hieu luc den",
                         "quotation valid until"),
    "AUTH_EXPIRY": ("hieu luc uy quyen den", "ngay het han uy quyen"),
}
LOOKUP = {key: field for field, aliases in LABELS.items() for key in aliases}


def normalize(value: Any) -> str:
    value = unicodedata.normalize("NFKD", str(value or "").replace("Đ", "D").replace("đ", "d"))
    value = "".join(c for c in value if not unicodedata.combining(c)).casefold()
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", value)).strip()


def _cell(text: str) -> str:
    # For XLSX extract_document serializes 1:label | 2:value.
    return re.sub(r"^\d{1,2}:", "", text.strip()).strip()


def _evidence_pair(raw: str) -> list[tuple[str, str]]:
    """Return only exact *labeled* values; no speculative entity extraction."""
    bits = [_cell(x) for x in raw.split("|") if _cell(x)]
    pairs: list[tuple[str, str]] = []
    if len(bits) >= 2:
        # Do NOT interpret header rows with two or more field labels as values:
        # "Mã sản phẩm | Tên thuốc | Hoạt chất | ..." is NOT a product record.
        if sum(bool(LOOKUP.get(normalize(bit.strip(":： ")))) for bit in bits) >= 2:
            return []
        # A key-value row may have "Tên hoạt chất | Acenocoumarol".
        label = LOOKUP.get(normalize(bits[0].strip(":： ")))
        if label:
            pairs.append((label, bits[1]))
            return pairs
    for part in bits:
        found = re.match(r"^([^:：]{2,80})\s*[:：]\s*(.{1,600})$", part)
        if not found:
            continue
        field = LOOKUP.get(normalize(found.group(1)))
        if field:
            pairs.append((field, found.group(2).strip()))
    return pairs


def _consecutive_pdf_fields(lines: list[dict]):
    """Recognize strictly adjacent labeled two-line fields on one PDF page.

    The raw page and both original lines remain in evidence. A heading is
    never accepted as a value; distant lines, other pages or guessed product
    names from filenames are excluded. Final fields remain unverified.
    """
    allowed = {"PRODUCT_NAME", "REGISTRATION_NO", "INGREDIENT",
               "STRENGTH", "DOSAGE_FORM", "MANUFACTURER", "FACTORY",
               "SUPPLIER", "BIDDER", "REGISTRATION_HOLDER"}
    for index in range(len(lines) - 1):
        here, following = lines[index], lines[index + 1]
        if here["source"] != following["source"]:
            continue
        heading = str(here["text"]).strip().strip(":： ")
        field = LOOKUP.get(normalize(heading))
        if field not in allowed:
            continue
        value = str(following["text"]).strip()
        if (not 1 <= len(value) <= 140
            or LOOKUP.get(normalize(value.strip(":： ")))
            or re.match(r"^[^:：]{2,80}\s*[:：]", value)):
            continue
        yield field, value, {
            "source": here["source"] + " (hai dòng liền kề)",
            "raw": str(here["text"])[:150] + " | " + value,
        }


def _read_one(filename: str, kind: str, raw: bytes, ocr: bool, *,
              selected_product_index: int | None = None,
              source_document: dict | None = None) -> dict:
    if kind not in KINDS:
        raise ValueError("Loại tài liệu không hợp lệ")
    if not filename or filename.rsplit(".", 1)[-1].lower() not in ACCEPTED:
        raise ValueError("Chưa hỗ trợ loại tệp")
    source = source_document if source_document is not None else extract_file(filename, raw, ocr=ocr)
    fields, evidence, conflicts = {}, {}, defaultdict(list)
    warnings = list(source["warnings"])
    if not source["complete"]:
        warnings.append("Chưa đọc đủ văn bản hoặc kết quả OCR chưa thể xác minh")
    truncated = len(source["lines"]) > MAX_LINE_COUNT
    if truncated:
        warnings.append("Tài liệu nhiều dòng; chỉ xem 10.000 dòng đầu")

    for line in source["lines"][:MAX_LINE_COUNT]:
        for field, raw_value in _evidence_pair(line["text"]):
            value = raw_value.strip()
            if not value or len(value) > MAX_FIELD_LENGTH:
                warnings.append(f"Trường {field}: giá trị rỗng/quá dài, cần xem tệp gốc")
                continue
            citation = {
                "file": filename,
                "sha256": source["sha256"],
                "source": line["source"],
                "raw": line["text"][:700],
                "value": value,
                "verified": False,
            }
            if field in fields and normalize(fields[field]) != normalize(value):
                conflicts[field].append(citation)
            elif field not in fields:
                fields[field] = value
                evidence[field] = citation
    if source["type"] == "pdf":
        # Native PDF often places a field label and its value on separate
        # consecutive text lines, unlike Word table pairs or key:value text.
        for field, value, location in _consecutive_pdf_fields(source["lines"][:MAX_LINE_COUNT]):
            citation = {
                "file": filename, "sha256": source["sha256"],
                "source": location["source"], "raw": location["raw"][:700],
                "value": value, "verified": False,
            }
            if field in fields and normalize(fields[field]) != normalize(value):
                conflicts[field].append(citation)
            elif field not in fields:
                fields[field] = value
                evidence[field] = citation
    candidates: list[dict] = []
    chosen_index = None
    selection_required = False
    catalogue_complete = True
    if kind == "PRODUCT" and source["type"] in ("xlsx", "xls", "csv", "docx"):
        from tdv_nsx_adaptive import extract_product_rows
        catalogue = extract_product_rows(filename, raw)
        candidates = catalogue["candidates"]
        catalogue_complete = bool(catalogue["complete"])
        warnings.extend(catalogue["warnings"])
        if len(candidates) > 1 and selected_product_index is None:
            selection_required = True
            warnings.append(
                f"Có {len(candidates)} sản phẩm trong bảng; phải chọn đúng một dòng trước khi liên kết GMP/báo giá")
        if selected_product_index is not None:
            if type(selected_product_index) is not int or not 0 <= selected_product_index < len(candidates):
                raise ValueError("Chỉ số sản phẩm được chọn không còn hợp lệ, vui lòng đọc lại hồ sơ")
            chosen_index = selected_product_index
        elif len(candidates) == 1:
            chosen_index = 0
        if chosen_index is not None:
            picked = candidates[chosen_index]
            for name, value in picked["fields"].items():
                citation = picked["evidence"][name]
                if name in fields and normalize(fields[name]) != normalize(value):
                    conflicts[name].append(citation)
                elif name not in fields:
                    fields[name] = value
                    evidence[name] = citation
            warnings.append("Bảng sản phẩm đọc theo tiêu đề cột; thông tin CHƯA XÁC MINH với tệp gốc")
    if conflicts:
        warnings.append("Có trường ghi nhiều giá trị không thống nhất: " +
                        ", ".join(sorted(conflicts)))
    # Supplementary, unverified candidates never alter fields/evidence/conflicts.
    # They are shown for document review but cannot become link identifiers.
    proof_hints = extract_labeled_hints(kind, source, fields)
    proof_diagnostic = diagnose_proof_source(kind, source, fields, proof_hints)
    proof_contexts = collect_proof_contexts(kind, source) if proof_diagnostic and proof_diagnostic.get("fields_without_hint") else []
    if len(proof_hints) >= 18:
        warnings.append(
            "Đã tới giới hạn 18 gợi ý SĐK/GMP có nhãn; "
            "danh sách không chứng minh đã tìm hết thông tin trong tệp."
        )
    return {
        "kind": kind, "kind_label": KINDS[kind], "filename": filename,
        "sha256": source["sha256"], "type": source["type"], "pages": source["pages"],
        "complete_read": source["complete"] and not truncated and catalogue_complete,
        "fields": fields, "evidence": evidence, "conflicts": dict(conflicts),
        "proof_hints": proof_hints,
        "proof_diagnostic": proof_diagnostic,
        "proof_contexts": proof_contexts,
        "catalogue_candidates": candidates,
        "catalogue_selected_index": chosen_index,
        "catalogue_selection_required": selection_required,
        "warnings": warnings, "verified": False,
        "gmp_confirmed": False, "registration_confirmed": False,
        "quote_approved": False, "authorization_confirmed": False,
    }


ANCHORS = {
    "GMP_CERT": ("FACTORY",),
    "GMP_EXPLANATION": ("FACTORY",),
    "QUOTE": ("PRODUCT_ID", "REGISTRATION_NO"),
    "REGISTRATION": ("REGISTRATION_NO",),
    "AUTHORIZATION": ("PRODUCT_ID", "REGISTRATION_NO", "MANUFACTURER"),
    "OTHER": (),
}


def _link(product: dict, document: dict) -> dict:
    """Flag obvious conflicts; equal strings = only a candidate association."""
    kind = document["kind"]
    anchors = ANCHORS[kind]
    p_fields, d_fields = product["fields"], document["fields"]
    if product.get("catalogue_selection_required"):
        return {
            "filename": document["filename"], "kind": kind,
            "status": "CHUA_CHON_SAN_PHAM",
            "matching_fields": [], "conflicting_fields": [],
            "unavailable_fields": list(anchors),
            "reasons": ["Danh mục có nhiều sản phẩm; chưa chọn dòng được đối chiếu"],
            "source_sha256": document["sha256"], "verified": False,
            "gmp_confirmed": False, "quote_approved": False, "bid_eligible": False,
        }
    ambiguous = set(product["conflicts"]) | set(document["conflicts"])
    conflict, matches, missing = [], [], []
    for field in anchors:
        if field in ambiguous:
            missing.append(field + " (mâu thuẫn trong tài liệu)")
            continue
        expected, actual = p_fields.get(field), d_fields.get(field)
        if not expected or not actual:
            missing.append(field)
        elif normalize(expected) != normalize(actual):
            conflict.append(field)
        else:
            matches.append(field)
    # For Quote, either product ID or SĐK must agree; for GMP, factory is essential.
    necessary = {
        "GMP_CERT": ["FACTORY"],
        "GMP_EXPLANATION": ["FACTORY"],
        "QUOTE": ["PRODUCT_ID", "REGISTRATION_NO"],
        "REGISTRATION": ["REGISTRATION_NO"],
        "AUTHORIZATION": ["PRODUCT_ID", "REGISTRATION_NO"],
    }.get(kind, [])
    if conflict:
        status = "XUNG_DOT_DINH_DANH"
    elif kind == "OTHER":
        status = "CHUA_DU_CHUNG_CU"
    elif kind == "AUTHORIZATION" and "MANUFACTURER" not in matches:
        status = "CHUA_DU_CHUNG_CU"
    elif not any(f in matches for f in necessary):
        status = "CHUA_DU_CHUNG_CU"
    elif not product["complete_read"] or not document["complete_read"]:
        status = "CAN_KIEM_TRA_BAN_GOC"
    else:
        status = "LIEN_KET_SO_BO_CAN_XAC_MINH"

    reasons = []
    if conflict:
        reasons.append("Trường định danh khác nhau: " + ", ".join(conflict))
    if status == "CHUA_DU_CHUNG_CU":
        reasons.append("Chưa đủ bằng chứng liên kết chính xác sản phẩm/nhà máy")
    if kind in ("GMP_CERT", "GMP_EXPLANATION"):
        reasons.append("Cần kiểm tra cơ quan cấp, thời hạn, phạm vi GMP và nhà máy")
    if kind == "GMP_EXPLANATION":
        reasons.append("Công văn giải trình không thay thế chứng nhận GMP")
    if kind == "REGISTRATION":
        reasons.append("Chưa xác thực SĐK/hiệu lực tại nguồn quản lý dược chính thức")
    if kind == "QUOTE":
        reasons.append("Giá báo chỉ là thông tin nguồn; không phải giá dự thầu được duyệt")
    if kind == "AUTHORIZATION":
        reasons.append("Chưa kiểm chứng tư cách chủ thể, thời hạn và chuỗi ủy quyền")
    return {
        "filename": document["filename"], "kind": kind, "status": status,
        "matching_fields": matches, "conflicting_fields": conflict,
        "unavailable_fields": missing, "reasons": reasons,
        "source_sha256": document["sha256"], "verified": False,
        "gmp_confirmed": False, "quote_approved": False, "bid_eligible": False,
    }


def analyze_nsx_documents(
    files: list[tuple[str, str, bytes]], *, enable_ocr: bool = False,
    selected_product_index: int | None = None,
) -> dict:
    """One product dossier + zero/many independent evidence documents.

    Files: (kind, filename, exact bytes). A document that cannot be parsed is
    reported, never silently made into a valid source.
    """
    if not files or len(files) > MAX_DOCUMENTS:
        raise ValueError(f"Chọn từ 1 đến {MAX_DOCUMENTS} tài liệu")
    if sum(kind == "PRODUCT" for kind, _, _ in files) != 1:
        raise ValueError("Mỗi lượt phân tích cần đúng một hồ sơ sản phẩm")
    records = []
    warnings = []
    for kind, filename, data in files:
        try:
            parsed = _read_one(
                filename, kind, data, enable_ocr,
                selected_product_index=selected_product_index if kind == "PRODUCT" else None,
            )
            records.append(parsed)
            warnings.extend(f"{filename}: {w}" for w in parsed["warnings"])
        except Exception as exc:
            # Untrusted documents can fail with parser-specific exceptions
            # such as zipfile.BadZipFile. One corrupt file must not crash
            # ingestion of the remaining evidence. Do not echo the raw
            # exception message because it can contain document content.
            error_code = type(exc).__name__
            records.append({
                "kind": kind, "filename": filename,
                "sha256": hashlib.sha256(data).hexdigest(),
                "error": error_code, "verified": False,
            })
            warnings.append(f"{filename}: KHÔNG ĐỌC ĐƯỢC ({error_code}); cần kiểm tra tệp gốc")
    product = next(r for r in records if r["kind"] == "PRODUCT")
    attachments = [r for r in records if r["kind"] != "PRODUCT"]
    if "error" in product:
        warnings.append("Không đọc được hồ sơ sản phẩm: mọi liên kết đều chưa đủ chứng cứ")
        links = [{
            "filename": r["filename"], "kind": r["kind"],
            "status": "CHUA_DU_CHUNG_CU", "reasons": ["Hồ sơ sản phẩm không đọc được"],
            "verified": False, "gmp_confirmed": False, "quote_approved": False,
            "bid_eligible": False,
        } for r in attachments]
    else:
        links = [
            (_link(product, r) if "error" not in r else {
                "filename": r["filename"], "kind": r["kind"],
                "status": "CHUA_DU_CHUNG_CU",
                "reasons": ["Tài liệu chưa đọc được"],
                "verified": False, "gmp_confirmed": False, "quote_approved": False,
                "bid_eligible": False,
            }) for r in attachments
        ]
    seen = defaultdict(list)
    for r in records:
        seen[r["sha256"]].append(r["filename"])
    duplicates = [names for names in seen.values() if len(names) > 1]
    if duplicates:
        warnings.append("Có tệp trùng nội dung SHA256: " + str(duplicates))
    return {
        "mode": "NSX_REAL_SOURCE_READ_ONLY",
        "selected_product_index": selected_product_index,
        "product": product, "documents": attachments, "links": links,
        "warnings": warnings,
        "summary": {
            "files_received": len(files),
            "files_parsed": sum("error" not in r for r in records),
            "links_provisional": sum(l["status"] == "LIEN_KET_SO_BO_CAN_XAC_MINH" for l in links),
            "links_conflicted": sum(l["status"] == "XUNG_DOT_DINH_DANH" for l in links),
            "links_missing": sum(l["status"] in ("CHUA_DU_CHUNG_CU", "CAN_KIEM_TRA_BAN_GOC") for l in links),
        },
        "verified": False, "gmp_confirmed": False, "quote_approved": False,
        "registration_confirmed": False, "bid_eligible": False,
    }


def export_nsx_xlsx(report: dict) -> bytes:
    """Audit workbook with no formula injection; reference values only."""
    from openpyxl import Workbook
    from see_pro_traceability import safe_cell, _format

    wb = Workbook()
    ws = wb.active
    ws.title = "DANH_SACH_TAI_LIEU"
    ws.append(["Vai trò", "Tệp nguồn", "Định dạng", "SHA256", "Trang", "Đọc đủ", "Xác minh"])
    sources = [report["product"]] + report["documents"]
    for item in sources:
        ws.append([safe_cell(x) for x in (
            item["kind"], item["filename"], item.get("type", ""),
            item["sha256"], item.get("pages"), item.get("complete_read"),
            "CHUA",
        )])

    checks = wb.create_sheet("DOI_CHIEU_DINH_DANH")
    checks.append(["Tệp", "Loại tài liệu", "Trạng thái", "Khớp mã", "Mâu thuẫn",
                   "Thiếu", "Giải thích", "Đã duyệt"])
    for item in report["links"]:
        checks.append([safe_cell(x) for x in (
            item["filename"], item["kind"], item["status"],
            item.get("matching_fields"), item.get("conflicting_fields"),
            item.get("unavailable_fields"), item.get("reasons"), "KHONG",
        )])

    prov = wb.create_sheet("CHUNG_CU_TUNG_TRUONG")
    prov.append(["Tệp", "Vai trò", "Tên trường", "Giá trị", "Đường dẫn nguồn",
                 "SHA256", "Trích đoạn gốc", "Xác minh"])
    for item in sources:
        for field, evidence in item.get("evidence", {}).items():
            prov.append([safe_cell(x) for x in (
                item["filename"], item["kind"], field, evidence["value"],
                evidence["source"], evidence["sha256"], evidence["raw"], "CHUA",
            )])
        for field, conflict_evidence in item.get("conflicts", {}).items():
            for evidence in conflict_evidence:
                prov.append([safe_cell(x) for x in (
                    item["filename"], item["kind"], field + " (MÂU THUẪN)",
                    evidence["value"], evidence["source"], evidence["sha256"],
                    evidence["raw"], "CHUA",
                )])

    meta = wb.create_sheet("GIOI_HAN_KET_LUAN")
    meta.append(["Nội dung", "Kết luận"])
    for a, b in (
        ("Loại báo cáo", "Trích xuất và liên kết nguồn sơ bộ, CHƯA PHÊ DUYỆT"),
        ("GMP", "CHƯA xác minh cơ quan cấp, phạm vi, nhà máy hoặc thời hạn"),
        ("SĐK", "CHƯA xác minh hiệu lực đăng ký lưu hành"),
        ("Báo giá", "Giá nguồn KHÔNG phải đơn giá chào thầu đã phê duyệt"),
        ("Ủy quyền", "CHƯA xác minh tư cách, hiệu lực và chuỗi ủy quyền"),
        ("Google Drive/Sheets", "KHÔNG tự ghi dữ liệu"),
        ("Đối chiếu", "Mã khớp chuỗi KHÔNG chứng minh đáp ứng kỹ thuật hoặc pháp lý"),
        ("Cảnh báo", report["warnings"]),
    ):
        meta.append([safe_cell(a), safe_cell(b)])
    for sh in wb:
        sh.freeze_panes = "B2"
        _format(sh, status_field="Trạng thái" if sh.title == "DOI_CHIEU_DINH_DANH" else None)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
